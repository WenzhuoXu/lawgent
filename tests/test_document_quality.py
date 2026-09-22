"""The deterministic half of the visual review, which no model can skip.

Rendering a deck and asking the model to look at it is advisory: nothing knows
whether it looked. Every case here is a defect measured on a real generated
artifact in outputs/harness_reviews/2026-09-21-document-visual-review.md — a
shape written past the canvas edge, a statute title that cannot fit its cell, a
label set below the readable floor, navy text on its own navy fill, a slide
carrying 8% ink — and each one is arithmetic, so it is caught before the reader
sees the file rather than after.

Fixtures are built with python-pptx and python-docx directly so a test states
its own defect instead of inheriting one from a writer.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.oxml.ns import qn as w_qn
from docx.shared import Pt as DocxPt
from lxml import etree
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from legal_helper.documents.quality import (
    BODY_FONT_FLOOR_PT,
    NOTE_FONT_FLOOR_PT,
    RENDER_CHECKS,
    Finding,
    QualityReport,
    contrast_floor,
    contrast_ratio,
    estimate_text_block,
    lint_artifact,
    lint_docx,
    lint_pptx,
    lint_renders,
    page_ink,
)

#: Sixty characters either way — the Han string is twice as wide, which is the
#: whole reason the estimator measures columns rather than characters.
HAN_BODY = "合同当事人应当遵循诚实信用原则" * 4
LATIN_BODY = "The parties shall observe good faith in all dealings herein."

NAVY = RGBColor(0x1F, 0x36, 0x64)
NEAR_NAVY = RGBColor(0x2A, 0x44, 0x76)
PALE = RGBColor(0xE8, 0xEE, 0xF6)


def _deck():
    """An empty 13.333x7.5in slide, the canvas every writer here targets."""
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    return prs, prs.slides.add_slide(prs.slide_layouts[6])


def _textbox(slide, x, y, w, h, text="", *, pt=18.0, wrap=True):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    return _set_text(shape, text, pt=pt, wrap=wrap)


def _card(slide, x, y, w, h, text="", *, pt=18.0, fill=PALE, wrap=True):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    return _set_text(shape, text, pt=pt, wrap=wrap)


def _set_text(shape, text, *, pt, wrap):
    """Write text the way a writer does: explicit size, wrap, no autofit.

    python-pptx's own add_textbox ships ``wrap="none"`` and ``<a:spAutoFit/>``,
    both of which switch the overflow check off — so a fixture that left them
    alone would prove nothing.
    """
    frame = shape.text_frame
    frame.text = text
    frame.word_wrap = wrap
    frame.auto_size = None
    for para in frame.paragraphs:
        for run in para.runs:
            run.font.size = Pt(pt)
    return shape


def _translucent(shape, alpha_pct=30):
    """Fill the shape at ``alpha_pct`` opacity, as the flowchart emitter does."""
    colour = shape.fill._xPr.find(".//" + qn("a:srgbClr"))
    etree.SubElement(colour, qn("a:alpha")).set("val", str(alpha_pct * 1000))
    return shape


def _colour_text(shape, rgb):
    for para in shape.text_frame.paragraphs:
        for run in para.runs:
            run.font.color.rgb = rgb
    return shape


def _lint(prs, tmp_path: Path, name: str = "case.pptx") -> list[Finding]:
    path = tmp_path / name
    prs.save(str(path))
    findings, _ = lint_pptx(path)
    return findings


def _checks(findings) -> list[str]:
    return [f.check for f in findings]


def _only(findings, check: str) -> Finding:
    matches = [f for f in findings if f.check == check]
    assert len(matches) == 1, f"expected one {check}, got {_checks(findings)}"
    return matches[0]


# ---------------------------------------------------------------------------
# Geometry
# ---------------------------------------------------------------------------


def test_a_shape_past_the_canvas_edge_is_critical(tmp_path):
    """Coordinates past the edge are written, not clamped: the shape is gone."""
    prs, slide = _deck()
    _card(slide, 12.0, 1.0, 3.0, 1.0, "Off the edge")
    findings = _lint(prs, tmp_path)
    off = _only(findings, "shape_off_canvas")
    assert off.severity == "critical"
    assert off.slide == 1 and off.shape == 1
    assert "13.33x7.50in canvas" in off.detail


def test_a_shape_inside_the_canvas_is_not_reported(tmp_path):
    prs, slide = _deck()
    _card(slide, 0.5, 1.2, 3.0, 1.0, "Inside")
    _card(slide, 9.8, 6.0, 3.0, 1.4, "Also inside")
    assert _lint(prs, tmp_path) == []


def test_a_shape_fully_containing_another_is_layering_not_overlap(tmp_path):
    """A card behind its own text block is the design, not a collision."""
    prs, slide = _deck()
    _card(slide, 1.0, 1.0, 5.0, 3.0)
    _card(slide, 1.4, 1.4, 4.2, 2.2, "Body copy on its own card", fill=RGBColor(0xFF, 0xFF, 0xFF))
    assert "shape_overlap" not in _checks(_lint(prs, tmp_path))


def test_two_stacked_shapes_of_the_same_size_are_an_overlap(tmp_path):
    prs, slide = _deck()
    _card(slide, 1.0, 1.0, 4.0, 2.0)
    _card(slide, 1.6, 1.6, 4.0, 2.0)
    overlap = _only(_lint(prs, tmp_path), "shape_overlap")
    assert overlap.severity == "major"
    assert "60%" in overlap.detail


def test_a_hairline_rule_crossing_a_card_is_not_an_overlap(tmp_path):
    """Rules, dividers and connectors are drawn under what they join."""
    prs, slide = _deck()
    _card(slide, 1.0, 1.0, 4.0, 2.0)
    _card(slide, 0.6, 2.9, 12.0, 0.05)
    assert "shape_overlap" not in _checks(_lint(prs, tmp_path))


def test_an_empty_text_box_is_reported_when_it_is_big_in_both_directions(tmp_path):
    prs, slide = _deck()
    _textbox(slide, 1.0, 1.0, 4.0, 2.5)
    empty = _only(_lint(prs, tmp_path), "empty_text_box")
    assert empty.severity == "major"
    assert "10.0in²" in empty.detail


def test_a_thin_full_width_empty_box_is_a_rule_not_a_hole(tmp_path):
    prs, slide = _deck()
    _textbox(slide, 0.6, 5.0, 12.0, 0.2)
    assert _lint(prs, tmp_path) == []


# ---------------------------------------------------------------------------
# Text fit
# ---------------------------------------------------------------------------


def test_text_that_cannot_fit_its_box_is_reported(tmp_path):
    prs, slide = _deck()
    _card(slide, 1.0, 1.0, 3.0, 0.6, HAN_BODY)
    overflow = _only(_lint(prs, tmp_path), "text_overflow")
    # A filled shape clips its text, so a box this far over is lossy.
    assert overflow.severity == "critical"
    assert "0.60in box" in overflow.detail


def test_text_spilling_out_of_an_unfilled_box_is_ugly_rather_than_lossy(tmp_path):
    prs, slide = _deck()
    _textbox(slide, 1.0, 1.0, 3.0, 0.6, HAN_BODY)
    assert _only(_lint(prs, tmp_path), "text_overflow").severity == "major"


def test_the_overflow_estimate_is_cjk_aware(tmp_path):
    """Same box, same type size, same character count, twice the columns."""
    assert len(HAN_BODY) == len(LATIN_BODY)
    prs, slide = _deck()
    _card(slide, 1.0, 1.0, 3.0, 1.2, LATIN_BODY)
    _card(slide, 6.0, 1.0, 3.0, 1.2, HAN_BODY)
    overflow = _only(_lint(prs, tmp_path), "text_overflow")
    assert overflow.shape == 2


def test_estimate_text_block_measures_han_at_double_width():
    han_w, han_h = estimate_text_block(HAN_BODY, font_pt=18.0, width_in=3.0)
    latin_w, latin_h = estimate_text_block(LATIN_BODY, font_pt=18.0, width_in=3.0)
    assert han_h > latin_h
    # Unwrapped, the width is the whole measure: two glyphs against two
    # characters is exactly double.
    assert estimate_text_block("合同", font_pt=18.0, width_in=3.0, wrap=False)[0] == (
        2 * estimate_text_block("ab", font_pt=18.0, width_in=3.0, wrap=False)[0]
    )
    assert han_w <= 3.0 and latin_w <= 3.0  # both wrapped inside the box


def test_unwrapped_text_wider_than_its_box_is_reported(tmp_path):
    prs, slide = _deck()
    _textbox(slide, 1.0, 1.0, 1.5, 1.2,
            "An unwrapped line far wider than the box holding it", wrap=False)
    assert "unwrapped" in _only(_lint(prs, tmp_path), "text_overflow").detail


# ---------------------------------------------------------------------------
# Typography
# ---------------------------------------------------------------------------


def test_a_run_under_the_type_floor_is_reported(tmp_path):
    prs, slide = _deck()
    _card(slide, 1.0, 1.0, 4.0, 1.5, "Body copy nobody can read", pt=10.0)
    small = _only(_lint(prs, tmp_path), "type_too_small")
    assert small.severity == "major"
    assert f"floor {BODY_FONT_FLOOR_PT:.0f}pt" in small.detail


def test_a_short_box_gets_the_note_floor_not_the_body_floor(tmp_path):
    """A 0.4in box is a footer, a chip or a diagram node, not body copy."""
    prs, slide = _deck()
    _card(slide, 1.0, 1.0, 4.0, 0.4, "source note", pt=10.0)
    assert "type_too_small" not in _checks(_lint(prs, tmp_path, "note.pptx"))

    prs, slide = _deck()
    _card(slide, 1.0, 1.0, 4.0, 0.4, "source note", pt=8.0)
    below = _only(_lint(prs, tmp_path, "micro.pptx"), "type_too_small")
    assert f"floor {NOTE_FONT_FLOOR_PT:.0f}pt" in below.detail


# ---------------------------------------------------------------------------
# Contrast
# ---------------------------------------------------------------------------


def test_contrast_ratio_spans_one_to_twenty_one():
    white, black = (255, 255, 255), (0, 0, 0)
    assert contrast_ratio(white, black) == 21.0
    assert contrast_ratio(black, white) == 21.0
    assert contrast_ratio(white, white) == 1.0


def test_the_large_text_exemption_drops_the_contrast_floor():
    assert contrast_floor(18.0, False) == 3.0
    assert contrast_floor(14.0, True) == 3.0
    assert contrast_floor(14.0, False) == 4.5
    assert contrast_floor(17.5, False) == 4.5


def test_low_contrast_text_on_its_own_fill_is_reported(tmp_path):
    prs, slide = _deck()
    _colour_text(_card(slide, 1.0, 1.0, 4.0, 1.5, "Navy on navy", pt=14.0, fill=NAVY), NEAR_NAVY)
    low = _only(_lint(prs, tmp_path), "low_contrast")
    assert low.severity == "major"
    assert "WCAG needs 4.5:1" in low.detail


def test_a_translucent_fill_yields_no_contrast_finding(tmp_path):
    """Its effective colour is whatever is behind it, so the ratio is fiction.

    The flowchart emitter fills nodes with the accent at 70% transparency; read
    nominally that measured 1.9:1 against the navy label and manufactured a
    finding for every node on the slide.
    """
    prs, slide = _deck()
    node = _card(slide, 1.0, 1.0, 4.0, 1.5, "Navy on pale navy", pt=14.0, fill=NAVY)
    _colour_text(_translucent(node), NEAR_NAVY)
    assert "low_contrast" not in _checks(_lint(prs, tmp_path))


def test_an_inherited_text_colour_is_skipped_rather_than_guessed(tmp_path):
    """Resolving it needs the whole theme chain; guessing invents findings."""
    prs, slide = _deck()
    _card(slide, 1.0, 1.0, 4.0, 1.5, "Theme colour", pt=14.0, fill=NAVY)
    assert "low_contrast" not in _checks(_lint(prs, tmp_path))


# ---------------------------------------------------------------------------
# Renders
# ---------------------------------------------------------------------------


def _page(tmp_path: Path, name: str, ink_boxes, *, size=(1280, 720), bg=255) -> Path:
    path = tmp_path / name
    image = Image.new("L", size, bg)
    draw = ImageDraw.Draw(image)
    for box in ink_boxes:
        draw.rectangle(box, fill=30)
    image.save(str(path))
    return path


def test_page_ink_reads_the_background_off_the_border_ring(tmp_path):
    """One wrong corner inverts every number downstream.

    This _page's top-left corner sits inside a dark panel. Sampling corners
    would call the background dark and report the whole _page as ink; the modal
    value of the border ring is still white.
    """
    metrics = page_ink(_page(tmp_path, "corner.png", [(0, 0, 240, 120)]))
    assert metrics.coverage < 0.1
    assert metrics.left == 0.0 and metrics.top == 0.0
    assert metrics.right < 0.25 and metrics.bottom < 0.25


def test_page_ink_reports_coverage_and_the_dead_band(tmp_path):
    thin = page_ink(_page(tmp_path, "thin.png", [(100, 40, 400, 200)]))
    assert thin.coverage < 0.2
    assert thin.dead_band > 0.6  # everything below the block is empty

    full = page_ink(_page(tmp_path, "full.png", [(40, 30, 1240, 690)]))
    assert full.coverage > 0.8
    assert full.dead_band < 0.1
    assert full.ink > thin.ink


def test_a_blank_page_reports_no_ink_rather_than_dividing_by_zero(tmp_path):
    blank = page_ink(_page(tmp_path, "blank.png", []))
    assert blank.coverage == 0.0 and blank.ink == 0.0
    assert blank.dead_rows == 1.0 and blank.dead_band == 1.0


def test_lint_renders_measures_but_grades_nothing(tmp_path):
    """Occupancy is a proxy for effort that padding satisfies.

    Tested over 98 slides of real work: the underfill check ranked our
    emptiest slide above the reference set's densest statutory slide, because
    a pale card is ink, and every page the overfill check ever flagged was
    correct work. The numbers stay — 0.08 is how the degraded writer looks —
    and nothing is graded on them.
    """
    thin = _page(tmp_path, "thin.png", [(100, 40, 400, 200)])
    full = _page(tmp_path, "full.png", [(40, 30, 1240, 690)])
    findings, metrics = lint_renders([thin, full])
    assert findings == []
    assert RENDER_CHECKS == frozenset()
    assert metrics["slides_measured"] == 2
    assert metrics["coverage_min"] < metrics["coverage_mean"]


def test_a_render_that_cannot_be_read_is_skipped_not_raised(tmp_path):
    broken = tmp_path / "broken.png"
    broken.write_bytes(b"not a png")
    good = _page(tmp_path, "full.png", [(40, 30, 1240, 690)])
    findings, metrics = lint_renders([broken, good])
    assert metrics["slides_measured"] == 1
    assert findings == []


# ---------------------------------------------------------------------------
# DOCX
# ---------------------------------------------------------------------------


def test_lint_docx_catches_a_first_line_indent_inside_a_table_cell(tmp_path):
    """A cell is one short field; the CJK body indent breaks a title in two."""
    doc = Document()
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "《民用航空法》第八十八条"
    table.cell(0, 1).text = "60 日"
    table.cell(0, 0).paragraphs[0].paragraph_format.first_line_indent = DocxPt(24)
    path = tmp_path / "indent.docx"
    doc.save(str(path))
    findings, metrics = lint_docx(path)
    indent = _only(findings, "cell_first_line_indent")
    assert indent.where == "table 1"
    assert indent.severity == "major"
    assert metrics["tables"] == 1


def test_lint_docx_reads_a_character_based_indent_too(tmp_path):
    """The CJK writer sets w:firstLineChars, which never reaches python-docx."""
    doc = Document()
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "《民用航空法》第八十八条"
    table.cell(0, 1).text = "60 日"
    pPr = table.cell(0, 0).paragraphs[0]._p.get_or_add_pPr()
    etree.SubElement(pPr, w_qn("w:ind")).set(w_qn("w:firstLineChars"), "200")
    path = tmp_path / "chars.docx"
    doc.save(str(path))
    assert _only(lint_docx(path)[0], "cell_first_line_indent")


def test_a_clean_memo_yields_nothing(tmp_path):
    doc = Document()
    doc.add_heading("一、结论", level=1)
    doc.add_paragraph("承租人应当在交付前完成登记。")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "事项"
    table.cell(0, 1).text = "时限"
    path = tmp_path / "clean.docx"
    doc.save(str(path))
    assert lint_docx(path)[0] == []


# ---------------------------------------------------------------------------
# Report
# ---------------------------------------------------------------------------


def _report() -> QualityReport:
    findings = [
        Finding("text_overflow", "major", f"slide {i}", f"detail {i}", "Shorten it.", slide=i)
        for i in range(1, 6)
    ]
    findings.append(Finding("shape_off_canvas", "critical", "slide 2", "gone", "Move it.", slide=2))
    return QualityReport(findings=findings, metrics={"slides": 6})


def test_to_payload_groups_repeated_checks_into_one_row():
    """Twenty overflowing boxes say the type scale is wrong, not 'slide 7 again'."""
    payload = _report().to_payload()
    rows = {row["check"]: row for row in payload["findings"]}
    assert rows["text_overflow"]["count"] == 5
    assert len(rows["text_overflow"]["where"]) == 3  # examples, not the list
    assert payload["counts"] == {"critical": 1, "major": 5}
    # Critical first, because that is the order a repair pass should work in.
    assert payload["findings"][0]["check"] == "shape_off_canvas"


def test_to_payload_says_blocked_when_anything_is_critical():
    assert _report().to_payload()["conformance"] == "blocked"
    assert _report().ok is False
    warnings = QualityReport(findings=[Finding("type_too_small", "major", "slide 1", "d", "f")])
    assert warnings.to_payload()["conformance"] == "warnings"
    assert warnings.ok is True
    assert QualityReport().to_payload()["conformance"] == "pass"
    # The wording is load-bearing. A pass reported as "quality: clean" is what
    # let an empty deck be handed over: these checks are a conformance floor
    # and cannot say whether the document reads well.
    assert QualityReport().summary().startswith("conformance: pass")
    assert "quality" not in QualityReport().summary()
    assert _report().summary() == "conformance: 1 critical, 5 major"


def test_to_payload_respects_its_character_cap():
    """It rides along with an image-bearing tool result, which skips the budget."""
    full = _report().to_payload()
    capped = _report().to_payload(cap=300)
    assert len(repr(capped)) <= 300
    assert capped["truncated"] is True
    assert len(capped["findings"]) < len(full["findings"])
    # The verdict and the counts survive the cap; only examples are dropped.
    assert capped["conformance"] == "blocked"
    assert capped["counts"] == full["counts"]


def test_lint_artifact_reports_an_unreadable_file_instead_of_raising(tmp_path):
    """A linter that can fail a document write would be worse than no linter."""
    broken = tmp_path / "broken.pptx"
    broken.write_bytes(b"not a deck")
    report = lint_artifact(broken)
    assert report.findings == []
    assert "file_check_error" in report.metrics


def test_lint_artifact_joins_the_file_and_its_renders(tmp_path):
    prs, slide = _deck()
    _card(slide, 12.0, 1.0, 3.0, 1.0, "Off the edge")
    path = tmp_path / "joined.pptx"
    prs.save(str(path))
    render = _page(tmp_path, "slide-1.png", [(100, 40, 400, 200)])
    report = lint_artifact(path, renders=[render])
    # The file contributes findings; the render contributes only numbers.
    assert set(_checks(report.findings)) == {"shape_off_canvas"}
    assert report.metrics["slides"] == 1 and report.metrics["slides_measured"] == 1
    assert "coverage_mean" in report.metrics
    assert report.ok is False
