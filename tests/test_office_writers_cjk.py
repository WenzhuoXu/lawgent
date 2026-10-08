"""CJK + safety contracts for the pdf writer and run-aware pptx edits."""

from __future__ import annotations

import zipfile

from legal_helper.documents.office import edit_pptx_text, extract_pptx_text
from legal_helper.documents.writers import write_pdf_document
from legal_helper.documents.writers.pdf import _detect_lang


# ---------------------------------------------------------------------------
# PDF writer
# ---------------------------------------------------------------------------


def test_pdf_lang_detection():
    assert _detect_lang("法律意见书") == "zh-CN"
    assert _detect_lang("Lease Review", None, "plain body") == "en"
    assert _detect_lang("Review", None, "含中文正文的备忘录") == "zh-CN"


def test_pdf_escapes_title_and_renders_cjk(tmp_path):
    out = tmp_path / "memo.pdf"
    write_pdf_document(
        output_path=out,
        title="法律意见书 <草稿> & Review",
        body_markdown="## 一、结论\n\n拟议交易符合《公司法》相关规定。",
        subtitle="仅供讨论 <internal>",
        disclaimer="非正式法律意见",
    )
    assert out.is_file()

    from pdfminer.high_level import extract_text

    text = extract_text(str(out))
    # Escaped markup renders literally instead of being swallowed as HTML.
    assert "<草稿>" in text and "& Review" in text
    assert "<internal>" in text
    # CJK glyphs actually embedded (SC-capable font stack), not tofu.
    assert "法律意见书" in text
    assert "拟议交易符合" in text


# ---------------------------------------------------------------------------
# edit_pptx_text: run-aware replacement with per-find counts
# ---------------------------------------------------------------------------


def _deck_with_split_runs(path):
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(2))
    para = box.text_frame.paragraphs[0]
    r1 = para.add_run()
    r1.text = "Hello "
    r1.font.bold = True
    r2 = para.add_run()
    r2.text = "World"
    para2 = box.text_frame.add_paragraph()
    r3 = para2.add_run()
    r3.text = "World peace"
    prs.save(str(path))
    return path


def test_edit_pptx_text_run_aware_with_counts(tmp_path):
    src = _deck_with_split_runs(tmp_path / "src.pptx")
    out = tmp_path / "edited.pptx"
    result = edit_pptx_text(
        src,
        out,
        [
            {"slide": 1, "find": "Hello World", "replace": "Goodbye Moon"},  # spans two runs
            {"slide": 1, "find": "peace", "replace": "harmony"},  # single run
            {"slide": 1, "find": "absent", "replace": "x"},  # miss -> count 0
        ],
    )
    assert result["replacements"] == {
        "slide 1: Hello World": 1,
        "slide 1: peace": 1,
        "slide 1: absent": 0,
    }
    text = extract_pptx_text(out)
    assert "Goodbye Moon" in text
    assert "World harmony" in text
    assert "Hello" not in text


def test_edit_pptx_text_reopens_with_python_pptx(tmp_path):
    from pptx import Presentation

    src = _deck_with_split_runs(tmp_path / "src.pptx")
    out = tmp_path / "edited.pptx"
    edit_pptx_text(src, out, [{"slide": 1, "find": "World", "replace": "全球"}])
    prs = Presentation(str(out))
    texts = [
        shape.text
        for slide in prs.slides
        for shape in slide.shapes
        if getattr(shape, "has_text_frame", False)
    ]
    joined = "\n".join(texts)
    assert "全球" in joined
    assert "World" not in joined
