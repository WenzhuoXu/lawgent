"""reshape_xlsx / copy_xlsx_sheet on real-world workbooks (review item D-5).

- ``set_styles`` rebuilt ``Font`` from ``base.color.rgb``; for theme, indexed
  or auto colours — openpyxl's default font among them — ``Font()`` raised.
  19 of 45 recorded calls failed this way.
- Row/column inserts moved values but not merges, validations, conditional
  formats or formula references, and reported nothing.
- ``copy_xlsx_sheet`` dropped validations and conditional formats, contrary
  to its own docstring.
"""

from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Font, PatternFill
from openpyxl.styles.colors import Color
from openpyxl.worksheet.datavalidation import DataValidation

from legal_helper.tools import copy_xlsx_sheet, reshape_xlsx


def _workbook(path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "A类"
    for row in range(1, 11):
        ws.cell(row, 1, f"item{row}")
        ws.cell(row, 2, row * 10)
    ws["B11"] = "=SUM(B1:B10)"
    ws["C5"] = "=B5*2"
    ws["D1"] = "=$B$7"
    ws["A1"].font = Font(name="Calibri", color=Color(theme=1))  # theme colour
    ws.merge_cells("A12:B13")
    ws.merge_cells("E3:E8")
    dv = DataValidation(type="list", formula1='"Y,N"')
    dv.add("C2:C9")
    ws.add_data_validation(dv)
    ws.conditional_formatting.add(
        "B2:B10", CellIsRule(operator="greaterThan", formula=["50"], fill=PatternFill("solid", start_color="FFFF0000"))
    )
    ws.conditional_formatting.add(
        "A2:A10", FormulaRule(formula=["$B2>50"], fill=PatternFill("solid", start_color="FF00FF00"))
    )
    ws.row_dimensions[8].height = 30
    summary = wb.create_sheet("Summary")
    summary["A1"] = "='A类'!B11"
    summary["A2"] = "='A类'!B5"
    wb.save(path)
    return path


def _reshape(src: Path, operations: list[dict], name: str) -> dict:
    raw = reshape_xlsx.call({"source_path": str(src), "filename": name, "operations": operations})
    assert not raw.startswith("ERROR"), raw
    return json.loads(raw)


def test_set_styles_on_theme_and_default_fonts(tmp_path):
    src = _workbook(tmp_path / "src.xlsx")
    result = _reshape(
        src,
        [{"op": "set_styles", "sheet": "A类", "range": "A1:D3", "font_bold": True, "font_color": "1F4E79"}],
        "styled.xlsx",
    )
    ws = load_workbook(result["path"])["A类"]
    assert ws["A1"].font.bold and ws["A1"].font.name == "Calibri"  # copy kept the rest
    assert ws["C3"].font.bold  # openpyxl's default font, never styled before
    assert ws["A1"].font.color.rgb.endswith("1F4E79")


def test_insert_rows_carries_merges_validations_formats_and_formulas(tmp_path):
    src = _workbook(tmp_path / "src.xlsx")
    result = _reshape(src, [{"op": "insert_rows", "sheet": "A类", "at": 5, "count": 2}], "ins.xlsx")
    effects = result["applied"][0]["side_effects"]
    assert effects["merged_ranges_shifted"] == 2
    assert effects["data_validations_shifted"] == 1
    assert effects["conditional_formats_shifted"] == 2
    assert effects["formulas_rewritten"] >= 5
    assert effects["row_heights_moved"] == 1
    wb = load_workbook(result["path"])
    ws = wb["A类"]
    assert ws["B13"].value == "=SUM(B1:B12)"  # the range grew around the insert
    assert ws["C7"].value == "=B7*2"  # the moved formula follows its row
    assert ws["D1"].value == "=$B$9"  # absolute references move too
    assert wb["Summary"]["A1"].value == "='A类'!B13"  # and on other sheets
    assert sorted(str(m) for m in ws.merged_cells.ranges) == ["A14:B15", "E3:E10"]
    assert [str(dv.sqref) for dv in ws.data_validations.dataValidation] == ["C2:C11"]
    cf = {str(c.sqref): [rule.formula for rule in c.rules] for c in ws.conditional_formatting}
    assert cf == {"B2:B12": [["50"]], "A2:A12": [["$B2>50"]]}
    assert ws.row_dimensions[10].height == 30


def test_delete_rows_reports_references_into_deleted_cells(tmp_path):
    src = _workbook(tmp_path / "src.xlsx")
    result = _reshape(src, [{"op": "delete_rows", "sheet": "A类", "at": 5, "count": 3}], "del.xlsx")
    effects = result["applied"][0]["side_effects"]
    assert "Summary!A2" in effects["formulas_now_ref_error"]  # ='A类'!B5 pointed into the cut
    # The recalculation that follows every edit reports the broken cells too.
    assert any(err["cell"] == "A2" and err["error"] == "#REF!" for err in result["formula_errors"])
    wb = load_workbook(result["path"])
    assert wb["Summary"]["A2"].value == "=#REF!"
    assert wb["A类"]["B8"].value == "=SUM(B1:B7)"
    assert wb["A类"]["D1"].value == "=#REF!"  # $B$7 was deleted


def test_unknown_op_lists_the_valid_ops(tmp_path):
    src = _workbook(tmp_path / "src.xlsx")
    raw = reshape_xlsx.call({"source_path": str(src), "filename": "bad.xlsx", "operations": [{"op": "frobnicate", "sheet": "A类"}]})
    assert raw.startswith("ERROR") and "frobnicate" in raw
    assert "insert_rows" in raw and "set_styles" in raw


def test_copy_row_translates_formulas(tmp_path):
    src = _workbook(tmp_path / "src.xlsx")
    result = _reshape(src, [{"op": "copy_row", "sheet": "A类", "src_row": 5, "dest_row": 9}], "copy.xlsx")
    assert load_workbook(result["path"])["A类"]["C9"].value == "=B9*2"
    assert result["applied"][0]["formulas_translated"] == 1


def test_copy_sheet_keeps_validations_and_conditional_formats(tmp_path):
    src = _workbook(tmp_path / "src.xlsx")
    raw = copy_xlsx_sheet.call({"source_path": str(src), "filename": "copy.xlsx", "sheet": "A类", "new_name": "B类"})
    payload = json.loads(raw)
    assert payload["copied"]["data_validations"] == 1
    assert payload["copied"]["conditional_formats"] == 2
    ws = load_workbook(payload["path"])["B类"]
    assert [str(dv.sqref) for dv in ws.data_validations.dataValidation] == ["C2:C9"]
    assert len(list(ws.conditional_formatting)) == 2


def test_names_tables_print_area_and_columns_move_too(tmp_path):
    from openpyxl.workbook.defined_name import DefinedName
    from openpyxl.worksheet.table import Table

    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["Name", "Amount", "Tax"])
    for i in range(2, 8):
        ws.append([f"n{i}", i, f"=B{i}*0.1"])
    ws["E1"] = "=SUM(Amounts)"
    ws["E2"] = "=SUM(C:C)"
    wb.defined_names["Amounts"] = DefinedName("Amounts", attr_text="Data!$B$2:$B$7")
    ws.add_table(Table(displayName="T1", ref="A1:C7"))
    ws.print_area = "A1:C7"
    ws.merge_cells("A9:C9")
    src = tmp_path / "names.xlsx"
    wb.save(src)
    rows = _reshape(src, [{"op": "insert_rows", "sheet": "Data", "at": 4, "count": 2}], "rows.xlsx")
    assert rows["applied"][0]["side_effects"]["defined_names_rewritten"] == 1
    out = load_workbook(rows["path"])
    assert out.defined_names["Amounts"].attr_text == "Data!$B$2:$B$9"
    table = out["Data"].tables["T1"]
    assert table.ref == "A1:C9" and table.autoFilter.ref == "A1:C9"
    assert out["Data"].print_area == "'Data'!$A$1:$C$9"
    cols = _reshape(src, [{"op": "insert_cols", "sheet": "Data", "at": 2, "count": 1}], "cols.xlsx")
    ws2 = load_workbook(cols["path"])["Data"]
    assert ws2["D3"].value == "=C3*0.1"  # the Tax formula moved right and follows Amount
    assert ws2["F2"].value == "=SUM(D:D)"  # whole-column references move too
    assert [str(m) for m in ws2.merged_cells.ranges] == ["A9:D9"]
