"""Round-trip tests for native Word tracked-changes redlining (w:ins/w:del)."""

from __future__ import annotations

import zipfile
from xml.etree import ElementTree as ET

import pytest
from docx import Document

from legal_helper.documents.redline import (
    extract_docx_revisions,
    preview_docx_revisions,
    redline_docx_document,
)


_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _make_source(path):
    doc = Document()
    p = doc.add_paragraph()
    p.add_run("The lessee shall pay ")
    amount = p.add_run("USD 1,000,000")
    amount.bold = True
    p.add_run(" within thirty (30) days.")
    doc.add_paragraph("Governing law: the laws of England.")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Deposit"
    table.rows[0].cells[1].text = "USD 50,000"
    doc.save(str(path))
    return path


@pytest.fixture()
def redlined(tmp_path):
    src = _make_source(tmp_path / "src.docx")
    out = tmp_path / "redlined.docx"
    result = redline_docx_document(
        src,
        out,
        [
            {"op": "replace", "find": "USD 1,000,000", "replace": "USD 850,000",
             "comment": "Reduced per term sheet."},
            {"op": "delete", "find": "thirty (30) "},
            {"op": "insert_after", "find": "days.",
             "text": " Payment shall be in immediately available funds."},
            {"op": "replace", "find": "USD 50,000", "replace": "USD 25,000"},
            {"op": "comment", "find": "Governing law", "text": "Confirm dispute-resolution seat."},
            {"op": "replace", "find": "not present anywhere", "replace": "x"},
        ],
        timestamp="2026-07-21T09:00:00Z",
    )
    return src, out, result


def test_redline_reports_counts_and_misses(redlined):
    _, out, result = redlined
    assert out.is_file()
    counts = {item["find"]: item["count"] for item in result["applied"]}
    assert counts["USD 1,000,000"] == 1
    assert counts["USD 50,000"] == 1  # table-cell paragraphs are covered
    assert counts["not present anywhere"] == 0
    # A miss is reported with a reason and a hint, never silently dropped.
    assert [m["find"] for m in result["not_found"]] == ["not present anywhere"]
    assert result["not_found"][0]["reason"] == "not_found"
    assert result["not_found"][0]["hint"]
    assert result["revision_count"] == 6  # 2 + 1 + 1 + 2
    assert result["comment_count"] == 2


def test_redline_emits_real_ins_del_marks(redlined):
    _, out, _ = redlined
    with zipfile.ZipFile(out) as zf:
        xml = zf.read("word/document.xml").decode("utf-8")
    root = ET.fromstring(xml)
    ins = list(root.iter(f"{_W}ins"))
    dels = list(root.iter(f"{_W}del"))
    assert len(ins) == 3 and len(dels) == 3
    ids = [el.get(f"{_W}id") for el in ins + dels]
    assert len(ids) == len(set(ids)), "revision ids must be unique"
    for el in ins + dels:
        assert el.get(f"{_W}author") == "Legal AI Helper"
        assert el.get(f"{_W}date") == "2026-07-21T09:00:00Z"
    # Deleted text must use w:delText, never w:t (Word rejects the latter).
    for del_el in dels:
        assert list(del_el.iter(f"{_W}delText"))
        assert not list(del_el.iter(f"{_W}t"))


def test_redline_round_trips_accept_and_reject(redlined):
    src, out, _ = redlined
    accepted = preview_docx_revisions(out, "accept")
    rejected = preview_docx_revisions(out, "reject")
    assert (
        "The lessee shall pay USD 850,000 within days."
        " Payment shall be in immediately available funds." in accepted
    )
    assert "USD 25,000" in accepted
    # Rejecting every mark restores the original document text exactly.
    original = Document(str(src))
    original_text = "\n".join(p.text for p in original.paragraphs)
    for line in original_text.splitlines():
        assert line in rejected
    assert "USD 850,000" not in rejected
    assert "USD 1,000,000" in rejected


def test_extract_docx_revisions_lists_marks(redlined):
    _, out, _ = redlined
    revisions = extract_docx_revisions(out)
    by_type = {}
    for rev in revisions:
        by_type.setdefault(rev["type"], []).append(rev["text"])
    assert "USD 1,000,000" in by_type["deletion"]
    assert "USD 850,000" in by_type["insertion"]
    assert " Payment shall be in immediately available funds." in by_type["insertion"]
    assert all(rev["date"].endswith("Z") for rev in revisions)


def test_redline_preserves_run_formatting(redlined):
    _, out, _ = redlined
    with zipfile.ZipFile(out) as zf:
        root = ET.fromstring(zf.read("word/document.xml"))
    # The replacement for the bold "USD 1,000,000" run copies its rPr.
    for ins in root.iter(f"{_W}ins"):
        text = "".join(t.text or "" for t in ins.iter(f"{_W}t"))
        if text == "USD 850,000":
            run = ins.find(f"{_W}r")
            rpr = run.find(f"{_W}rPr")
            assert rpr is not None and rpr.find(f"{_W}b") is not None
            break
    else:
        pytest.fail("inserted replacement run not found")


def test_redline_comments_round_trip_via_python_docx(redlined):
    _, out, _ = redlined
    doc = Document(str(out))
    comments = {c.text for c in doc.comments}
    assert "Reduced per term sheet." in comments
    assert "Confirm dispute-resolution seat." in comments
    for c in doc.comments:
        assert c.author == "Legal AI Helper"


def test_redline_output_reopens_cleanly(redlined):
    _, out, _ = redlined
    # python-docx must parse the revised body without error and still see
    # the untouched text as regular runs.
    doc = Document(str(out))
    texts = [p.text for p in doc.paragraphs]
    assert any("The lessee shall pay" in t for t in texts)


def test_redline_cross_run_match(tmp_path):
    src = tmp_path / "src.docx"
    doc = Document()
    p = doc.add_paragraph()
    p.add_run("Liquidated ")
    p.add_run("damages ")
    p.add_run("apply here.")
    doc.save(str(src))

    out = tmp_path / "out.docx"
    result = redline_docx_document(
        src, out,
        [{"op": "replace", "find": "Liquidated damages apply",
          "replace": "违约金条款适用"}],
    )
    assert result["applied"][0]["count"] == 1
    assert "违约金条款适用 here." in preview_docx_revisions(out, "accept")
    assert "Liquidated damages apply here." in preview_docx_revisions(out, "reject")


def test_redline_multiple_occurrences_all_marked(tmp_path):
    src = tmp_path / "src.docx"
    doc = Document()
    doc.add_paragraph("Party A pays Party A's costs; Party A indemnifies.")
    doc.save(str(src))

    out = tmp_path / "out.docx"
    result = redline_docx_document(
        src, out, [{"op": "replace", "find": "Party A", "replace": "Party B"}]
    )
    assert result["applied"][0]["count"] == 3
    accepted = preview_docx_revisions(out, "accept")
    assert accepted.count("Party B") == 3
    assert "Party A" not in accepted


def test_redline_rejects_empty_find(tmp_path):
    src = _make_source(tmp_path / "src.docx")
    with pytest.raises(ValueError):
        redline_docx_document(src, tmp_path / "out.docx", [{"op": "replace", "find": ""}])


# ---------------------------------------------------------------------------
# Review item D-3 / proposal 5: the engine as a tool, and what it got wrong
# ---------------------------------------------------------------------------

import json  # noqa: E402
from pathlib import Path  # noqa: E402

from docx_fixtures import PRC_CONTRACT_NUMBERING, build_docx, document_xml, ins as ins_xml, p, part_xml, r  # noqa: E402


def _tab_source(tmp_path):
    body = p("Term:", runs='<w:r><w:tab/><w:t>three years</w:t></w:r>')
    return build_docx(tmp_path / "tab.docx", body)


def test_tabs_survive_a_run_split(tmp_path):
    """Reproduced: "Term:\\tthree years" redlined to five accepted as
    "Term:five years" — the split collapsed the run to one w:t."""
    out = tmp_path / "out.docx"
    redline_docx_document(_tab_source(tmp_path), out, [{"find": "three", "replace": "five"}])
    assert preview_docx_revisions(out, "accept") == "Term:\tfive years"
    assert preview_docx_revisions(out, "reject") == "Term:\tthree years"


def _contract(tmp_path):
    body = (
        p("总则", num=(1, 0))
        + p("本合同双方为甲方与乙方。")
        + p("租期", num=(1, 0))
        + p("租期为三年，自交付之日起算。")
        + p("租金", num=(1, 0))
        + p("", runs=r("每月租金人民币壹万元") + '<w:r><w:footnoteReference w:id="1"/></w:r>' + r("，按月支付。"))
        + p("Pending ", runs=ins_xml("counterparty text"))
    )
    return build_docx(
        tmp_path / "contract.docx",
        body,
        numbering=PRC_CONTRACT_NUMBERING,
        footnotes='<w:footnote w:id="1"><w:p><w:r><w:t>不含税</w:t></w:r></w:p></w:footnote>',
    )


def test_round_trip_contract_over_every_op(tmp_path):
    """Accept-all equals the intended revision; reject-all equals the original."""
    src = _contract(tmp_path)
    out = tmp_path / "out.docx"
    result = redline_docx_document(
        src,
        out,
        [
            {"op": "replace", "clause": "第二条", "find": "三年", "replace": "五年"},
            {"op": "delete", "paragraph": 1},
            {"op": "insert_after", "clause": "第三条", "text": "租金应于每月五日前支付。"},
            {"op": "insert_before", "find": "每月租金", "text": "（含物业费）"},
            {"op": "comment", "clause": "第一条", "text": "Check the parties."},
        ],
        timestamp="2026-10-07T00:00:00Z",
    )
    assert not result["not_found"], result["not_found"]
    # The source carries a counterparty's pending insertion, so each view is
    # compared with the same view of the source.
    assert preview_docx_revisions(out, "reject") == preview_docx_revisions(src, "reject")
    original = preview_docx_revisions(src, "accept")
    # 第三条 runs to the end of the document, so the new paragraph lands last.
    expected = (
        original.replace("本合同双方为甲方与乙方。\n", "")
        .replace("租期为三年", "租期为五年")
        .replace("每月租金", "（含物业费）每月租金")
        + "\n租金应于每月五日前支付。"
    )
    assert preview_docx_revisions(out, "accept") == expected


def test_round_trip_holds_in_libreoffice_too(tmp_path):
    """The same contract checked by LibreOffice's own accept-all / reject-all,
    not by this module's preview."""
    from legal_helper.documents.compare import accept_all, reject_all, uno_python
    from legal_helper.documents.soffice import soffice_binary

    if not soffice_binary() or not uno_python():
        pytest.skip("LibreOffice with a uno-capable Python is not available")
    src = _contract(tmp_path)
    out = tmp_path / "out.docx"
    redline_docx_document(
        src,
        out,
        [
            {"op": "replace", "clause": "第二条", "find": "三年", "replace": "五年"},
            {"op": "delete", "paragraph": 1},
            {"find": "Term", "replace": "x"},  # not found: must not disturb anything
        ],
    )

    def lines(path, mode="accept"):
        return [line for line in preview_docx_revisions(path, mode).splitlines() if line.strip()]

    accepted = lines(accept_all(out, tmp_path / "accepted.docx"))
    rejected = lines(reject_all(out, tmp_path / "rejected.docx"))
    assert rejected == lines(src, "reject")
    assert "Pending counterparty text" in accepted
    assert "租期为五年，自交付之日起算。" in accepted
    assert "本合同双方为甲方与乙方。" not in accepted


def test_footnote_reference_is_kept_out_of_a_deletion(tmp_path):
    out = tmp_path / "out.docx"
    redline_docx_document(_contract(tmp_path), out, [{"find": "人民币壹万元，按月", "replace": "人民币贰万元，按季"}])
    xml = document_xml(out)
    deleted = xml[xml.index("<w:del "):]
    deleted = deleted[: deleted.index("</w:del>")]
    assert "footnoteReference" not in deleted
    assert "footnoteReference" in xml


def test_finds_that_open_with_a_numbering_label(tmp_path):
    """3 of 5 misses on the review fixture were strings copied from
    read_document output — numbering labels that are not in the XML. Such a
    find anchors on the one paragraph carrying the label; when it cannot, the
    miss says why and how to anchor instead."""
    out = tmp_path / "out.docx"
    result = redline_docx_document(
        _contract(tmp_path),
        out,
        [
            {"find": "第2条 租期", "replace": "第2条 期限"},
            {"find": "第一条 不在文中", "replace": "x"},
            {"clause": "第九条", "find": "x", "replace": "y"},
        ],
    )
    first = result["applied"][0]
    assert first["count"] == 1 and first["anchored_after_label"] == "第二条"
    accepted = preview_docx_revisions(out, "accept").splitlines()
    assert "期限" in accepted and "租期" not in accepted  # the label itself stays Word's
    missing, unknown_clause = result["not_found"]
    assert missing["reason"] == "label_is_auto_numbering"
    assert "clause='第一条'" in missing["hint"]
    assert unknown_clause["reason"] == "clause_not_found" and "第一条 总则" in unknown_clause["hint"]


def test_existing_revisions_are_never_rematched(tmp_path):
    result = redline_docx_document(
        _contract(tmp_path), tmp_path / "out.docx", [{"find": "counterparty text", "replace": "ours"}]
    )
    assert result["not_found"][0]["reason"] == "inside_tracked_change"


def test_track_changes_is_switched_on(tmp_path):
    out = tmp_path / "out.docx"
    result = redline_docx_document(_tab_source(tmp_path), out, [{"find": "three", "replace": "five"}])
    assert result["track_changes_turned_on"] is True
    settings = part_xml(out, "word/settings.xml")
    assert "<w:trackRevisions/>" in settings
    # Schema order: trackRevisions precedes defaultTabStop.
    if "defaultTabStop" in settings:
        assert settings.index("trackRevisions") < settings.index("defaultTabStop")


def test_protected_documents_are_refused_but_tracked_changes_protection_is_not(tmp_path):
    from legal_helper.documents.redline import RedlineRefused

    body = p("Term: three years")
    locked = build_docx(
        tmp_path / "locked.docx", body,
        settings_extra='<w:documentProtection w:edit="readOnly" w:enforcement="1"/>',
    )
    with pytest.raises(RedlineRefused, match="readOnly"):
        redline_docx_document(locked, tmp_path / "x.docx", [{"find": "three", "replace": "five"}])
    tracked_only = build_docx(
        tmp_path / "tracked.docx", body,
        settings_extra='<w:documentProtection w:edit="trackedChanges" w:enforcement="1"/>',
    )
    result = redline_docx_document(tracked_only, tmp_path / "y.docx", [{"find": "three", "replace": "five"}])
    assert result["revision_count"] == 2


def test_per_category_authors(tmp_path):
    out = tmp_path / "out.docx"
    redline_docx_document(
        _contract(tmp_path),
        out,
        [
            {"find": "三年", "replace": "五年", "category": "substantive"},
            {"find": "，按月支付", "replace": "，按月付清", "category": "drafting"},
        ],
        author="Counsel",
        category_authors={"substantive": "Counsel (substantive)"},
    )
    authors = {rev["author"] for rev in extract_docx_revisions(out) if rev["author"] != "Counterparty"}
    assert authors == {"Counsel (substantive)", "Counsel (drafting)"}


def test_redline_docx_tool_end_to_end(tmp_path):
    from legal_helper.tools import redline_docx

    src = _contract(tmp_path)
    raw = redline_docx.call(
        {
            "source_path": str(src),
            "filename": "contract.docx",
            "changes": [
                {"op": "replace", "clause": "第2条", "find": "三年", "replace": "五年", "comment": "Per term sheet."},
                {"op": "replace", "find": "第一条 不在文中", "replace": "x"},
            ],
            "author": "律师甲",
        }
    )
    payload = json.loads(raw)
    assert Path(payload["output_path"]).is_file()
    assert Path(payload["output_path"]) != src
    assert payload["revision_count"] == 2 and payload["comment_count"] == 1
    assert payload["not_found"][0]["reason"] == "label_is_auto_numbering"
    assert "re-anchor" in payload["next_step"].lower()
    assert payload["integrity"]["status"] == "ok", payload["integrity"]
    from legal_helper.tools import read_document

    text = read_document.call({"path": payload["output_path"]})
    assert "租期为{--三年--}{++五年++}" in text


def test_redline_and_compare_are_on_every_authoring_surface_and_mutating():
    """D-3: no tool exposed the engine; the skill told the model to call a
    Python function it could not reach."""
    from legal_helper.tool_policy import is_mutating, tool_name_of
    from legal_helper.tools import skill_tools_for_task
    from legal_helper.workflow import _document_io_tools

    for surface in (_document_io_tools(), skill_tools_for_task("review-contract")):
        names = {tool_name_of(t) for t in surface}
        assert {"redline_docx", "compare_docx"} <= names
    assert is_mutating("redline_docx") and is_mutating("compare_docx")
    auditor = {tool_name_of(t) for t in skill_tools_for_task("cite-check")}
    assert "redline_docx" not in auditor and "compare_docx" not in auditor
    skill = (Path(__file__).resolve().parent.parent / "legal_helper/skills/docx-redline/SKILL.md").read_text(encoding="utf-8")
    assert "redline_docx(" in skill and "redline_docx_document" not in skill
    assert "layout=" not in skill


def test_markdown_emphasis_copied_from_read_document_still_anchors(tmp_path):
    body = p("", runs=r("Lessor", bold=True) + r(": Aurora Capital"))
    out = tmp_path / "out.docx"
    result = redline_docx_document(
        build_docx(tmp_path / "bold.docx", body), out, [{"find": "**Lessor**: Aurora", "replace": "Owner: Aurora"}]
    )
    assert result["applied"][0]["count"] == 1
    assert result["applied"][0]["matched_without_markdown"] == "Lessor: Aurora"
    assert preview_docx_revisions(out, "accept") == "Owner: Aurora Capital"
