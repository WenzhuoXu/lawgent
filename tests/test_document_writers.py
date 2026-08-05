"""Round-trip tests for the .docx and .pdf writers."""

from __future__ import annotations

from pathlib import Path

from legal_helper.documents.writers import write_docx_document, write_pdf_document


SAMPLE_BODY = """
This is the introductory paragraph.

## Risk Findings

- **GREEN**: Insurance minimums met
- **YELLOW**: AD threshold higher than playbook
- **RED**: IDERA missing

| Clause | Severity | Note |
|---|---|---|
| Insurance | GREEN | OK |
| AD | YELLOW | $2M cap |
| IDERA | RED | Absent |

### Next Steps

1. Add IDERA package
2. Lower AD threshold
3. Confirm AVN52E
""".strip()


def test_docx_writer_round_trip(tmp_path):
    out = tmp_path / "lease_review.docx"
    p = write_docx_document(
        output_path=out,
        title="Lease Review — MSN 7842",
        sections=[("Summary", SAMPLE_BODY)],
        subtitle="Aviation Contract Review (Sample)",
        disclaimer="Not legal advice. For testing only.",
    )
    assert p.is_file()

    from docx import Document

    doc = Document(str(p))
    full = "\n".join(par.text for par in doc.paragraphs)
    assert "Lease Review" in full
    assert "Not legal advice" in full
    assert "Risk Findings" in full
    # Table rendered with 4 rows (header + 3) and 3 columns.
    assert doc.tables, "table not rendered"
    table = doc.tables[0]
    assert len(table.rows) == 4
    assert len(table.rows[0].cells) == 3


def test_pdf_writer_round_trip(tmp_path):
    out = tmp_path / "risk_memo.pdf"
    p = write_pdf_document(
        output_path=out,
        title="Risk Memo — Emergency AD 2026-09-51",
        body_markdown=SAMPLE_BODY,
        subtitle="Aviation Risk Assessment (Sample)",
        disclaimer="Not legal advice. For testing only.",
    )
    assert p.is_file()

    from pdfminer.high_level import extract_text

    text = extract_text(str(p))
    assert "Risk Memo" in text
    assert "Next Steps" in text
    assert "IDERA" in text
