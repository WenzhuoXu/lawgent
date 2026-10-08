"""PRC-ready Word output: the 2026-10-07 capability review, D-10 … D-14.

Every test reproduces a measured defect and reads the produced OOXML:

* D-10 — 115 of 115 harness DOCX were US Letter with no page number;
* D-11 — ordered lists never restarted (first item of a later list: median 19)
  and nesting was stripped;
* D-12 — an English AI disclaimer was the first paragraph of 88 of 104
  Chinese documents, filings included;
* D-13 — links, footnotes, H4-H6, strikethrough and ``<br>`` reached the page
  as literal Markdown;
* D-14 — one fullwidth comma switched an English memo onto Chinese
  typography, and a fresh draft could not start from a firm template.
"""

from __future__ import annotations

import io
import zipfile

import pytest
from docx import Document
from docx.oxml.ns import qn

from legal_helper.documents.quality import lint_docx, markdown_residue
from legal_helper.documents.writers import write_docx_document

ZH_BODY = """本所受贵司委托，就拟议交易出具本法律意见书。

## 一、诉讼请求

1. 判令被告支付货款；
2. 判令被告支付违约金；
3. 判令被告承担诉讼费用。

## 二、事实与理由

1. 双方于2026年签订买卖合同；
2. 被告逾期未付款。
"""


def _xml(path, part):
    with zipfile.ZipFile(path) as zf:
        return zf.read(part).decode("utf-8")


def _footer_parts(path):
    with zipfile.ZipFile(path) as zf:
        return {n: zf.read(n).decode("utf-8") for n in zf.namelist() if n.startswith("word/footer")}


def _write(tmp_path, body, *, title="法律意见书", name="out.docx", heading="", **kwargs):
    out = tmp_path / name
    write_docx_document(output_path=out, title=title, sections=[(heading, body)], **kwargs)
    return out


def _all_text(doc):
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


# ---------------------------------------------------------------------------
# D-10 · page and footer
# ---------------------------------------------------------------------------


def test_a_chinese_document_is_a4_with_a_chinese_page_number_footer(tmp_path):
    out = _write(tmp_path, ZH_BODY)
    section = Document(str(out)).sections[0]
    assert (section.page_width.twips, section.page_height.twips) == (11906, 16838)
    footers = "".join(_footer_parts(out).values())
    assert 'w:instr="PAGE"' in footers and 'w:instr="NUMPAGES"' in footers
    assert "第 " in footers and " 页 共 " in footers


def test_an_english_document_stays_letter_with_page_x_of_y(tmp_path):
    out = _write(tmp_path, "All conditions are met.", title="Lease Review")
    section = Document(str(out)).sections[0]
    assert (section.page_width.twips, section.page_height.twips) == (12240, 15840)
    footers = "".join(_footer_parts(out).values())
    assert "Page " in footers and " of " in footers
    assert 'w:instr="PAGE"' in footers and 'w:instr="NUMPAGES"' in footers


def test_page_size_can_be_overridden(tmp_path):
    out = _write(tmp_path, "All conditions are met.", title="Lease Review", page_size="A4")
    assert Document(str(out)).sections[0].page_width.twips == 11906
    out = _write(tmp_path, ZH_BODY, name="letter.docx", page_size="Letter")
    assert Document(str(out)).sections[0].page_width.twips == 12240


def test_one_cjk_character_does_not_make_an_english_memo_chinese(tmp_path):
    # A fullwidth comma and one party name: 18 of 27 low-Han documents were
    # put on the Chinese track by exactly this.
    body = (
        "The lessee, China Eastern Airlines (中国东方航空)， shall deliver the aircraft "
        "on the agreed date and maintain it in accordance with the approved programme. "
        "All notices shall be in writing and delivered to the addresses set out below."
    )
    out = _write(tmp_path, body, title="Lease Review")
    doc = Document(str(out))
    assert doc.sections[0].page_width.twips == 12240
    rfonts = doc.styles["Normal"].element.rPr.find(qn("w:rFonts"))
    assert rfonts is None or rfonts.get(qn("w:eastAsia")) is None
    assert doc.core_properties.language == "en-US"


# ---------------------------------------------------------------------------
# D-11 · lists
# ---------------------------------------------------------------------------


def _numbered(doc):
    """(text, numId, ilvl) for every numbered body paragraph."""
    out = []
    for para in doc.paragraphs:
        ppr = para._p.pPr
        num_pr = ppr.find(qn("w:numPr")) if ppr is not None else None
        if num_pr is None:
            continue
        out.append((
            para.text,
            num_pr.find(qn("w:numId")).get(qn("w:val")),
            int(num_pr.find(qn("w:ilvl")).get(qn("w:val"))),
        ))
    return out


def _start_overrides(doc):
    root = doc.part.numbering_part.element
    out = {}
    for num in root.findall(qn("w:num")):
        for override in num.findall(qn("w:lvlOverride")):
            start = override.find(qn("w:startOverride"))
            if start is not None:
                out[num.get(qn("w:numId"))] = (int(override.get(qn("w:ilvl"))), int(start.get(qn("w:val"))))
    return out


def test_every_ordered_list_restarts_at_its_own_first_number(tmp_path):
    doc = Document(str(_write(tmp_path, ZH_BODY)))
    items = _numbered(doc)
    first_list = {num for text, num, _ in items if "判令" in text}
    second_list = {num for text, num, _ in items if text.startswith(("双方", "被告逾期"))}
    assert len(first_list) == 1 and len(second_list) == 1
    assert first_list != second_list, "a second list must not share the first list's counter"
    overrides = _start_overrides(doc)
    for num in first_list | second_list:
        assert overrides[num] == (0, 1)
    # The lint models Word's counting and agrees: nothing continues.
    findings, _ = lint_docx(tmp_path / "out.docx")
    assert "list_numbering_continues" not in [f.check for f in findings]


def test_a_list_keeps_the_number_its_author_started_it_at(tmp_path):
    doc = Document(str(_write(tmp_path, "3. 第三项请求\n4. 第四项请求")))
    (num,) = {n for _, n, _ in _numbered(doc)}
    assert _start_overrides(doc)[num] == (0, 3)


def test_nesting_is_kept_as_ilvl(tmp_path):
    body = "1. 一级\n   - 二级甲\n   - 二级乙\n2. 一级二\n   1. 编号二级\n"
    items = _numbered(Document(str(_write(tmp_path, body))))
    levels = {text: ilvl for text, _, ilvl in items}
    assert levels["一级"] == 0 and levels["一级二"] == 0
    assert levels["二级甲"] == 1 and levels["二级乙"] == 1 and levels["编号二级"] == 1


def test_list_paragraphs_carry_a_hanging_indent_not_the_body_indent(tmp_path):
    doc = Document(str(_write(tmp_path, ZH_BODY)))
    para = next(p for p in doc.paragraphs if p.text.startswith("判令被告支付货款"))
    ind = para._p.pPr.find(qn("w:ind"))
    assert int(ind.get(qn("w:hanging"))) > 0
    # Word ranks a character-based indent over a twips one; without this the
    # style's 2-character first-line indent would beat the hanging indent.
    assert ind.get(qn("w:hangingChars")) is not None


# ---------------------------------------------------------------------------
# D-12 · disclaimer
# ---------------------------------------------------------------------------


def test_the_disclaimer_is_a_document_property_not_part_of_the_instrument(tmp_path):
    notice = "本文件由法律 AI 助手生成，不构成法律意见。"
    doc = Document(str(_write(tmp_path, ZH_BODY, disclaimer=notice)))
    assert notice not in _all_text(doc)
    assert doc.core_properties.comments == notice
    assert doc.core_properties.author == ""  # not "python-docx"
    assert doc.core_properties.title == "法律意见书"
    assert doc.core_properties.language == "zh-CN"


def test_no_disclaimer_leaves_no_trace(tmp_path):
    doc = Document(str(_write(tmp_path, ZH_BODY, title="民事起诉状")))
    assert doc.core_properties.comments == ""
    assert "generated by python-docx" not in _xml(tmp_path / "out.docx", "docProps/core.xml")


# ---------------------------------------------------------------------------
# Tables
# ---------------------------------------------------------------------------


def test_table_header_repeats_and_rows_never_split(tmp_path):
    body = "| 事项 | 依据 |\n|---|---|\n" + "\n".join(f"| 事项{i} | 第{i}条 |" for i in range(40))
    table = Document(str(_write(tmp_path, body))).tables[0]
    first = table.rows[0]._tr.trPr
    assert first.find(qn("w:tblHeader")) is not None
    for row in table.rows:
        assert row._tr.trPr.find(qn("w:cantSplit")) is not None
    for row in table.rows[1:]:
        assert row._tr.trPr.find(qn("w:tblHeader")) is None


def test_table_column_alignment_is_kept(tmp_path):
    body = "| 事项 | 金额 |\n|:---|---:|\n| 货款 | 100 |"
    table = Document(str(_write(tmp_path, body))).tables[0]
    assert table.rows[1].cells[1].paragraphs[0].alignment == 2  # right


# ---------------------------------------------------------------------------
# D-13 · Markdown
# ---------------------------------------------------------------------------

RICH = """参见[《民法典》第五百七十七条](https://flk.npc.gov.cn/detail)。

依据如下[^1]，另见注释[^2]。

#### 四级标题

##### 五级标题

###### 六级标题

~~已删除的表述~~ 与 `code` 与 第一行<br>第二行。

> 引用条文原文。

[^1]: 《民法典》第一百四十三条，见[国家法律法规数据库](https://flk.npc.gov.cn/)。
[^2]: 第二个注释。
"""


def test_links_are_live_hyperlinks(tmp_path):
    out = _write(tmp_path, RICH)
    document = _xml(out, "word/document.xml")
    rels = _xml(out, "word/_rels/document.xml.rels")
    assert "<w:hyperlink" in document
    assert 'Target="https://flk.npc.gov.cn/detail"' in rels and 'TargetMode="External"' in rels
    doc = Document(str(out))
    assert "](http" not in _all_text(doc)


def test_footnotes_are_real_word_footnotes(tmp_path):
    out = _write(tmp_path, RICH)
    document = _xml(out, "word/document.xml")
    assert document.count("<w:footnoteReference") == 2
    notes = _xml(out, "word/footnotes.xml")
    assert "第一百四十三条" in notes and "第二个注释" in notes
    assert 'w:type="separator"' in notes
    # A hyperlink inside a footnote is related from the footnotes part.
    assert "flk.npc.gov.cn" in _xml(out, "word/_rels/footnotes.xml.rels")
    assert "[^1]" not in _all_text(Document(str(out)))


def test_deep_headings_strike_br_code_and_quotes_render(tmp_path):
    out = _write(tmp_path, RICH)
    doc = Document(str(out))
    styles = {p.text: p.style.name for p in doc.paragraphs}
    assert styles["四级标题"].startswith("Heading")
    assert styles["五级标题"].startswith("Heading")
    assert styles["六级标题"].startswith("Heading")
    document = _xml(out, "word/document.xml")
    assert "<w:strike/>" in document
    assert "<w:br/>" in document
    assert "Courier New" in document
    text = _all_text(doc)
    assert markdown_residue(text) == [], markdown_residue(text)


def test_cjk_emphasis_next_to_fullwidth_punctuation_is_bold(tmp_path):
    # CommonMark keeps the asterisks of **结论：**本交易; the CJK-friendly
    # delimiter rule closes them.
    doc = Document(str(_write(tmp_path, "**结论：**本交易可行。前文**「引用」**后文。")))
    para = next(p for p in doc.paragraphs if "本交易可行" in p.text)
    assert "**" not in para.text
    bold = [r.text for r in para.runs if r.bold]
    assert "结论：" in bold and "「引用」" in bold


def test_each_line_of_a_party_block_stays_its_own_paragraph(tmp_path):
    doc = Document(str(_write(tmp_path, "原告：某某公司\n被告：另一公司\n日期：2026年10月7日")))
    texts = [p.text for p in doc.paragraphs]
    assert "原告：某某公司" in texts and "被告：另一公司" in texts


def test_raw_html_other_than_line_breaks_is_printed_as_text(tmp_path):
    doc = Document(str(_write(tmp_path, "<甲方名称> 与 <Party A> 及 <script>x</script>")))
    text = _all_text(doc)
    assert "<甲方名称>" in text and "<Party A>" in text and "<script>" in text


def test_images_outside_the_allowed_roots_render_as_alt_text(tmp_path):
    from PIL import Image

    allowed = tmp_path / "outputs"
    allowed.mkdir()
    Image.new("RGB", (40, 20), (200, 0, 0)).save(allowed / "chart.png")
    elsewhere = tmp_path / "secret.png"
    Image.new("RGB", (40, 20), (0, 0, 0)).save(elsewhere)
    out = _write(
        tmp_path,
        f"![流程图](chart.png)\n\n![外部]({elsewhere})",
        image_roots=[allowed],
    )
    document = _xml(out, "word/document.xml")
    assert document.count("<pic:pic") == 1
    assert "外部" in _all_text(Document(str(out)))


# ---------------------------------------------------------------------------
# D-14 · templates
# ---------------------------------------------------------------------------


def _template(tmp_path, suffix=".docx"):
    from docx.shared import Cm

    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21.0), Cm(29.7)
    section.left_margin = Cm(3.0)
    header = section.header
    header.is_linked_to_previous = False
    header.paragraphs[0].text = "某某律师事务所"
    footer = section.footer
    footer.is_linked_to_previous = False
    footer.paragraphs[0].text = "地址：北京市"
    doc.add_paragraph("模板占位正文，应被清除。")
    path = tmp_path / f"firm{suffix}"
    if suffix == ".dotx":
        buffer = io.BytesIO()
        doc.save(buffer)
        source = zipfile.ZipFile(io.BytesIO(buffer.getvalue()))
        with zipfile.ZipFile(path, "w") as out:
            for item in source.infolist():
                data = source.read(item.filename)
                if item.filename == "[Content_Types].xml":
                    data = data.replace(
                        b"wordprocessingml.document.main+xml", b"wordprocessingml.template.main+xml"
                    )
                out.writestr(item, data)
    else:
        doc.save(str(path))
    return path


@pytest.mark.parametrize("suffix", [".docx", ".dotx"])
def test_a_reference_template_keeps_its_page_header_and_footer(tmp_path, suffix):
    template = _template(tmp_path, suffix)
    out = _write(tmp_path, ZH_BODY, reference_docx=template)
    doc = Document(str(out))
    section = doc.sections[0]
    assert round(section.left_margin.cm, 1) == 3.0
    assert section.header.paragraphs[0].text == "某某律师事务所"
    assert section.footer.paragraphs[0].text == "地址：北京市"
    text = _all_text(doc)
    assert "模板占位正文" not in text
    assert "判令被告支付货款" in text


# ---------------------------------------------------------------------------
# GB/T 9704
# ---------------------------------------------------------------------------

GB_BODY = """现将有关情况报告如下。

一、基本情况

（一）项目背景

1.项目概况

（1）具体事项

正文自然段。
"""


def test_gbt9704_page_matches_the_standard(tmp_path):
    out = _write(tmp_path, GB_BODY, title="关于某事项的请示", layout="gbt9704")
    document = _xml(out, "word/document.xml")
    assert '<w:pgSz w:w="11906" w:h="16838"/>' in document
    for attr in ('w:top="2098"', 'w:right="1474"', 'w:bottom="1984"', 'w:left="1588"', 'w:footer="1134"'):
        assert attr in document, attr
    grid = Document(str(out)).sections[0]._sectPr.find(qn("w:docGrid"))
    assert grid.get(qn("w:type")) == "linesAndChars"
    assert grid.get(qn("w:linePitch")) == "560"
    assert int(grid.get(qn("w:charSpace"))) < 0  # 28 characters a line, not 27


def test_gbt9704_structure_levels_use_the_standard_fonts(tmp_path):
    doc = Document(str(_write(tmp_path, GB_BODY, title="关于某事项的请示", layout="gbt9704")))
    expected = {"一、基本情况": ("Heading 1", "黑体"), "（一）项目背景": ("Heading 2", "楷体"),
                "1.项目概况": ("Heading 3", "仿宋"), "（1）具体事项": ("Heading 4", "仿宋")}
    for para in doc.paragraphs:
        if para.text in expected:
            style, font = expected.pop(para.text)
            assert para.style.name == style
            assert f'w:eastAsia="{font}"' in para._p.xml
            assert 'w:firstLine="640"' in para._p.xml  # 空二字
    assert not expected
    for level in range(1, 5):
        assert doc.styles[f"Heading {level}"].font.bold is False
    title = doc.styles["Title"]
    assert 'w:eastAsia="方正小标宋简体"' in title.element.xml
    assert title.font.size.pt == 22 and title.font.bold is False


def test_gbt9704_page_numbers_are_dashed_and_outside_aligned(tmp_path):
    out = _write(tmp_path, GB_BODY, title="关于某事项的请示", layout="gbt9704")
    footers = _footer_parts(out)
    odd, even = footers["word/footerOdd.xml"], footers["word/footerEven.xml"]
    for xml in (odd, even):
        assert "— " in xml and " —" in xml and 'w:instr="PAGE"' in xml
        assert 'w:sz w:val="28"' in xml  # 四号
    assert '<w:jc w:val="right"/>' in odd and 'w:right="280"' in odd
    assert '<w:jc w:val="left"/>' in even and 'w:left="280"' in even
    assert "<w:evenAndOddHeaders/>" in _xml(out, "word/settings.xml")


def test_a_template_without_lists_still_takes_a_numbered_list(tmp_path):
    # python-docx cannot create a missing numbering part on its own.
    import re

    source = _write(tmp_path, "正文。", name="plain.docx")
    stripped = tmp_path / "no-numbering.docx"
    with zipfile.ZipFile(source) as src, zipfile.ZipFile(stripped, "w") as out:
        for item in src.infolist():
            if item.filename == "word/numbering.xml":
                continue
            data = src.read(item.filename)
            if item.filename == "word/_rels/document.xml.rels":
                data = re.sub(rb"<Relationship [^>]*numbering[^>]*/>", b"", data)
            if item.filename == "[Content_Types].xml":
                data = re.sub(rb'<Override PartName="/word/numbering.xml"[^>]*/>', b"", data)
            out.writestr(item, data)
    doc = Document(str(_write(tmp_path, "1. 甲\n2. 乙", name="out2.docx", reference_docx=stripped)))
    assert [t for t, _, _ in _numbered(doc)] == ["甲", "乙"]


def test_a_templates_own_notes_and_comments_go_with_its_body(tmp_path):
    template = _write(tmp_path, "模板正文[^1]。\n\n[^1]: 模板注释。", name="tpl.docx")
    out = _write(tmp_path, "新正文[^1]。\n\n[^1]: 新注释。", name="new.docx", reference_docx=template)
    notes = _xml(out, "word/footnotes.xml")
    assert "新注释" in notes and "模板注释" not in notes
    assert 'w:type="separator"' in notes


def test_a_chinese_title_has_no_accent_rule_under_it(tmp_path):
    doc = Document(str(_write(tmp_path, ZH_BODY)))
    ppr = doc.styles["Title"].element.find(qn("w:pPr"))
    assert ppr is None or ppr.find(qn("w:pBdr")) is None
