"""create_deck: the one deck path, driven through ppt-master.

The sub-agent itself is a live model run, so these tests replace it with a
fake that does on disk what ppt-master does — a project under the work
directory, per-page SVGs, a postflight report beside an export — and check the
part this tool owns: what the sub-agent is told, where its files go, how the
export is verified, and what the caller and the UI are shown. Nothing here
calls a model.
"""

from __future__ import annotations

import json
import os
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Inches, Pt

from legal_helper.config import load_settings, override_current_settings
from legal_helper.logging_setup import agent_name_var, log_workflow_event, workflow_event_sink_var
from legal_helper.tools import decks


def _call(tool, **kwargs):
    return (getattr(tool, "func", None) or tool)(**kwargs)


@pytest.fixture()
def env(tmp_path, monkeypatch):
    """Request-scoped settings with every directory in tmp, plus an event log."""
    settings = load_settings(refresh=True).model_copy(
        update={
            "outputs_dir": tmp_path / "out",
            "state_dir": tmp_path / "state",
            "logs_dir": tmp_path / "logs",
        }
    )
    settings.outputs_dir.mkdir(parents=True)
    fake_skill = tmp_path / "skills" / "ppt-master"
    fake_skill.mkdir(parents=True)
    monkeypatch.setattr(decks, "_deck_skill_dir", lambda: fake_skill)
    events: list[dict] = []
    token = workflow_event_sink_var.set(events.append)
    with override_current_settings(settings):
        yield settings, events
    workflow_event_sink_var.reset(token)


def _ppt_master_deck(path: Path, slides: int = 3) -> Path:
    """A deck shaped like ppt-master's: every text frame inside a group."""
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    for index in range(slides):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        group = slide.shapes.add_group_shape()
        box = group.shapes.add_textbox(Inches(0.6), Inches(0.5), Inches(10), Inches(0.8))
        box.text_frame.text = f"第{index + 1}页：个人信息出境合规"
        box.text_frame.paragraphs[0].runs[0].font.size = Pt(28)
    path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(path))
    return path


def _emit_call(tool: str, arguments: dict, output: str) -> None:
    """One tool call as both providers log it, attributed to the sub-agent."""
    token = agent_name_var.set("ppt-master")
    try:
        log_workflow_event("local_tool_call_started", {"tool_name": tool, "arguments": arguments})
        log_workflow_event("local_tool_call_finished", {"tool_name": tool, "output_preview": output})
    finally:
        agent_name_var.reset(token)


REPORT = """Deck exported.

- PPTX: {pptx}
- Final quality gate: passed

## Decisions under delegation
- Route: Default Generate; the request named no quick intent.
- Free design: no template was supplied.
- Canvas ppt169, five pages, one per finding.
"""


def _fake_agent(*, slides: int = 3, status: str = "passed-with-warnings", export: bool = True):
    """Stand-in for the ppt-master sub-agent: does on disk what the skill does."""
    seen: dict = {}

    def run(task: str, settings) -> str:
        seen["task"] = task
        seen["outputs_dir"] = settings.outputs_dir
        project = settings.outputs_dir / "pi_export_20261007"
        for page in range(1, slides + 1):
            svg = project / "svg_output" / f"{page:02d}_page.svg"
            _emit_call("write_text_file", {"path": str(svg), "content": "<svg/>"}, str(svg))
            svg.parent.mkdir(parents=True, exist_ok=True)
            svg.write_text("<svg/>", encoding="utf-8")
        _emit_call(
            "run_skill_script",
            {"skill": "ppt-master", "script": "scripts/svg_quality_checker.py", "args": [str(project), "--stage", "final"]},
            json.dumps({"exit_code": 0, "stdout": "ok"}),
        )
        if not export:
            return "Stopped: the final quality gate kept failing."
        pptx = _ppt_master_deck(project / "exports" / "pi_export_20261007_120000.pptx", slides)
        # A copied source deck is a .pptx too, and must never be taken for the export.
        _ppt_master_deck(project / "sources" / "uploaded.pptx", 1)
        (project / "validation").mkdir(parents=True, exist_ok=True)
        (project / "validation" / f"{pptx.stem}.report.json").write_text(
            json.dumps({"status": status, "checks": {"quality_gate": "passed"}}), encoding="utf-8"
        )
        (project / "validation" / "svg_quality_report.json").write_text(
            json.dumps({"summary": {"total": slides, "passed": slides, "warnings": 0, "errors": 0}}),
            encoding="utf-8",
        )
        _emit_call(
            "run_skill_script",
            {"skill": "ppt-master", "script": "scripts/svg_to_pptx.py", "args": [str(project)]},
            json.dumps({"exit_code": 0}),
        )
        return REPORT.format(pptx=pptx)

    return run, seen


# ---------------------------------------------------------------------------
# create_deck
# ---------------------------------------------------------------------------


def test_without_ppt_master_the_tool_says_so_and_writes_nothing(env, monkeypatch):
    settings, _ = env
    monkeypatch.setattr(decks, "_deck_skill_dir", lambda: None)
    called = []
    monkeypatch.setattr(decks, "run_deck_agent", lambda *a, **k: called.append(1) or "")

    result = json.loads(_call(decks.create_deck, brief="内容", title="汇报"))

    assert result["status"] == "unavailable"
    assert result["pptx_path"] is None
    assert "git clone https://github.com/hugohe3/ppt-master" in result["next_step"]
    assert not called, "no fallback writer, and no sub-agent either"
    assert list(settings.outputs_dir.iterdir()) == []


def test_only_the_verified_export_reaches_the_chat(env, monkeypatch):
    settings, events = env
    fake, seen = _fake_agent(slides=3)
    monkeypatch.setattr(decks, "run_deck_agent", fake)

    result = json.loads(
        _call(decks.create_deck, brief="个人信息出境的三条路径……", title="个人信息出境合规汇报", page_count=3)
    )

    assert result["status"] == "ok", result
    delivered = Path(result["pptx_path"])
    assert delivered == (settings.outputs_dir / "个人信息出境合规汇报.pptx").resolve()
    assert delivered.is_file()
    assert result["slides"] == 3
    assert result["gate_verdict"]["export"] == "passed-with-warnings"
    assert result["gate_verdict"]["quality_gate"] == "passed"
    assert result["gate_verdict"]["pages_checked"]["errors"] == 0
    assert result["decisions_under_delegation"][0].startswith("Route: Default Generate")
    assert result["lint_summary"]["conformance"] in {"pass", "warnings", "blocked"}
    # The project is internal: under _decks/, where the artifact scan does not look,
    # and the sub-agent's work directory is that project, not the chat's outputs.
    project = Path(result["project_dir"])
    assert project.parent == (settings.outputs_dir / "_decks").resolve()
    assert seen["outputs_dir"] == project
    top_level = sorted(p.name for p in settings.outputs_dir.iterdir())
    assert top_level == ["_decks", "个人信息出境合规汇报.pptx"]
    # One progress event per sub-agent tool call, with pages counted.
    progress = [e["data"] for e in events if e["event_type"] == "deck_progress"]
    assert [p["stage"] for p in progress] == ["page_authored"] * 3 + ["quality_gate", "export"]
    assert progress[2]["current"] == 3 and progress[2]["total"] == 3
    assert progress[3]["exit_code"] == 0 and progress[3]["ok"] is True
    assert result["tool_calls"] == 5 and result["pages_authored"] == 3
    # The sub-agent's own tool events still reach the stream unchanged.
    assert sum(e["event_type"] == "local_tool_call_finished" for e in events) == 5


def test_a_repeated_title_never_overwrites_a_delivered_deck(env, monkeypatch):
    settings, _ = env
    fake, _ = _fake_agent()
    monkeypatch.setattr(decks, "run_deck_agent", fake)
    first = json.loads(_call(decks.create_deck, brief="b", title="Board update"))
    second = json.loads(_call(decks.create_deck, brief="b", title="Board update"))
    assert first["pptx_path"] != second["pptx_path"]
    assert Path(first["pptx_path"]).is_file() and Path(second["pptx_path"]).is_file()


def test_no_export_is_a_failed_deck_with_the_reason(env, monkeypatch):
    settings, _ = env
    fake, _ = _fake_agent(export=False)
    monkeypatch.setattr(decks, "run_deck_agent", fake)

    result = json.loads(_call(decks.create_deck, brief="b", title="T"))

    assert result["status"] == "failed"
    assert result["pptx_path"] is None
    assert "final quality gate kept failing" in result["sub_agent_report"]
    assert not list(settings.outputs_dir.glob("*.pptx"))


def test_a_postflight_that_did_not_pass_is_a_failed_deck(env, monkeypatch):
    fake, _ = _fake_agent(status="failed")
    monkeypatch.setattr(decks, "run_deck_agent", fake)
    result = json.loads(_call(decks.create_deck, brief="b", title="T"))
    assert result["status"] == "failed"
    assert "postflight" in result["next_step"]


def test_a_sub_agent_crash_is_reported_not_raised(env, monkeypatch):
    def boom(task, settings):
        raise RuntimeError("provider unavailable")

    monkeypatch.setattr(decks, "run_deck_agent", boom)
    result = json.loads(_call(decks.create_deck, brief="b", title="T"))
    assert result["status"] == "failed"
    assert "provider unavailable" in result["error"]


def test_missing_inputs_are_refused_before_any_work(env, monkeypatch):
    called = []
    monkeypatch.setattr(decks, "run_deck_agent", lambda *a, **k: called.append(1) or "")
    empty = json.loads(_call(decks.create_deck, brief="  ", title="T"))
    assert empty["status"] == "failed" and "cannot see this conversation" in empty["next_step"]
    gone = json.loads(_call(decks.create_deck, brief="b", title="T", source_paths=["/nope/x.pdf"]))
    assert gone["status"] == "failed" and "/nope/x.pdf" in gone["next_step"]
    assert not called


def test_the_task_is_self_contained(env, monkeypatch, tmp_path):
    settings, _ = env
    fake, seen = _fake_agent()
    monkeypatch.setattr(decks, "run_deck_agent", fake)
    source = tmp_path / "uploads" / "评估报告.pdf"
    source.parent.mkdir()
    source.write_bytes(b"%PDF-1.4")
    template = tmp_path / "uploads" / "firm.potx"
    template.write_bytes(b"PK")

    result = json.loads(
        _call(
            decks.create_deck,
            brief="结论：标准合同路径适用于本项目。",
            title="出境路径",
            audience="管理层",
            purpose="汇报",
            quick=True,
            source_paths=[str(source)],
            template_path=str(template),
        )
    )

    task = seen["task"]
    project = result["project_dir"]
    assert "结论：标准合同路径适用于本项目。" in task
    assert f"--dir {project}" in task, "init must be pointed at the work directory"
    assert "`zh-CN`" in task, "deck language defaults to the brief's"
    assert "管理层" in task and "汇报" in task
    assert "quick generation" in task
    assert str(source.resolve()) in task and str(template.resolve()) in task
    assert "closed corpus" in task
    assert "## Decisions under delegation" in task
    brief_file = Path(project) / "brief.md"
    assert brief_file.read_text(encoding="utf-8").startswith("# 出境路径")
    assert str(brief_file) in task
    # The user's uploaded source and template are never taken for the export.
    assert Path(result["pptx_path"]).parent == settings.outputs_dir.resolve()


def test_deck_language_follows_the_brief():
    assert decks.deck_language("个人信息出境合规：三条路径的适用条件与时间表") == "zh-CN"
    assert decks.deck_language("Board update on the PIPL export rules (《个人信息保护法》 art. 38)") == "en"


def test_find_export_prefers_exports_and_ignores_copied_sources(tmp_path):
    workspace = tmp_path / "ws"
    newest_source = _ppt_master_deck(workspace / "proj" / "sources" / "upload.pptx", 1)
    export = _ppt_master_deck(workspace / "proj" / "exports" / "proj_1.pptx", 2)
    os.utime(newest_source, (export.stat().st_mtime + 60,) * 2)
    assert decks.find_export(workspace.resolve(), "", set()) == export.resolve()
    # A path the sub-agent names outside the workspace is not trusted.
    stray = _ppt_master_deck(tmp_path / "elsewhere.pptx", 1)
    assert decks.find_export(workspace.resolve(), f"saved to {stray}", set()) == export.resolve()


def test_progress_names_the_step_each_call_is():
    classify = decks.classify_deck_call
    assert classify("write_text_file", {"path": "/w/p/svg_output/03_risk.svg"}) == ("page_authored", "03_risk.svg")
    assert classify("write_text_file", {"path": "/w/p/design_spec.md"}) == ("plan_written", "design_spec.md")
    assert classify("run_skill_script", {"script": "scripts/project_manager.py", "args": ["init", "x"]})[0] == "project_init"
    assert classify(
        "run_skill_script", {"script": "scripts/svg_quality_checker.py", "args": ["/p", "--stage", "early"]}
    ) == ("quality_gate", "svg_quality_checker.py (early)")
    assert classify("run_skill_script", {"script": "scripts/svg_to_pptx.py"})[0] == "export"
    assert classify("read_document", {"path": "/s/references/executor-base.md"}) == ("reading", "executor-base.md")


def test_progress_reports_a_failing_gate_as_not_ok():
    seen: list[dict] = []
    sink = decks.DeckProgress(seen.append)
    token = workflow_event_sink_var.set(sink)
    try:
        _emit_call(
            "run_skill_script",
            {"script": "scripts/svg_quality_checker.py", "args": ["/p", "--stage", "final"]},
            json.dumps({"exit_code": 1}),
        )
    finally:
        workflow_event_sink_var.reset(token)
    (progress,) = [e["data"] for e in seen if e["event_type"] == "deck_progress"]
    assert progress["ok"] is False and progress["exit_code"] == 1


# ---------------------------------------------------------------------------
# Surfaces
# ---------------------------------------------------------------------------


def _names(tools) -> set[str]:
    return {getattr(t, "name", getattr(t, "__name__", "")) for t in tools}


def test_create_deck_is_on_every_authoring_surface():
    from legal_helper.tools import skill_tools_for_task
    from legal_helper.workflow import _document_io_tools

    assert "create_deck" in _names(_document_io_tools())
    for skill in ("brief", "litigation-analysis", "flowchart"):
        assert "create_deck" in _names(skill_tools_for_task(skill)), skill


def test_create_deck_is_mutating_and_kept_from_auditors_and_from_ppt_master_itself():
    from legal_helper.tool_policy import is_mutating
    from legal_helper.tools import _external_skill_tools, skill_tools_for_task

    assert is_mutating("create_deck")
    assert "create_deck" not in _names(skill_tools_for_task("cite-check"))
    # The deck sub-agent drives ppt-master directly; handing it create_deck
    # would let it recurse into another deck sub-agent.
    assert "create_deck" not in _names(_external_skill_tools())


def test_the_tool_description_says_it_is_the_way_to_make_slides():
    description = decks.create_deck.to_dict()["description"]
    for phrase in ("THE way", "汇报材料", "课件", "single slide", "does not see this conversation"):
        assert phrase in description, phrase


def test_the_host_contract_keeps_work_out_of_the_skill_and_asks_for_decisions(tmp_path, monkeypatch):
    from legal_helper.agent import _external_skill_prompt

    skill = tmp_path / "root" / "probe-proc"
    skill.mkdir(parents=True)
    (skill / "SKILL.md").write_text("---\nname: probe-proc\n---\nprocedure\n")
    monkeypatch.setenv("LEGAL_HELPER_SKILL_ROOTS", str(tmp_path / "root"))
    prompt = _external_skill_prompt("probe-proc", tmp_path / "work")
    assert "read-only" in prompt
    assert "## Decisions under delegation" in prompt
    assert "offset=<line>" in prompt, "a long reference must be read to the end"
    assert "environment variables" in prompt


def test_the_planner_never_routes_a_deck_to_a_specialist():
    """ppt-master is reached only through create_deck: the planner does not
    offer it as a task, and a plan that names it becomes general work."""
    from legal_helper.workflow import WorkflowExecutor, _routable_skills

    assert "ppt-master" not in {name for name, _ in _routable_skills()}

    class _P:
        name = "mock"
        model = "mock"

    plan = WorkflowExecutor(provider=_P())._parse_plan_data(
        {"agent_tasks": [{"skill_name": "ppt-master", "title": "Deck", "task": "make slides"}]}
    )
    assert plan.agent_tasks[0].skill_name == "general-answer"


def test_deck_docs_name_only_tools_the_reader_can_call():
    """The documented default once named tools absent from the surface reading it.

    The HTML deck path was "Preferred for new decks" in a recipe the main author
    could read but could not act on — it was never called. Every tool the deck
    recipe or the flowchart skill names must be on the surface that reads it,
    and none of the deleted deck writers may be named at all.
    """
    import re
    from pathlib import Path

    import legal_helper
    from legal_helper.skills import internal_skill_names, load_skill_text
    from legal_helper.tools import _external_skill_tools, skill_tools_for_task
    from legal_helper.workflow import _document_io_tools

    known = _names(_document_io_tools()) | _names(_external_skill_tools())
    for skill in internal_skill_names():
        known |= _names(skill_tools_for_task(skill))

    def named_tools(text: str) -> set[str]:
        return set(re.findall(r"`([a-z][a-z0-9_]+)(?:\(|`)", text)) & known

    root = Path(legal_helper.__file__).resolve().parent
    recipe = (root / "playbook" / "format_recipes" / "pptx.md").read_text(encoding="utf-8")
    flowchart = load_skill_text("flowchart")
    for text, surface in (
        (recipe, _names(_document_io_tools())),
        (flowchart, _names(skill_tools_for_task("flowchart"))),
    ):
        referenced = named_tools(text)
        assert "create_deck" in referenced
        assert referenced <= surface, sorted(referenced - surface)
        for gone in ("write_pptx", "write_pptx_from_html", "read_deck_stylesheet", "SlideFlowchart"):
            assert gone not in text, gone
