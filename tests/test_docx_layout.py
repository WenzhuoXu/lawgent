"""Table and indent layout contract for the .docx writer.

The 首行缩进 is a body-prose rule: headings, table cells and centred lines
cancel it, and markdown tables get content-weighted columns instead of an
equal split. Every assertion here corresponds to a defect measured at 150dpi
in outputs/harness_reviews/2026-09-21-document-visual-review.md.
"""

from __future__ import annotations

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Emu

from legal_helper.documents.writers import write_docx_document
from legal_helper.documents.writers.docx import _CELL_PAD_TWIPS, _column_widths


TABLE_BODY = """| 事项 | 依据 | 时限 |
|---|---|---|
| 经营许可 | 《民用航空法》第九十二条 | 60 日 |
| 航线许可 | 《国际航线管理规定》第十二条 | 30 日 |
"""

LATIN_BODY = """### Risk Findings

| Clause | Rating | Note |
|---|---|---|
| Warranty | Green | Conforms |
| Indemnity | Amber | Cap missing |
"""


def _ind(p_or_style):
    ppr = p_or_style.find(qn("w:pPr"))
    return None if ppr is None else ppr.find(qn("w:ind"))


def _jc(paragraph):
    ppr = paragraph._p.find(qn("w:pPr"))
    if ppr is None:
        return None
    jc = ppr.find(qn("w:jc"))
    return None if jc is None else jc.get(qn("w:val"))


def _grid_widths(table):
    grid = table._tbl.find(qn("w:tblGrid"))
    return [int(c.get(qn("w:w"))) for c in grid.findall(qn("w:gridCol"))]


def _text_width_twips(doc):
    s = doc.sections[0]
    return Emu(s.page_width - s.left_margin - s.right_margin).twips


def _cjk_memo(tmp_path, name="memo.docx", **kwargs):
    out = tmp_path / name
    write_docx_document(
        output_path=out,
        title="法律意见书",
        sections=[("一、结论", "拟议交易不存在实质性法律障碍。\n\n" + TABLE_BODY)],
        **kwargs,
    )
    return Document(str(out))


def test_table_cells_do_not_inherit_the_body_indent(tmp_path):
    # The regression that broke 《民用航空法》第九十二条 across two lines: the
    # cell paragraph inherited Normal's 2-character indent from the style
    # hierarchy, which a table style cannot override.
    doc = _cjk_memo(tmp_path)
    cells = [
        cell
        for row in doc.tables[0].rows
        for cell in row.cells
    ]
    assert cells
    for cell in cells:
        for para in cell.paragraphs:
            ind = _ind(para._p)
            assert ind is not None, "cell paragraph needs direct w:ind"
            assert ind.get(qn("w:firstLineChars")) == "0"
            assert ind.get(qn("w:firstLine")) == "0"


def test_table_is_fixed_layout_and_fills_the_text_width(tmp_path):
    doc = _cjk_memo(tmp_path)
    table = doc.tables[0]
    tbl_pr = table._tbl.tblPr
    assert tbl_pr.find(qn("w:tblLayout")).get(qn("w:type")) == "fixed"

    total = _text_width_twips(doc)
    tbl_w = tbl_pr.find(qn("w:tblW"))
    assert tbl_w.get(qn("w:type")) == "dxa"
    assert int(tbl_w.get(qn("w:w"))) == total

    widths = _grid_widths(table)
    assert sum(widths) == total
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            # A stale per-cell tcW overrides the grid in Word.
            assert cell.width.twips == width


def test_long_citation_column_is_wider_than_its_siblings(tmp_path):
    doc = _cjk_memo(tmp_path)
    widths = _grid_widths(doc.tables[0])
    assert widths[1] > widths[0]
    assert widths[1] > 2 * widths[2], "依据 must beat 时限 by more than the old equal split"


def test_heading_styles_do_not_inherit_the_body_indent(tmp_path):
    doc = _cjk_memo(tmp_path)
    for level in range(1, 5):
        ind = _ind(doc.styles[f"Heading {level}"].element)
        assert ind is not None, f"Heading {level} needs its own w:ind"
        assert ind.get(qn("w:firstLineChars")) == "0"
        assert ind.get(qn("w:firstLine")) == "0"
    # Body prose keeps the indent — the hierarchy is the point.
    assert _ind(doc.styles["Normal"].element).get(qn("w:firstLineChars")) == "200"


def test_gbt9704_headings_keep_the_two_character_indent(tmp_path):
    # GB/T 9704-2012 sets 各层次标题 空两格, so that layout keeps the
    # inheritance while the ordinary memo track goes 顶格.
    doc = _cjk_memo(tmp_path, name="gb.docx", layout="gbt9704")
    for level in (1, 2):
        assert _ind(doc.styles[f"Heading {level}"].element) is None
    assert _ind(doc.styles["Normal"].element).get(qn("w:firstLineChars")) == "200"


def test_centred_paragraphs_have_no_first_line_indent(tmp_path):
    out = tmp_path / "watermarked.docx"
    write_docx_document(
        output_path=out,
        title="法律意见书",
        sections=[("一、结论", "第一段。\n\n---\n\n第二段。")],
        subtitle="仅供内部讨论",
        watermark="DO NOT FILE · 禁止对外提交",
    )
    doc = Document(str(out))
    centred = [
        p for p in doc.paragraphs if _jc(p) == "center" and p.text.strip()
    ]
    centred += [
        p for p in doc.sections[0].header.paragraphs if _jc(p) == "center"
    ]
    assert len(centred) >= 4, "banner, title, subtitle, rule and header expected"
    for p in centred:
        ind = _ind(p._p) if _ind(p._p) is not None else _ind(p.style.element)
        assert ind is not None and ind.get(qn("w:firstLine")) == "0", p.text


def test_column_widths_clamp_and_normalise():
    em = 12.0 * 20
    floor = 3.0 * em + _CELL_PAD_TWIPS

    wide_row = ["x" * 80, "a", "b", "c", "dd", "e"]  # 40 ems in column 0
    widths = _column_widths([wide_row], 6, 9360, 12.0)
    assert sum(widths) == 9360
    assert min(widths) >= int(floor) - 1

    widths = _column_widths([["经" * 20, "甲"]], 2, 9360, 12.0)
    assert widths[0] > 3 * widths[1]
    assert sum(widths) == 9360

    widths = _column_widths([["甲", "乙", "丙"]], 3, 9360, 12.0)
    assert sum(widths) == 9360, "slack is shared out, not dropped"


def test_latin_tables_are_content_weighted_too(tmp_path):
    out = tmp_path / "latin.docx"
    write_docx_document(
        output_path=out,
        title="Lease Review",
        sections=[("Summary", LATIN_BODY)],
    )
    doc = Document(str(out))
    widths = _grid_widths(doc.tables[0])
    assert len(set(widths)) > 1, "equal columns are the defect"
    assert sum(widths) == _text_width_twips(doc)
    assert _ind(doc.styles["Normal"].element) is None, "no first-line indent for Latin docs"
    assert doc.tables[0].style.name == "Light Grid Accent 1"


def test_blockquote_keeps_its_left_indent_without_the_first_line_indent(tmp_path):
    out = tmp_path / "quote.docx"
    write_docx_document(
        output_path=out,
        title="法律意见书",
        sections=[("一、结论", "> 引用条文原文。")],
    )
    doc = Document(str(out))
    quoted = [p for p in doc.paragraphs if p.text.strip() == "引用条文原文。"]
    assert quoted
    ind = _ind(quoted[0]._p)
    assert int(ind.get(qn("w:left"))) == 576  # 0.4in
    assert ind.get(qn("w:firstLine")) == "0"
