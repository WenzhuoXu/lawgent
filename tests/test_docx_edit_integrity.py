"""DOCX edits that keep what they do not touch, and say what they lost.

Repros from review item D-4 ("Edit tools corrupt text silently") and D-9:
- a cross-run match collapsed into ``runs[0]`` and deleted the runs holding
  footnote and comment references, while reporting success;
- ``replace_paragraph`` blanked only direct runs, so hyperlink and ``w:ins``
  text survived at the end of the "replaced" paragraph;
- ``edit_docx_text`` walked ``row.cells``, editing a merged cell twice;
- 7 of 121 edit calls overwrote their own source file.
"""

from __future__ import annotations

import json
from pathlib import Path

from docx import Document

from docx_fixtures import build_docx, document_xml, ins, p, r
from legal_helper.documents.integrity import compare_docx_integrity
from legal_helper.tools import edit_docx_text, reshape_docx

_COMMENT = (
    '<w:comment w:id="1" w:author="Reviewer" w:date="2026-09-01T00:00:00Z">'
    "<w:p><w:r><w:t>check amount</w:t></w:r></w:p></w:comment>"
)
_FOOTNOTE = '<w:footnote w:id="1"><w:p><w:r><w:t>Deposit note</w:t></w:r></w:p></w:footnote>'


def _source(tmp_path: Path) -> Path:
    body = (
        p("Heading One", style="Heading1")
        + p(
            "",
            runs=r("The Lessee shall pay ")
            + '<w:commentRangeStart w:id="1"/>'
            + r("rent")
            + '<w:commentRangeEnd w:id="1"/><w:r><w:commentReference w:id="1"/></w:r>'
            + r(" monthly in ")
            + r("advance."),
        )
        + p("", runs=r("See note") + '<w:r><w:footnoteReference w:id="1"/></w:r>' + r(" for the ") + r("deposit") + r(" terms."))
        + p("", runs=r("Visit ") + '<w:hyperlink r:id="rIdH1"><w:r><w:t>our site</w:t></w:r></w:hyperlink>' + r(" today."))
        + p("", runs=r("Pending ") + ins("insertion") + r(" here."))
        + '<w:tbl><w:tr><w:tc><w:tcPr><w:gridSpan w:val="2"/></w:tcPr><w:p><w:r><w:t>Merged Party A</w:t></w:r></w:p></w:tc></w:tr>'
        '<w:tr><w:tc><w:p><w:r><w:t>Party A</w:t></w:r></w:p></w:tc><w:tc><w:p><w:r><w:t>x</w:t></w:r></w:p></w:tc></w:tr></w:tbl>'
        + p("Last paragraph.")
    )
    return build_docx(
        tmp_path / "src.docx",
        body,
        footnotes=_FOOTNOTE,
        comments=_COMMENT,
        hyperlinks={"rIdH1": "https://example.com"},
    )


def _edit(tmp_path: Path, replacements: list[dict], filename: str = "out.docx") -> dict:
    return json.loads(
        edit_docx_text.call(
            {"source_path": str(_source(tmp_path)), "filename": filename, "replacements": replacements}
        )
    )


def test_cross_run_edit_keeps_footnote_and_comment_references(tmp_path):
    payload = _edit(
        tmp_path,
        [
            {"find": "rent monthly", "replace": "RENT QUARTERLY"},
            {"find": "note for the deposit", "replace": "NOTE FOR THE DEPOSIT"},
        ],
    )
    assert payload["replacements"] == {"rent monthly": 1, "note for the deposit": 1}
    xml = document_xml(Path(payload["output_path"]))
    assert '<w:commentReference w:id="1"/>' in xml
    assert '<w:footnoteReference w:id="1"/>' in xml
    texts = [para.text for para in Document(payload["output_path"]).paragraphs]
    assert texts[1] == "The Lessee shall pay RENT QUARTERLY in advance."
    assert texts[2] == "See NOTE FOR THE DEPOSIT terms."
    assert payload["integrity"]["status"] == "ok", payload["integrity"]


def test_merged_cell_is_edited_once(tmp_path):
    # A replacement that contains its own find shows a double edit: python-docx
    # repeats a merged cell in row.cells, and the old loop gave
    # "Merged Party A Ltd Ltd".
    payload = _edit(tmp_path, [{"find": "Party A", "replace": "Party A Ltd"}])
    assert payload["replacements"]["Party A"] == 2  # merged cell + plain cell
    table = Document(payload["output_path"]).tables[0]
    assert table.rows[0].cells[0].text == "Merged Party A Ltd"
    assert table.rows[1].cells[0].text == "Party A Ltd"


def test_paragraph_with_tracked_changes_is_refused_not_rewritten(tmp_path):
    payload = _edit(tmp_path, [{"find": "insertion", "replace": "X"}])
    assert payload["replacements"]["insertion"] == 0
    refused = payload["refused"][0]
    assert refused["paragraph"] == 4
    assert "redline_docx" in refused["reason"] and "Counterparty" in refused["reason"]
    assert '<w:ins w:id="901"' in document_xml(Path(payload["output_path"]))


def test_not_found_carries_a_hint(tmp_path):
    payload = _edit(tmp_path, [{"find": "nowhere to be found", "replace": "x"}])
    assert payload["not_found"][0]["find"] == "nowhere to be found"
    assert payload["not_found"][0]["hint"]


def test_replace_paragraph_clears_hyperlink_and_leaves_no_stale_text(tmp_path):
    src = _source(tmp_path)
    out = json.loads(
        reshape_docx.call(
            {
                "source_path": str(src),
                "filename": "rp.docx",
                "operations": [{"op": "replace_paragraph", "index": 3, "text": "Visit the new site."}],
            }
        )
    )
    para = Document(out["path"]).paragraphs[3]
    xml = document_xml(Path(out["path"]))
    assert para.text == "Visit the new site."
    assert "our site" not in xml and " today." not in xml


def test_reshape_refuses_to_rewrite_tracked_revisions(tmp_path):
    result = reshape_docx.call(
        {
            "source_path": str(_source(tmp_path)),
            "filename": "bad.docx",
            "operations": [{"op": "replace_paragraph", "index": 4, "text": "nope"}],
        }
    )
    assert result.startswith("ERROR") and "redline_docx" in result


def test_edit_never_overwrites_its_own_source(tmp_path, monkeypatch):
    """D-9: 7 of 121 edit calls named their source as the output."""
    from legal_helper.config import current_settings

    outputs = current_settings().outputs_dir
    outputs.mkdir(parents=True, exist_ok=True)
    src = build_docx(outputs / "contract.docx", p("Party A pays."))
    before = src.read_bytes()
    payload = json.loads(
        edit_docx_text.call(
            {"source_path": str(src), "filename": "contract.docx", "replacements": [{"find": "Party A", "replace": "Party B"}]}
        )
    )
    assert src.read_bytes() == before
    assert payload["output_path"].endswith("contract_v2.docx")
    assert "contract_v2.docx" in payload["note"]
    assert Document(payload["output_path"]).paragraphs[0].text == "Party B pays."


def test_integrity_reports_what_a_lossy_edit_lost(tmp_path):
    """The gate itself: deleting the run that holds a comment reference
    orphans the comment, and the report must say so."""
    src = _source(tmp_path)
    doc = Document(str(src))
    para = doc.paragraphs[1]
    for run in list(para.runs):
        run._r.getparent().remove(run._r)
    para.add_run("The Lessee shall pay rent.")
    lossy = tmp_path / "lossy.docx"
    doc.save(str(lossy))
    report = compare_docx_integrity(src, lossy)
    assert report["status"] == "issues"
    messages = " ".join(f["message"] for f in report["findings"])
    assert "comment anchors" in messages
    clean = compare_docx_integrity(src, src)
    assert clean["status"] == "ok" and not clean["findings"]


def test_in_tree_counts_cover_bookmarks_fields_and_sections(tmp_path):
    """Constructs ooxml-integrity does not count are counted here."""
    body = (
        p("", runs='<w:bookmarkStart w:id="5" w:name="Clause2"/>' + r("Bookmarked") + '<w:bookmarkEnd w:id="5"/>')
        + p("", runs='<w:fldSimple w:instr=" PAGE "><w:r><w:t>1</w:t></w:r></w:fldSimple>')
        + '<w:p><w:pPr><w:sectPr/></w:pPr><w:r><w:t>End of section one</w:t></w:r></w:p>'
        + p("Tail")
    )
    src = build_docx(tmp_path / "marks.docx", body)
    doc = Document(str(src))
    for para in doc.paragraphs[:3]:
        para._p.getparent().remove(para._p)
    stripped = tmp_path / "stripped.docx"
    doc.save(str(stripped))
    findings = {f["check"]: f for f in compare_docx_integrity(src, stripped)["findings"]}
    assert findings["bookmarks"]["lost"] == ["Clause2"]
    assert "simple_fields" in findings
    assert findings["section_breaks"]["severity"] == "error"


def test_inserted_ordered_list_starts_at_one(tmp_path):
    """A paragraph inserted with the List Number style continued the
    document's last ordered list (Word counts per abstract numbering), so a
    new list opened at 4; it now gets its own w:num unless the paragraph above
    is an item of the same list."""
    from legal_helper.documents.docx_reader import read_docx

    doc = Document()
    doc.add_paragraph("Requests", style="Heading 1")
    for text in ("first", "second", "third"):
        doc.add_paragraph(text, style="List Number")
    doc.add_paragraph("Facts", style="Heading 1")
    doc.add_paragraph("Body text.")
    src = tmp_path / "lists.docx"
    doc.save(str(src))
    out = json.loads(
        reshape_docx.call(
            {
                "source_path": str(src),
                "filename": "lists_out.docx",
                "operations": [
                    {"op": "insert_paragraph", "after": 5, "text": "new one", "style": "List Number"},
                    {"op": "insert_paragraph", "after": 6, "text": "new two", "style": "List Number"},
                    {"op": "insert_paragraph", "after": 3, "text": "fourth", "style": "List Number"},
                ],
            }
        )
    )
    assert [op["numbering"] for op in out["applied"]] == [
        "starts a new list at 1", "continues the list above", "continues the list above",
    ]
    labels = {b.accepted_text(): b.label() for b in read_docx(Path(out["path"])).body_paragraphs()}
    assert labels["fourth"] == "4."
    assert labels["new one"] == "1." and labels["new two"] == "2."
