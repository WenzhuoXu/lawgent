"""read_document and attachment intake for every format users upload.

Review items D-1 (legacy formats read as binary garbage: 64 reads, 4.5M
characters, in all 9 chats that uploaded a .doc), D-9 (17 HEIC photos never
reached the model) and proposal 6 (OFD).
"""

from __future__ import annotations

import io
import zipfile
from pathlib import Path

import pytest

from docx_fixtures import PRC_CONTRACT_NUMBERING, build_docx, p
from legal_helper import attachments as att
from legal_helper.tools import read_document


def _contract(tmp_path: Path) -> Path:
    body = p("总则", num=(1, 0)) + p("本合同自双方签字之日起生效。") + p("租金", num=(1, 0)) + p("每月人民币壹万元。")
    return build_docx(tmp_path / "合同.docx", body, numbering=PRC_CONTRACT_NUMBERING)


def test_binary_file_is_an_error_not_garbage(tmp_path):
    blob = tmp_path / "scan.bin"
    blob.write_bytes(bytes(range(256)) * 64)
    out = read_document.call({"path": str(blob)})
    assert out.startswith("ERROR: unsupported binary format .bin")
    assert "supported:" in out and ".docx" in out


def test_gb18030_text_reads_as_chinese(tmp_path):
    note = tmp_path / "说明.txt"
    note.write_bytes("第一条 甲方应于每月五日前支付租金。".encode("gb18030"))
    assert read_document.call({"path": str(note)}) == "第一条 甲方应于每月五日前支付租金。"


def test_text_without_an_extension_is_still_read(tmp_path):
    note = tmp_path / "README"
    note.write_text("plain notes", encoding="utf-8")
    assert read_document.call({"path": str(note)}) == "plain notes"


def test_legacy_word_formats(tmp_path):
    from legal_helper.documents.soffice import convert, soffice_binary

    if not soffice_binary():
        pytest.skip("LibreOffice not installed")
    source = _contract(tmp_path)
    for target in ("doc", "rtf", "odt"):
        legacy = convert(source, tmp_path / target, target)
        text = read_document.call({"path": str(legacy)})
        assert "本合同自双方签字之日起生效" in text, (target, text[:200])
        # LibreOffice's own RTF export writes the 第/条 of "第%1条" as "?", so
        # only the counted number survives in that fixture.
        assert ("一? 总则" if target == "rtf" else "第一条 总则") in text, (target, text[:200])
    # Kingsoft .wps is the Word 97 binary layout under another name.
    wps = tmp_path / "合同.wps"
    wps.write_bytes((tmp_path / "doc" / "合同.doc").read_bytes())
    assert "本合同自双方签字之日起生效" in read_document.call({"path": str(wps)})


def test_legacy_spreadsheet(tmp_path):
    from openpyxl import Workbook

    from legal_helper.documents.soffice import convert, soffice_binary

    if not soffice_binary():
        pytest.skip("LibreOffice not installed")
    wb = Workbook()
    wb.active.append(["条款", "金额"])
    wb.active.append(["租金", 10000])
    xlsx = tmp_path / "book.xlsx"
    wb.save(xlsx)
    xls = convert(xlsx, tmp_path / "xls", "xls")
    text = read_document.call({"path": str(xls)})
    assert "租金" in text and "10000" in text


def test_zip_lists_members_and_reads_one(tmp_path):
    archive = tmp_path / "材料.zip"
    with zipfile.ZipFile(archive, "w") as zf:
        zf.write(_contract(tmp_path), "附件/合同.docx")
        zf.writestr("notes.txt", "see clause 2")
        zf.writestr("XXXX.txt", "中文说明".encode("utf-8"))
    # A member named the way Chinese Windows zips name them: GBK bytes with no
    # UTF-8 flag (zipfile would flag a non-ASCII name, so patch the bytes).
    raw = archive.read_bytes().replace(b"XXXX.txt", "说明.txt".encode("gbk"))
    archive.write_bytes(raw)
    listing = read_document.call({"path": str(archive)})
    assert "附件/合同.docx" in listing and "notes.txt" in listing and "说明.txt" in listing
    member = read_document.call({"path": str(archive), "member": "附件/合同.docx"})
    assert "第一条 总则" in member
    assert read_document.call({"path": f"{archive}::notes.txt"}) == "see clause 2"
    assert read_document.call({"path": f"{archive}::说明.txt"}) == "中文说明"
    assert read_document.call({"path": str(archive), "member": "../../etc/passwd"}).startswith("ERROR")


def _ofd(path: Path) -> Path:
    ns = 'xmlns:ofd="http://www.ofdspec.org/2016"'
    with zipfile.ZipFile(path, "w") as zf:
        zf.writestr(
            "OFD.xml",
            f'<ofd:OFD {ns} Version="1.0" DocType="OFD"><ofd:DocBody><ofd:DocInfo><ofd:Title>民事判决书</ofd:Title>'
            "</ofd:DocInfo><ofd:DocRoot>Doc_0/Document.xml</ofd:DocRoot></ofd:DocBody></ofd:OFD>",
        )
        zf.writestr(
            "Doc_0/Document.xml",
            f"<ofd:Document {ns}><ofd:CommonData><ofd:PageArea><ofd:PhysicalBox>0 0 210 297</ofd:PhysicalBox>"
            '</ofd:PageArea><ofd:TemplatePage ID="9" BaseLoc="Tpls/Tpl_0/Content.xml"/></ofd:CommonData>'
            '<ofd:Pages><ofd:Page ID="1" BaseLoc="Pages/Page_0/Content.xml"/>'
            '<ofd:Page ID="2" BaseLoc="Pages/Page_1/Content.xml"/></ofd:Pages></ofd:Document>',
        )
        zf.writestr(
            "Doc_0/Tpls/Tpl_0/Content.xml",
            f'<ofd:Page {ns}><ofd:Content><ofd:Layer ID="1"><ofd:TextObject ID="2" Boundary="10 10 50 6" Size="4">'
            '<ofd:TextCode X="0" Y="4">北京市朝阳区人民法院</ofd:TextCode></ofd:TextObject></ofd:Layer></ofd:Content></ofd:Page>',
        )
        zf.writestr(
            "Doc_0/Pages/Page_0/Content.xml",
            f'<ofd:Page {ns}><ofd:Template TemplateID="9" ZOrder="Background"/><ofd:Content><ofd:Layer ID="3">'
            # Written out of reading order: the second line first, and one
            # line split into two objects.
            '<ofd:TextObject ID="5" Boundary="10 40 80 6" Size="4"><ofd:TextCode X="0" Y="4">原告诉称，</ofd:TextCode></ofd:TextObject>'
            '<ofd:TextObject ID="6" Boundary="40 40 80 6" Size="4"><ofd:TextCode X="0" Y="4">被告应支付租金。</ofd:TextCode></ofd:TextObject>'
            '<ofd:TextObject ID="4" Boundary="10 25 80 6" Size="4"><ofd:TextCode X="0" Y="4">（2026）京0105民初12345号</ofd:TextCode></ofd:TextObject>'
            "</ofd:Layer></ofd:Content></ofd:Page>",
        )
        zf.writestr(
            "Doc_0/Pages/Page_1/Content.xml",
            f'<ofd:Page {ns}><ofd:Content><ofd:Layer ID="7"><ofd:TextObject ID="8" Boundary="10 20 80 6" Size="4">'
            '<ofd:TextCode X="0" Y="4">本判决为终审判决。</ofd:TextCode></ofd:TextObject></ofd:Layer></ofd:Content></ofd:Page>',
        )
    return path


def test_ofd_reads_in_reading_order_with_page_markers(tmp_path):
    text = read_document.call({"path": str(_ofd(tmp_path / "判决书.ofd"))})
    assert text.startswith("[OFD document: 民事判决书]")
    page1 = text[text.index("[page 1]"):text.index("[page 2]")]
    assert page1.index("北京市朝阳区人民法院") < page1.index("（2026）京0105民初12345号") < page1.index("原告诉称，被告应支付租金。")
    assert "[page 2]\n本判决为终审判决。" in text
    only_two = read_document.call({"path": str(tmp_path / "判决书.ofd"), "pages": "2"})
    assert "本判决为终审判决" in only_two and "原告诉称" not in only_two


def test_ofd_attachment_is_inlined_as_text(tmp_path):
    info = att.describe(_ofd(tmp_path / "invoice.ofd"))
    assert info.category == "text_inline"
    assert "原告诉称，被告应支付租金。" in att.extract_inline_text(info)


def _heic(tmp_path: Path) -> Path:
    pillow_heif = pytest.importorskip("pillow_heif", reason="writing a HEIC fixture needs pillow-heif's encoder")
    from PIL import Image

    pillow_heif.register_heif_opener()
    path = tmp_path / "IMG_0001.HEIC"
    Image.new("RGB", (640, 480), (200, 30, 30)).save(path, format="HEIF")
    return path


def test_heic_attachment_reaches_the_model_as_jpeg(tmp_path):
    info = att.describe(_heic(tmp_path))
    assert info.category == "native_image"
    block = att.build_anthropic_block(info)
    assert block["type"] == "image" and block["source"]["media_type"] == "image/jpeg"
    openai_block = att.build_openai_block(info)
    assert openai_block["image_url"].startswith("data:image/jpeg;base64,")
    from PIL import Image
    import base64

    decoded = Image.open(io.BytesIO(base64.b64decode(block["source"]["data"])))
    assert decoded.format == "JPEG" and decoded.size == (640, 480)


def test_view_image_opens_heic(tmp_path):
    import json

    from legal_helper.tools import view_image

    payload = json.loads(view_image.call({"path": str(_heic(tmp_path))}))
    assert payload["width"] == 640 and payload["height"] == 480


def test_legacy_attachment_is_extracted_not_stubbed(tmp_path):
    from legal_helper.documents.soffice import convert, soffice_binary

    if not soffice_binary():
        pytest.skip("LibreOffice not installed")
    doc = convert(_contract(tmp_path), tmp_path / "doc", "doc")
    info = att.describe(doc)
    assert info.category == "text_inline"
    blocks = att.build_user_content("review", [doc], "openai")
    assert "第一条 总则" in blocks[0]["text"]


def test_pages_on_a_format_without_pages_says_so(tmp_path):
    text = read_document.call({"path": str(_contract(tmp_path)), "pages": "2"})
    assert text.startswith("[pages='2' applies to PDF, OFD and presentations")
    assert "第一条 总则" in text


def test_slides_are_selectable_by_pages(tmp_path):
    from pptx import Presentation

    prs = Presentation()
    for i in range(1, 4):
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = f"Title {i}"
        slide.placeholders[1].text = f"Body {i}"
    deck = tmp_path / "deck.pptx"
    prs.save(deck)
    text = read_document.call({"path": str(deck), "pages": "2"})
    assert "## Slide 2" in text and "Body 2" in text
    assert "Body 1" not in text and "Body 3" not in text


def test_zip_attachment_is_inlined_as_its_listing(tmp_path):
    archive = tmp_path / "bundle.zip"
    with zipfile.ZipFile(archive, "w") as zf:
        zf.writestr("a.txt", "alpha")
    info = att.describe(archive)
    assert info.category == "text_inline"
    text = att.extract_inline_text(info)
    assert "- a.txt (5 bytes)" in text and "member=" in text
