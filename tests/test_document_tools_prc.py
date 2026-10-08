"""The document tools and the export path, at the surface the model and UI use.

* ``write_docx``/``write_pdf`` take ``lang``, ``layout``, ``page_size``,
  ``disclaimer`` and (Word) ``reference_docx``; the disclaimer goes into
  document properties in the document's language, and a filing can carry none
  — with a stated reason (D-12, proposal 2).
* The export path titles an answer by its own first heading in its own
  language: 18 of 19 exported DOCX and 5 of 5 PDFs used to open with an
  English "Legal Analysis" and an English disclaimer.
* A review render maps the GB fonts to open stand-ins and says so, and lints
  the file with its PDF (proposal 14 font policy).
"""

from __future__ import annotations

import json
import shutil
from pathlib import Path

import pytest
from docx import Document

from legal_helper.config import load_settings
from legal_helper.documents.writers._text import DISCLAIMERS, export_title, han_share


@pytest.fixture()
def outputs(tmp_path, monkeypatch):
    out = tmp_path / "out"
    monkeypatch.setenv("LEGAL_HELPER_OUTPUTS_DIR", str(out))
    from legal_helper import config as cfg

    cfg._CACHED = None
    return out


def _body_text(doc):
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.extend(c.text for c in row.cells)
    return "\n".join(parts)


ZH_SECTIONS = [{"heading": "一、诉讼请求", "body_markdown": "1. 判令被告支付货款；\n2. 判令被告承担诉讼费用。"}]


def test_write_docx_puts_a_chinese_notice_in_properties(outputs):
    from legal_helper.tools import write_docx

    path = write_docx.call({"filename": "起诉状.docx", "title": "民事起诉状", "sections": ZH_SECTIONS})
    doc = Document(path)
    assert doc.core_properties.comments == DISCLAIMERS["zh"]
    text = _body_text(doc)
    assert "not legal advice" not in text.lower() and "AI" not in text
    assert doc.sections[0].page_width.twips == 11906


def test_a_filing_can_carry_no_notice_but_must_say_why(outputs):
    from legal_helper.tools import write_docx

    refused = write_docx.call({
        "filename": "起诉状.docx", "title": "民事起诉状", "sections": ZH_SECTIONS, "disclaimer": "none",
    })
    assert refused.startswith("ERROR") and "disclaimer_reason" in refused
    path = write_docx.call({
        "filename": "起诉状.docx", "title": "民事起诉状", "sections": ZH_SECTIONS,
        "disclaimer": "none", "disclaimer_reason": "court filing to be signed by counsel",
    })
    doc = Document(path)
    assert doc.core_properties.comments == ""
    assert "AI" not in _body_text(doc)


def test_write_docx_exposes_layout_lang_and_page_size(outputs):
    from legal_helper.tools import write_docx

    props = write_docx.input_schema["properties"]
    assert props["layout"]["enum"] == ["memo", "gbt9704"]
    assert props["lang"]["enum"] == ["auto", "zh", "en"]
    assert props["page_size"]["enum"] == ["auto", "A4", "Letter"]
    assert props["disclaimer"]["enum"] == ["properties", "none"]
    assert "reference_docx" in props
    path = write_docx.call({
        "filename": "brief.docx", "title": "Brief", "sections": [{"heading": "Facts", "body_markdown": "Text."}],
        "page_size": "A4",
    })
    assert Document(path).sections[0].page_width.twips == 11906
    path = write_docx.call({
        "filename": "gb.docx", "title": "关于某事项的请示",
        "sections": [{"heading": "一、背景", "body_markdown": "现将有关情况报告如下。"}], "layout": "gbt9704",
    })
    assert Document(path).sections[0].top_margin.twips == 2098


def test_a_template_build_is_recorded_as_derived_from_the_template(outputs, tmp_path, monkeypatch):
    from legal_helper.documents import lineage
    from legal_helper.tools import write_docx

    state = tmp_path / "state"
    monkeypatch.setattr(lineage, "_ledger", lambda state_dir=None: state / "lineage.jsonl")
    template = tmp_path / "firm.docx"
    doc = Document()
    doc.sections[0].header.is_linked_to_previous = False
    doc.sections[0].header.paragraphs[0].text = "某某律师事务所"
    doc.add_paragraph("占位")
    doc.save(str(template))
    path = write_docx.call({
        "filename": "letter.docx", "title": "律师函", "sections": ZH_SECTIONS,
        "reference_docx": str(template), "disclaimer": "none", "disclaimer_reason": "律师函 on letterhead",
    })
    assert Document(path).sections[0].header.paragraphs[0].text == "某某律师事务所"
    assert lineage.derived_from(path) == [str(template.resolve())]
    missing = write_docx.call({
        "filename": "x.docx", "title": "t", "sections": ZH_SECTIONS, "reference_docx": str(tmp_path / "nope.docx"),
    })
    assert missing.startswith("ERROR")


def test_write_pdf_puts_the_notice_in_metadata(outputs):
    import pypdfium2 as pdfium

    from legal_helper.tools import write_pdf

    path = write_pdf.call({"filename": "memo.pdf", "title": "Memo", "body_markdown": "All conditions are met."})
    doc = pdfium.PdfDocument(path)
    try:
        assert doc.get_metadata_dict()["Subject"] == DISCLAIMERS["en"]
        assert "not legal advice" not in doc[0].get_textpage().get_text_range()
    finally:
        doc.close()


# ---------------------------------------------------------------------------
# Export path
# ---------------------------------------------------------------------------


def test_export_titles_a_chinese_answer_by_its_own_heading(tmp_path):
    from legal_helper.workflow import write_followup_docx

    settings = load_settings(refresh=True).model_copy(update={"outputs_dir": tmp_path / "out"})
    answer = "# 关于货款纠纷的法律意见\n\n## 一、结论\n\n被告应当支付货款。"
    artifact = write_followup_docx(settings, answer)
    doc = Document(artifact["path"])
    texts = [p.text for p in doc.paragraphs]
    assert texts.count("关于货款纠纷的法律意见") == 1  # the title, not printed twice
    assert "Legal Analysis" not in texts
    assert doc.core_properties.comments == DISCLAIMERS["zh"]
    assert doc.sections[0].page_width.twips == 11906


def test_export_without_a_heading_uses_a_title_in_the_answers_language(tmp_path):
    from legal_helper.workflow import write_followup_pdf

    settings = load_settings(refresh=True).model_copy(update={"outputs_dir": tmp_path / "out"})
    assert export_title("被告应当支付货款。") == "法律分析"
    assert export_title("The lessee must pay.") == "Legal Analysis"
    assert export_title("被告应当支付货款。", chat_title="货款纠纷咨询") == "货款纠纷咨询"
    assert export_title("被告应当支付货款。", chat_title="Payment dispute") == "法律分析"
    artifact = write_followup_pdf(settings, "The lessee must pay the rent on time.")
    assert Path(artifact["path"]).is_file()


def test_export_endpoint_derives_the_title(tmp_path, monkeypatch):
    from fastapi.testclient import TestClient

    import legal_helper.server as server_mod
    from legal_helper.chat_store import ChatStore

    settings = load_settings(refresh=True).model_copy(update={"outputs_dir": tmp_path / "out"})
    store = ChatStore(settings, tmp_path / "chat.sqlite3")
    monkeypatch.setattr(server_mod, "_store", lambda: store)
    monkeypatch.setattr(server_mod, "_settings_for_chat", lambda chat_settings, chat_id=None: settings)
    chat = store.create_chat(title="货款纠纷咨询")
    msg = store.add_message(chat.id, "assistant", "被告应当支付货款。")
    response = TestClient(server_mod.app).post(
        f"/api/chats/{chat.id}/messages/{msg.id}/export", json={"format": "docx"},
    )
    assert response.status_code == 200
    started = store.list_events(chat.id)[0]
    assert started.event_type == "local_tool_call_started"
    assert started.data["arguments"]["title"] == "货款纠纷咨询"
    doc = Document(response.json()["path"])
    assert doc.paragraphs[0].text == "货款纠纷咨询"
    assert doc.core_properties.comments == DISCLAIMERS["zh"]


def test_han_share_decides_the_language():
    assert han_share("法律意见书") == 1.0
    assert han_share("The lessee shall pay，") == 0.0  # a fullwidth comma is not Han
    assert han_share("Lease with 东航 as lessee and a long English text body here") < 0.5


# ---------------------------------------------------------------------------
# Review renders
# ---------------------------------------------------------------------------


@pytest.mark.skipif(not shutil.which("soffice"), reason="LibreOffice not installed")
def test_a_review_render_names_its_font_stand_ins_and_lints_with_the_pdf(outputs, tmp_path):
    from legal_helper.documents.review_fonts import installed_families
    from legal_helper.tools import render_docx_pages, write_docx
    from legal_helper.tools.multimodal import split_inline_images

    path = write_docx.call({
        "filename": "gb.docx", "title": "关于某事项的请示",
        "sections": [{"heading": "一、基本情况", "body_markdown": "| 事项 | 依据 |\n|---|---|\n| 许可 | 第九十二条 |"}],
        "layout": "gbt9704",
    })
    raw = render_docx_pages.call({"path": path, "filename_prefix": "qa"})
    assert not raw.startswith("ERROR"), raw
    text, images = split_inline_images(raw)
    body = json.loads(text)
    assert images
    assert body["lint"]["conformance"] == "pass"
    assert body["lint"]["metrics"].get("pages") == 1  # the PDF reached the linter
    if "仿宋" not in installed_families():
        assert body["font_stand_ins"]["fonts"]["仿宋"] == "Noto Serif CJK SC"
        assert "do not change fonts" in body["font_stand_ins"]["note"]


@pytest.mark.skipif(not shutil.which("soffice"), reason="LibreOffice not installed")
def test_review_renders_give_each_gb_font_role_its_own_face(tmp_path):
    """Without the stand-ins every role came out in one fallback face."""
    import pypdfium2 as pdfium
    import pypdfium2.raw as pdfium_c

    from legal_helper.documents.review_fonts import review_env
    from legal_helper.documents.soffice import convert

    doc = Document()
    for font in ("黑体", "仿宋", "方正小标宋简体"):
        run = doc.add_paragraph().add_run(f"{font}：中华人民共和国民法典")
        run._r.get_or_add_rPr().get_or_add_rFonts().set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia", font)
    source = tmp_path / "fonts.docx"
    doc.save(str(source))
    pdf = convert(source, tmp_path / "pdf", "pdf", env=review_env())
    names = set()
    document = pdfium.PdfDocument(str(pdf))
    try:
        textpage = document[0].get_textpage()
        import ctypes

        for index in range(textpage.count_chars()):
            buffer = ctypes.create_string_buffer(256)
            flags = ctypes.c_int()
            pdfium_c.FPDFText_GetFontInfo(textpage.raw, index, buffer, 256, ctypes.byref(flags))
            names.add(buffer.value.decode("utf-8", "replace").split("+")[-1])
    finally:
        document.close()
    cjk = {n for n in names if "CJK" in n}
    assert len(cjk) >= 3, names  # a sans, a serif and a bold serif


def test_a_floating_table_the_preview_drops_is_reported_not_repaired(tmp_path):
    """LibreOffice 7.3 does not draw a table positioned with w:tblpPr (a 公文's
    版记); the render must say so, or the author re-runs the generator to fix
    a document that was never broken."""
    from docx import Document
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls
    from reportlab.pdfgen import canvas

    from legal_helper.tools.documents import _unrendered_floating_tables

    doc = Document()
    doc.add_paragraph("正文")
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "抄送：市安委办。某市交通运输局办公室2026年10月9日印发"
    table._tbl.tblPr.append(parse_xml(
        f'<w:tblpPr {nsdecls("w")} w:vertAnchor="margin" w:horzAnchor="margin" w:tblpYSpec="bottom"/>'
    ))
    src = tmp_path / "gw.docx"
    doc.save(str(src))

    pdf = tmp_path / "gw.pdf"
    c = canvas.Canvas(str(pdf))
    c.drawString(72, 720, "body text only")
    c.save()
    note = _unrendered_floating_tables(src, pdf)["floating_tables_not_rendered"]
    assert note["tables"][0]["table_index"] == 0
    assert "not a defect" in note["note"]

    # Drawn after all: nothing to report.
    c = canvas.Canvas(str(pdf))
    c.drawString(72, 720, "body")
    c.save()
    plain = Document()
    plain.add_table(rows=1, cols=1).cell(0, 0).text = "in flow"
    plain.save(str(src))
    assert _unrendered_floating_tables(src, pdf) == {}
