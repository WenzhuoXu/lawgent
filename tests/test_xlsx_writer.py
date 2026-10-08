"""The workbook writer: formats, dates, widths and stored formula values (D-14).

Measured before: 0 of 351 numeric cells formatted, 45 dates stored as text,
48 of 48 formulas with no cached value, 288 of 578 columns too narrow for their
Chinese text — and three docstrings promising a recalculation that never ran.
"""

from __future__ import annotations

import datetime as dt
import json
import shutil

import pytest
from openpyxl import load_workbook

from legal_helper.documents.writers import write_xlsx_workbook
from legal_helper.documents.writers._xlsx_recalc import recalc_xlsx_workbook


def _write(tmp_path, rows, **sheet):
    path = tmp_path / "book.xlsx"
    write_xlsx_workbook(path, [{"name": "台账", "rows": rows, **sheet}])
    return path


def test_iso_date_strings_become_real_dates(tmp_path):
    path = _write(tmp_path, [
        ["事项", "日期"],
        ["立案", "2026-10-07"],
        ["开庭", "2026-11-03 09:30"],
        ["文号", "'2026-10-08"],      # apostrophe: keep it text
        ["无效", "2026-02-30"],       # not a date: stays text
    ])
    ws = load_workbook(path).active
    assert ws["B2"].value == dt.datetime(2026, 10, 7)
    assert ws["B2"].number_format == "yyyy-mm-dd"
    assert ws["B3"].value == dt.datetime(2026, 11, 3, 9, 30)
    assert ws["B4"].value == "2026-10-08" and ws["B4"].data_type == "s"
    assert ws["B5"].value == "2026-02-30"


def test_number_formats_apply_to_columns_and_ranges(tmp_path):
    path = _write(
        tmp_path,
        [["项目", "金额", "比例"], ["甲", 1234567.5, 0.125], ["乙", 89, 0.5]],
        number_formats=[{"range": "B", "format": "#,##0.00"}, {"range": "C2:C3", "format": "0.0%"}],
    )
    ws = load_workbook(path).active
    assert ws["B2"].number_format == "#,##0.00" and ws["B3"].number_format == "#,##0.00"
    assert ws["C2"].number_format == "0.0%" and ws["C3"].number_format == "0.0%"


def test_column_widths_count_a_han_character_double(tmp_path):
    path = _write(tmp_path, [["说明说明说明说明说明说明", "abcdefghijkl"], ["x", "y"]])
    ws = load_workbook(path).active
    assert ws.column_dimensions["A"].width > ws.column_dimensions["B"].width * 1.5


def test_text_longer_than_the_widest_column_wraps(tmp_path):
    long_text = "根据《民用航空法》第九十二条之规定办理经营许可的申请手续" * 3
    path = _write(tmp_path, [["备注"], [long_text]])
    ws = load_workbook(path).active
    assert ws.column_dimensions["A"].width == 60
    assert ws["A2"].alignment.wrap_text is True


def test_formula_values_are_stored_in_the_file(tmp_path):
    path = _write(tmp_path, [["项目", "金额"], ["甲", 10], ["乙", 20], ["合计", "=SUM(B2:B3)"]])
    status = recalc_xlsx_workbook(path)
    assert status["recalculated"] is True
    assert status["cached"] == 1
    assert load_workbook(path, data_only=True).active["B4"].value == 30
    # The formula itself is still there for Excel.
    assert load_workbook(path).active["B4"].value == "=SUM(B2:B3)"


def test_formula_errors_are_cached_and_reported(tmp_path):
    path = _write(tmp_path, [["a", "b"], [1, "=A2/0"]])
    status = recalc_xlsx_workbook(path)
    assert status["errors"] == [{"sheet": "台账", "cell": "B2", "error": "#DIV/0!"}]
    assert load_workbook(path, data_only=True).active["B2"].value == "#DIV/0!"


def test_recalculation_preserves_what_openpyxl_wrote(tmp_path):
    """Only cached values change: styles and widths survive the LibreOffice pass."""
    path = _write(tmp_path, [["项目", "金额"], ["甲", 10], ["合计", "=SUM(B2:B2)"]],
                  number_formats=[{"range": "B", "format": "#,##0.00"}])
    recalc_xlsx_workbook(path)
    ws = load_workbook(path).active
    assert ws["A1"].font.bold is True
    assert ws["B3"].number_format == "#,##0.00"
    assert ws.freeze_panes == "A2"


def test_the_formulas_engine_stands_in_without_soffice(tmp_path):
    path = _write(tmp_path, [["a", "b"], [2, "=A2*21"]])
    status = recalc_xlsx_workbook(path, use_libreoffice=False)
    assert status["engine"] == "formulas"
    assert load_workbook(path, data_only=True).active["B2"].value == 42


@pytest.mark.skipif(not shutil.which("soffice"), reason="LibreOffice not installed")
def test_a_function_libreoffice_lacks_is_evaluated_by_the_second_engine(tmp_path):
    # LibreOffice 7.3 answers XLOOKUP with #NAME?; that is the engine's gap,
    # not the workbook's error.
    path = _write(tmp_path, [["k", "v"], [1, "甲"], [2, "乙"], ["查", "=XLOOKUP(2,A2:A3,B2:B3)"]])
    status = recalc_xlsx_workbook(path)
    assert status["errors"] == []
    assert load_workbook(path, data_only=True).active["B4"].value == "乙"


def test_write_xlsx_tool_reports_errors_and_stores_values(tmp_path, monkeypatch):
    monkeypatch.setenv("LEGAL_HELPER_OUTPUTS_DIR", str(tmp_path / "out"))
    from legal_helper import config as cfg
    from legal_helper.tools import write_xlsx

    cfg._CACHED = None
    clean = write_xlsx.call({
        "filename": "ok.xlsx",
        "sheets": [{"name": "S", "rows": [["a", "b"], [1, "=A2+1"]],
                    "number_formats": [{"range": "B", "format": "0.00"}]}],
    })
    assert clean.endswith(".xlsx")
    assert load_workbook(clean, data_only=True).active["B2"].value == 2
    broken = json.loads(write_xlsx.call({
        "filename": "bad.xlsx",
        "sheets": [{"name": "S", "rows": [["a", "b"], [1, "=A2/0"]]}],
    }))
    assert broken["formula_errors"][0]["error"] == "#DIV/0!"
    assert broken["path"].endswith("bad.xlsx")
