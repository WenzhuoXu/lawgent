"""Provable-breakage checks for Word, PDF and Excel output (proposal 2).

The 2026-10-07 review built a synthetic contract with the objective defects of
§2.2 and ``lint_docx`` found none of them, and an XLSX render reported
"conformance: pass" with no check having run. Each check here is one the
review measured on real output; each passes only when nothing is provably
broken, and none grades style.

Fixtures are built with python-docx / openpyxl / reportlab directly, the way
the old writer and the edit tools produce them, so every test states its own
defect.
"""

from __future__ import annotations

import shutil

import pytest
from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from openpyxl import Workbook

from legal_helper.documents.quality import (
    NO_CHECKS,
    QualityReport,
    lint_artifact,
    lint_docx,
    lint_pdf,
    lint_xlsx,
)
from legal_helper.documents.writers import write_docx_document


def _checks(findings):
    return [f.check for f in findings]


# ---------------------------------------------------------------------------
# The synthetic contract
# ---------------------------------------------------------------------------


def _synthetic_contract(path, *, long_table=True):
    """A contract carrying the eight objective defects of review §2.2."""
    doc = Document()
    doc.add_heading("买卖合同", level=0)
    doc.add_heading("第一部分 诉讼请求", level=1)
    for text in ("支付货款", "支付违约金", "承担诉讼费用"):
        doc.add_paragraph(text, style="List Number")  # the old writer's single list
    doc.add_heading("第二部分 事实", level=1)
    for text in ("签订合同", "逾期付款"):               # (1) numbered 4 and 5, not 1 and 2
        doc.add_paragraph(text, style="List Number")
    doc.add_paragraph("参见[《民法典》第五百七十七条](https://flk.npc.gov.cn/)。")  # (2)
    doc.add_paragraph("#### 违约责任")                                             # (3)
    doc.add_paragraph("**结论：**被告应承担违约责任。")                              # (4)
    doc.add_paragraph("| 事项 | 依据 |")
    doc.add_paragraph("|---|---|")                                                 # (5)
    doc.add_paragraph("依据如下[^1]。")                                            # (6)
    doc.add_paragraph("~~原表述~~ 第一行<br>第二行")                                 # (7)
    if long_table:                                                                  # (8)
        table = doc.add_table(rows=1, cols=2)
        table.style = "Table Grid"
        table.rows[0].cells[0].text = "条款编号表头"
        table.rows[0].cells[1].text = "条款内容表头"
        for i in range(70):
            row = table.add_row()
            row.cells[0].text = f"第{i + 1}行编号"
            row.cells[1].text = f"内容{i + 1}号条款说明文字"
    doc.save(str(path))
    return path


def test_the_synthetic_contract_now_gets_its_findings(tmp_path):
    findings, _ = lint_docx(_synthetic_contract(tmp_path / "contract.docx", long_table=False))
    checks = _checks(findings)
    continuing = [f for f in findings if f.check == "list_numbering_continues"]
    assert len(continuing) == 1
    assert continuing[0].severity == "critical"
    assert "numbered 4" in continuing[0].detail
    residue = " ".join(f.detail for f in findings if f.check == "markdown_residue")
    for kind in ("[text](url) link", "#### heading marks", "**bold** markers",
                 "|---| table rule", "[^1] footnote marker", "~~strike~~ markers", "<br> tag"):
        assert kind in residue, kind
    assert checks.count("markdown_residue") >= 6
    assert all(f.where.startswith("paragraph ") for f in findings if f.check == "markdown_residue")


@pytest.mark.skipif(not shutil.which("soffice"), reason="LibreOffice not installed")
def test_a_table_that_crosses_a_page_without_a_header_row_is_found(tmp_path):
    from legal_helper.documents.soffice import convert

    path = _synthetic_contract(tmp_path / "contract.docx")
    pdf = convert(path, tmp_path / "pdf", "pdf")
    findings, metrics = lint_docx(path, pdf=pdf)
    crossing = [f for f in findings if f.check == "table_header_not_repeated"]
    assert metrics["pages"] >= 2
    assert len(crossing) == 1 and crossing[0].where == "table 1"
    assert "page 1 to page 2" in crossing[0].detail or "to page" in crossing[0].detail

    # The writer's own tables repeat their header row, so the same table is clean.
    body = "| 条款编号表头 | 条款内容表头 |\n|---|---|\n" + "\n".join(
        f"| 第{i + 1}行编号 | 内容{i + 1}号条款说明文字 |" for i in range(70)
    )
    ours = tmp_path / "ours.docx"
    write_docx_document(output_path=ours, title="合同", sections=[("", body)])
    pdf = convert(ours, tmp_path / "pdf2", "pdf")
    findings, metrics = lint_docx(ours, pdf=pdf)
    assert metrics["pages"] >= 2
    assert "table_header_not_repeated" not in _checks(findings)


def test_the_writers_own_output_is_clean(tmp_path):
    body = (
        "## 一、诉讼请求\n\n1. 支付货款；\n2. 支付违约金。\n\n## 二、事实\n\n1. 签订合同；\n2. 逾期付款。\n\n"
        "参见[《民法典》](https://flk.npc.gov.cn/)[^1]。**结论：**成立。\n\n"
        "| 事项 | 依据 |\n|---|---|\n| 货款 | 第五百七十七条 |\n\n[^1]: 注释。"
    )
    out = tmp_path / "ours.docx"
    write_docx_document(output_path=out, title="民事起诉状", sections=[("", body)])
    findings, _ = lint_docx(out)
    assert findings == [], [(f.check, f.where, f.detail) for f in findings]


def test_clause_numbering_that_runs_across_chapters_is_not_a_defect(tmp_path):
    """第1条 … 第6条 across 第一章/第二章 is how contracts are numbered."""
    doc = Document()
    numbering = doc.part.numbering_part.element
    numbering.insert(0, parse_xml(
        f'<w:abstractNum {nsdecls("w")} w:abstractNumId="77"><w:lvl w:ilvl="0">'
        '<w:start w:val="1"/><w:numFmt w:val="chineseCounting"/><w:lvlText w:val="第%1条"/>'
        '<w:lvlJc w:val="left"/></w:lvl></w:abstractNum>'
    ))
    num_id = numbering.add_num(77).numId
    for chapter in ("第一章 总则", "第二章 价款"):
        doc.add_heading(chapter, level=1)
        for _ in range(3):
            para = doc.add_paragraph("条款内容")
            num_pr = para._p.get_or_add_pPr().get_or_add_numPr()
            num_pr.get_or_add_ilvl().val = 0
            num_pr.get_or_add_numId().val = num_id
    path = tmp_path / "clauses.docx"
    doc.save(str(path))
    assert "list_numbering_continues" not in _checks(lint_docx(path)[0])


def test_gbt9704_headings_are_meant_to_be_indented(tmp_path):
    out = tmp_path / "gb.docx"
    write_docx_document(
        output_path=out, title="关于某事项的请示", layout="gbt9704",
        sections=[("", "一、基本情况\n\n（一）项目背景\n\n1.项目概况\n\n2.另一项\n\n正文。")],
    )
    checks = _checks(lint_docx(out)[0])
    assert "heading_indent" not in checks and "empty_section" not in checks


def test_markdown_inside_a_code_listing_is_left_alone(tmp_path):
    out = tmp_path / "code.docx"
    write_docx_document(output_path=out, title="Guide", sections=[("", "```\n**not bold**\n#### not a heading\n```")])
    assert "markdown_residue" not in _checks(lint_docx(out)[0])


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------


def test_markdown_residue_in_a_pdf_is_found_by_page(tmp_path):
    from reportlab.pdfgen import canvas

    path = tmp_path / "residue.pdf"
    c = canvas.Canvas(str(path))
    c.drawString(72, 720, "Clean first page.")
    c.showPage()
    c.drawString(72, 720, "See [the Act](https://example.com) and **bold** text.")
    c.showPage()
    c.save()
    findings, metrics = lint_pdf(path)
    assert metrics["pages"] == 2
    assert [f.where for f in findings] == ["page 2"]
    assert "[text](url) link" in findings[0].detail and "**bold** markers" in findings[0].detail


# ---------------------------------------------------------------------------
# XLSX
# ---------------------------------------------------------------------------


def _workbook(path, *, cached: bool):
    wb = Workbook()
    ws = wb.active
    ws.title = "台账"
    ws.append(["a", "b", "c"])
    ws.append([1, "=A2/0", "=A2+1"])
    wb.save(path)
    if cached:
        from legal_helper.documents.writers._xlsx_recalc import recalc_xlsx_workbook

        recalc_xlsx_workbook(path)
    return path


def test_formulas_without_cached_values_are_found(tmp_path):
    # openpyxl writes a formula with an empty cached value; every reader but
    # Excel sees a blank.
    findings, metrics = lint_xlsx(_workbook(tmp_path / "raw.xlsx", cached=False))
    uncached = [f for f in findings if f.check == "formula_without_cached_value"]
    assert {f.where for f in uncached} == {"台账!B2", "台账!C2"}
    assert metrics["uncached_formulas"] == 2


def test_error_values_are_critical_findings(tmp_path):
    findings, metrics = lint_xlsx(_workbook(tmp_path / "computed.xlsx", cached=True))
    errors = [f for f in findings if f.check == "formula_error_value"]
    assert [(f.where, f.severity) for f in errors] == [("台账!B2", "critical")]
    assert "#DIV/0!" in errors[0].detail
    assert "formula_without_cached_value" not in _checks(findings)
    assert metrics["error_cells"] == 1


def test_a_clean_workbook_reports_pass_because_checks_ran(tmp_path):
    wb = Workbook()
    wb.active.append(["a", 1])
    path = tmp_path / "clean.xlsx"
    wb.save(path)
    report = lint_artifact(path)
    assert report.findings == [] and report.checked is True
    assert report.to_payload()["conformance"] == "pass"


# ---------------------------------------------------------------------------
# "no checks for this type"
# ---------------------------------------------------------------------------


def test_a_type_with_no_checks_never_reports_pass(tmp_path):
    path = tmp_path / "data.csv"
    path.write_text("a,b\n1,2\n", encoding="utf-8")
    report = lint_artifact(path)
    assert report.checked is False
    assert report.to_payload()["conformance"] == NO_CHECKS
    assert report.summary() == f"conformance: {NO_CHECKS}"


def test_an_unreadable_file_is_not_checked_rather_than_passed(tmp_path):
    path = tmp_path / "broken.xlsx"
    path.write_bytes(b"not a workbook")
    payload = lint_artifact(path).to_payload()
    assert payload["conformance"].startswith("not checked")
    assert "file_check_error" in payload["metrics"]


def test_a_hand_built_report_keeps_the_old_meaning():
    assert QualityReport().to_payload()["conformance"] == "pass"
