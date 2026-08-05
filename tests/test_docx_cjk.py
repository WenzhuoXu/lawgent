"""CJK-native typography contract for the .docx writer."""

from __future__ import annotations

from docx import Document
from docx.oxml.ns import qn

from legal_helper.documents.writers import write_docx_document


CJK_BODY = """
本所受贵司委托，就拟议交易出具本法律意见书。

## 一、结论

- 拟议交易不存在实质性法律障碍
- 需补办备案手续

| 条款 | 评级 | 说明 |
|---|---|---|
| 保证 | 绿 | 符合要求 |
""".strip()


def _east_asia(style):
    rpr = style.element.rPr
    if rpr is None:
        return None
    rfonts = rpr.find(qn("w:rFonts"))
    return rfonts.get(qn("w:eastAsia")) if rfonts is not None else None


def test_cjk_auto_detected_typography(tmp_path):
    out = tmp_path / "opinion.docx"
    write_docx_document(
        output_path=out,
        title="法律意见书",
        sections=[("结论", CJK_BODY)],
        subtitle="仅供内部讨论",
        disclaimer="本文件不构成正式法律意见。",
    )
    doc = Document(str(out))

    normal = doc.styles["Normal"]
    assert _east_asia(normal) == "宋体"
    rpr = normal.element.rPr
    assert rpr.find(qn("w:lang")).get(qn("w:eastAsia")) == "zh-CN"

    ppr = normal.element.pPr
    # CJK punctuation/line-break rules: kinsoku, hanging punctuation, and
    # CJK-Latin auto-spacing must be on.
    for tag in ("w:kinsoku", "w:overflowPunct", "w:autoSpaceDE", "w:autoSpaceDN"):
        assert ppr.find(qn(tag)) is not None, f"missing {tag}"
    ind = ppr.find(qn("w:ind"))
    assert ind.get(qn("w:firstLineChars")) == "200"  # 首行缩进 2 字符

    for level in (1, 2, 3):
        heading = doc.styles[f"Heading {level}"]
        assert _east_asia(heading) == "黑体"
        assert str(heading.font.color.rgb) == "000000", "headings must be black, not accent blue"
    assert _east_asia(doc.styles["Title"]) == "宋体"

    # Title must not inherit the body first-line indent.
    title_ind = doc.styles["Title"].element.pPr.find(qn("w:ind"))
    assert title_ind.get(qn("w:firstLineChars")) == "0"

    assert doc.tables[0].style.name == "Table Grid"


def test_cjk_paragraph_property_order_is_schema_valid(tmp_path):
    # kinsoku/overflowPunct/autoSpace* must precede w:spacing and w:ind in
    # pPr, or Word flags the file for repair.
    out = tmp_path / "opinion.docx"
    write_docx_document(output_path=out, title="意见", sections=[("", "正文内容。")])
    doc = Document(str(out))
    ppr = doc.styles["Normal"].element.pPr
    tags = [child.tag for child in ppr]
    kinsoku_idx = tags.index(qn("w:kinsoku"))
    spacing_idx = tags.index(qn("w:spacing"))
    ind_idx = tags.index(qn("w:ind"))
    assert kinsoku_idx < spacing_idx < ind_idx or kinsoku_idx < ind_idx


def test_gbt9704_layout(tmp_path):
    out = tmp_path / "memo.docx"
    write_docx_document(
        output_path=out,
        title="关于某事项的请示",
        sections=[("一、背景", "现将有关情况报告如下。")],
        layout="gbt9704",
    )
    doc = Document(str(out))
    section = doc.sections[0]
    assert round(section.top_margin.cm, 1) == 3.7
    assert round(section.bottom_margin.cm, 1) == 3.5
    assert round(section.left_margin.cm, 1) == 2.8
    assert round(section.right_margin.cm, 1) == 2.6

    normal = doc.styles["Normal"]
    assert _east_asia(normal) == "仿宋"
    assert normal.font.size.pt == 16  # 三号
    assert normal.paragraph_format.line_spacing.pt == 28

    assert _east_asia(doc.styles["Heading 1"]) == "黑体"
    assert _east_asia(doc.styles["Heading 2"]) == "楷体"
    assert doc.styles["Title"].font.size.pt == 22  # 二号


def test_explicit_lang_overrides_detection(tmp_path):
    out = tmp_path / "forced.docx"
    write_docx_document(
        output_path=out,
        title="Purely Latin title",
        sections=[("Summary", "English body only.")],
        lang="zh-CN",
    )
    doc = Document(str(out))
    assert _east_asia(doc.styles["Normal"]) == "宋体"


def test_latin_documents_unchanged(tmp_path):
    out = tmp_path / "latin.docx"
    write_docx_document(
        output_path=out,
        title="Lease Review",
        sections=[("Summary", "All conditions met.\n\n| A | B |\n|---|---|\n| 1 | 2 |")],
        subtitle="Sample",
    )
    doc = Document(str(out))
    assert _east_asia(doc.styles["Normal"]) is None
    ppr = doc.styles["Normal"].element.pPr
    assert ppr.find(qn("w:ind")) is None, "no first-line indent for Latin docs"
    assert ppr.find(qn("w:kinsoku")) is None
    assert doc.tables[0].style.name == "Light Grid Accent 1"
