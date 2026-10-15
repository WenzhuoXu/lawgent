"""Row/column insert and delete with Excel's bookkeeping, for ``reshape_xlsx``.

openpyxl's ``insert_rows`` / ``delete_rows`` move cell values and styles and
nothing else: merged ranges stay where they were (so a merge now covers the
wrong rows), data validations and conditional formats keep their old ranges,
and no formula is rewritten — ``=SUM(B2:B10)`` keeps summing the old rows
after a row is inserted at 5. ``reshape_xlsx`` reported success on all of it.

Excel's semantics are *structural*, not copy semantics: every reference that
points at or past the insertion point moves, absolute or relative, on every
sheet; a reference into deleted cells becomes ``#REF!``; a range spanning the
change grows or shrinks. openpyxl's ``Translator`` implements copy semantics
(relative references follow the formula, absolute ones stay), so it is the
right tool for ``copy_row`` and the wrong one here; this module tokenizes each
formula with openpyxl's tokenizer and rewrites the references itself.
"""

from __future__ import annotations

import re
from typing import Any, Optional

_CELL_RE = re.compile(r"^(\$?)([A-Za-z]{1,3})(\$?)([0-9]+)$")
_COL_RE = re.compile(r"^(\$?)([A-Za-z]{1,3})$")
_ROW_RE = re.compile(r"^(\$?)([0-9]+)$")


class _Shift:
    """One structural change along one axis: insert ``count`` before ``at``,
    or delete ``count`` starting at ``at`` (1-based)."""

    def __init__(self, axis: str, at: int, count: int, delete: bool):
        self.axis = axis
        self.at = int(at)
        self.count = int(count)
        self.delete = delete
        self.end = self.at + self.count - 1

    def point(self, x: int) -> Optional[int]:
        if x < self.at:
            return x
        if not self.delete:
            return x + self.count
        if x <= self.end:
            return None
        return x - self.count

    def span(self, lo: int, hi: int) -> Optional[tuple[int, int]]:
        lo, hi = min(lo, hi), max(lo, hi)
        if not self.delete:
            return (self.point(lo) or lo, self.point(hi) or hi)
        new_lo = lo if lo < self.at else (lo - self.count if lo > self.end else self.at)
        new_hi = hi if hi < self.at else (hi - self.count if hi > self.end else self.at - 1)
        if new_lo > new_hi:
            return None
        return new_lo, new_hi

    def touches(self, lo: int, hi: int) -> bool:
        return max(lo, hi) >= self.at


def _col_index(letters: str) -> int:
    from openpyxl.utils import column_index_from_string

    return column_index_from_string(letters.upper())


def _col_letters(index: int) -> str:
    from openpyxl.utils import get_column_letter

    return get_column_letter(index)


def _split_sheet(ref: str) -> tuple[str, str]:
    if "!" in ref:
        sheet, rest = ref.rsplit("!", 1)
        return sheet, rest
    return "", ref


def _unquote(sheet: str) -> str:
    if len(sheet) >= 2 and sheet[0] == "'" and sheet[-1] == "'":
        return sheet[1:-1].replace("''", "'")
    return sheet


def shift_reference(ref: str, shift: _Shift) -> Optional[str]:
    """Rewrite one A1 reference (no sheet part) for ``shift``; ``None`` when
    every cell it named was deleted (Excel shows ``#REF!``). References that
    are not plain A1 (names, structured refs) come back unchanged."""
    parts = ref.split(":")
    if len(parts) > 2:
        return ref
    parsed: list[tuple[str, Any]] = []
    for part in parts:
        m = _CELL_RE.match(part)
        if m:
            parsed.append(("cell", m))
            continue
        m = _COL_RE.match(part)
        if m and len(parts) == 2:
            parsed.append(("col", m))
            continue
        m = _ROW_RE.match(part)
        if m and len(parts) == 2:
            parsed.append(("row", m))
            continue
        return ref
    kinds = {k for k, _ in parsed}
    if len(kinds) != 1:
        return ref
    kind = parsed[0][0]
    if kind == "col" and shift.axis == "row" or kind == "row" and shift.axis == "col":
        return ref

    def coord(entry: tuple[str, Any]) -> int:
        k, m = entry
        if k == "cell":
            return int(m.group(4)) if shift.axis == "row" else _col_index(m.group(2))
        if k == "row":
            return int(m.group(2))
        return _col_index(m.group(2))

    def rebuild(entry: tuple[str, Any], value: int) -> str:
        k, m = entry
        if k == "cell":
            if shift.axis == "row":
                return f"{m.group(1)}{m.group(2)}{m.group(3)}{value}"
            return f"{m.group(1)}{_col_letters(value)}{m.group(3)}{m.group(4)}"
        if k == "row":
            return f"{m.group(1)}{value}"
        return f"{m.group(1)}{_col_letters(value)}"

    if len(parsed) == 1:
        new = shift.point(coord(parsed[0]))
        return None if new is None else rebuild(parsed[0], new)
    a, b = coord(parsed[0]), coord(parsed[1])
    mapped = shift.span(a, b)
    if mapped is None:
        return None
    lo, hi = mapped
    if a <= b:
        return f"{rebuild(parsed[0], lo)}:{rebuild(parsed[1], hi)}"
    return f"{rebuild(parsed[0], hi)}:{rebuild(parsed[1], lo)}"


def shift_formula(formula: str, *, target_sheet: str, formula_sheet: str, shift: _Shift) -> tuple[str, bool, bool]:
    """``(new_formula, changed, became_ref_error)`` for one ``=…`` formula."""
    from openpyxl.formula.tokenizer import Token, Tokenizer

    if not isinstance(formula, str) or not formula.startswith("="):
        return formula, False, False
    try:
        tokenizer = Tokenizer(formula)
    except Exception:  # noqa: BLE001 — an unparsable formula is left alone
        return formula, False, False
    changed = False
    ref_error = False
    out: list[str] = []
    for token in tokenizer.items:
        value = token.value
        if token.type == Token.OPERAND and token.subtype == Token.RANGE and "[" not in value:
            sheet, ref = _split_sheet(value)
            sheet_name = _unquote(sheet) if sheet else formula_sheet
            if sheet_name == target_sheet and ":" not in sheet:
                new_ref = shift_reference(ref, shift)
                if new_ref is None:
                    value = "#REF!"
                    ref_error = True
                    changed = True
                elif new_ref != ref:
                    value = (sheet + "!" if sheet else "") + new_ref
                    changed = True
        out.append(value)
    return "=" + "".join(out), changed, ref_error


def _shift_multi_range(sqref: Any, shift: _Shift) -> tuple[list[str], bool]:
    """Shift a MultiCellRange / sqref string; returns (ranges, changed)."""
    from openpyxl.utils.cell import range_boundaries

    ranges = str(sqref).split()
    out: list[str] = []
    changed = False
    for rng in ranges:
        try:
            min_col, min_row, max_col, max_row = range_boundaries(rng)
        except Exception:  # noqa: BLE001
            out.append(rng)
            continue
        if min_col is None or min_row is None:
            out.append(rng)
            continue
        lo, hi = (min_row, max_row) if shift.axis == "row" else (min_col, max_col)
        mapped = shift.span(lo, hi)
        if mapped is None:
            changed = True
            continue
        if mapped != (lo, hi):
            changed = True
        if shift.axis == "row":
            min_row, max_row = mapped
        else:
            min_col, max_col = mapped
        start = f"{_col_letters(min_col)}{min_row}"
        end = f"{_col_letters(max_col)}{max_row}"
        out.append(start if start == end else f"{start}:{end}")
    return out, changed


def structural_shift(wb: Any, ws: Any, *, axis: str, at: int, count: int, delete: bool) -> dict[str, Any]:
    """Insert or delete rows/columns on ``ws`` and carry every dependent
    structure with them. Returns the side effects, for the tool result."""
    from openpyxl.formatting.formatting import ConditionalFormattingList
    from openpyxl.utils.cell import range_boundaries
    from openpyxl.worksheet.formula import ArrayFormula

    shift = _Shift(axis, at, count, delete)
    effects: dict[str, Any] = {}

    # 1. Merged ranges: unmerge those the change touches, re-merge mapped.
    affected = []
    for merged in list(ws.merged_cells.ranges):
        min_col, min_row, max_col, max_row = range_boundaries(str(merged))
        lo, hi = (min_row, max_row) if axis == "row" else (min_col, max_col)
        if shift.touches(lo, hi):
            affected.append((str(merged), min_col, min_row, max_col, max_row))
            ws.unmerge_cells(str(merged))

    # 2. Row heights / column widths travel with their rows / columns.
    dims_moved = _shift_dimensions(ws, shift)

    # 3. The cell move itself.
    if axis == "row":
        (ws.delete_rows if delete else ws.insert_rows)(at, amount=count)
    else:
        (ws.delete_cols if delete else ws.insert_cols)(at, amount=count)

    removed_merges: list[str] = []
    remerged = 0
    for original, min_col, min_row, max_col, max_row in affected:
        lo, hi = (min_row, max_row) if axis == "row" else (min_col, max_col)
        mapped = shift.span(lo, hi)
        if mapped is None:
            removed_merges.append(original)
            continue
        if axis == "row":
            min_row, max_row = mapped
        else:
            min_col, max_col = mapped
        if (min_row, min_col) == (max_row, max_col):
            removed_merges.append(original)
            continue
        ws.merge_cells(start_row=min_row, start_column=min_col, end_row=max_row, end_column=max_col)
        remerged += 1
    if remerged:
        effects["merged_ranges_shifted"] = remerged
    if removed_merges:
        effects["merged_ranges_removed"] = removed_merges

    # 4. Data validations: ranges and list/limit formulas.
    dv_shifted = dv_removed = 0
    validations = getattr(ws, "data_validations", None)
    if validations is not None:
        keep = []
        for dv in list(validations.dataValidation):
            ranges, changed = _shift_multi_range(dv.sqref, shift)
            if not ranges:
                dv_removed += 1
                continue
            if changed:
                dv.sqref = " ".join(ranges)
                dv_shifted += 1
            for attr in ("formula1", "formula2"):
                value = getattr(dv, attr, None)
                if value:
                    new, f_changed, _ = shift_formula("=" + str(value), target_sheet=ws.title, formula_sheet=ws.title, shift=shift)
                    if f_changed:
                        setattr(dv, attr, new[1:])
            keep.append(dv)
        validations.dataValidation = keep
    if dv_shifted:
        effects["data_validations_shifted"] = dv_shifted
    if dv_removed:
        effects["data_validations_removed"] = dv_removed

    # 5. Conditional formats: ranges and rule formulas.
    cf_shifted = cf_removed = 0
    if ws.conditional_formatting:
        rebuilt = ConditionalFormattingList()
        for cf in ws.conditional_formatting:
            ranges, changed = _shift_multi_range(cf.sqref, shift)
            if not ranges:
                cf_removed += 1
                continue
            if changed:
                cf_shifted += 1
            for rule in cf.rules:
                formulas = list(getattr(rule, "formula", None) or [])
                if formulas:
                    rule.formula = [
                        shift_formula("=" + f, target_sheet=ws.title, formula_sheet=ws.title, shift=shift)[0][1:]
                        for f in formulas
                    ]
                rebuilt.add(" ".join(ranges), rule)
        ws.conditional_formatting = rebuilt
    if cf_shifted:
        effects["conditional_formats_shifted"] = cf_shifted
    if cf_removed:
        effects["conditional_formats_removed"] = cf_removed

    # 6. Formulas on every sheet that point into the changed sheet.
    rewritten = 0
    ref_errors: list[str] = []
    for sheet in wb.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                value = cell.value
                if isinstance(value, ArrayFormula):
                    new_text, changed, err = shift_formula(value.text, target_sheet=ws.title, formula_sheet=sheet.title, shift=shift)
                    if changed:
                        value.text = new_text
                        rewritten += 1
                    if err:
                        ref_errors.append(f"{sheet.title}!{cell.coordinate}")
                    continue
                if not (isinstance(value, str) and value.startswith("=")):
                    continue
                new_value, changed, err = shift_formula(value, target_sheet=ws.title, formula_sheet=sheet.title, shift=shift)
                if changed:
                    cell.value = new_value
                    rewritten += 1
                if err:
                    ref_errors.append(f"{sheet.title}!{cell.coordinate}")
    if rewritten:
        effects["formulas_rewritten"] = rewritten
    if ref_errors:
        effects["formulas_now_ref_error"] = ref_errors[:50]

    # 7. Defined names, tables and the autofilter.
    names = 0
    for holder in (getattr(wb, "defined_names", None), getattr(ws, "defined_names", None)):
        if not holder:
            continue
        values = holder.values() if hasattr(holder, "values") else holder
        for dn in list(values):
            text = getattr(dn, "attr_text", None)
            if not text:
                continue
            new, changed, _ = shift_formula("=" + text, target_sheet=ws.title, formula_sheet=ws.title, shift=shift)
            if changed:
                dn.attr_text = new[1:]
                names += 1
    if names:
        effects["defined_names_rewritten"] = names
    tables = 0
    for table in getattr(ws, "tables", {}).values():
        ranges, changed = _shift_multi_range(table.ref, shift)
        if changed and ranges:
            table.ref = ranges[0]
            # A table whose autoFilter range disagrees with its own range is
            # a file Excel offers to "repair".
            if getattr(table, "autoFilter", None) is not None and table.autoFilter.ref:
                table.autoFilter.ref = ranges[0]
            tables += 1
    if tables:
        effects["tables_shifted"] = tables
    print_area = getattr(ws, "print_area", None)
    if print_area:
        refs = [_split_sheet(part.strip())[1].replace("$", "") for part in str(print_area).split(",") if part.strip()]
        shifted, changed = _shift_multi_range(" ".join(refs), shift)
        if changed and shifted:
            ws.print_area = shifted
            effects["print_area_shifted"] = True
    if ws.auto_filter is not None and ws.auto_filter.ref:
        ranges, changed = _shift_multi_range(ws.auto_filter.ref, shift)
        if changed and ranges:
            ws.auto_filter.ref = ranges[0]
            effects["autofilter_shifted"] = True
    if dims_moved:
        effects["row_heights_moved" if axis == "row" else "column_widths_moved"] = dims_moved
    anchored = len(getattr(ws, "_images", []) or []) + len(getattr(ws, "_charts", []) or [])
    if anchored:
        effects["not_adjusted"] = f"{anchored} image(s)/chart(s) keep their original anchors"
    return effects


def _shift_dimensions(ws: Any, shift: _Shift) -> int:
    moved = 0
    if shift.axis == "row":
        dims = ws.row_dimensions
        saved = {}
        for idx in sorted(list(dims.keys())):
            if idx < shift.at:
                continue
            dim = dims[idx]
            saved[idx] = (dim.height, dim.hidden, dim.outlineLevel)
            del dims[idx]
        for idx, (height, hidden, outline) in saved.items():
            new = shift.point(idx)
            if new is None:
                continue
            target = dims[new]
            target.height = height
            target.hidden = hidden
            target.outlineLevel = outline
            if height is not None or hidden:
                moved += 1
        return moved
    dims = ws.column_dimensions
    saved_cols = {}
    for letter in list(dims.keys()):
        try:
            idx = _col_index(letter)
        except ValueError:
            continue
        if idx < shift.at:
            continue
        dim = dims[letter]
        saved_cols[idx] = (dim.width, dim.hidden)
        del dims[letter]
    for idx, (width, hidden) in saved_cols.items():
        new = shift.point(idx)
        if new is None:
            continue
        target = dims[_col_letters(new)]
        target.width = width
        target.hidden = hidden
        if width or hidden:
            moved += 1
    return moved


def translate_copied_formula(value: Any, origin: str, destination: str) -> Any:
    """A formula copied from ``origin`` to ``destination`` (copy semantics:
    relative references follow, absolute ones stay) — what Excel's copy/paste
    does, via openpyxl's ``Translator``."""
    if not (isinstance(value, str) and value.startswith("=")):
        return value
    from openpyxl.formula.translate import Translator

    try:
        return Translator(value, origin=origin).translate_formula(destination)
    except Exception:  # noqa: BLE001 — leave an untranslatable formula as is
        return value


def copy_sheet_extras(source: Any, target: Any) -> dict[str, Any]:
    """What ``Workbook.copy_worksheet`` leaves behind: data validations,
    conditional formats, freeze panes, the autofilter and print titles.
    Formulas that name the source sheet explicitly are pointed at the copy, as
    Excel's Move or Copy does."""
    from copy import deepcopy

    from openpyxl.formatting.formatting import ConditionalFormattingList

    copied: dict[str, Any] = {}
    validations = getattr(source, "data_validations", None)
    if validations is not None and validations.dataValidation:
        for dv in validations.dataValidation:
            target.add_data_validation(deepcopy(dv))
        copied["data_validations"] = len(validations.dataValidation)
    if source.conditional_formatting:
        rebuilt = ConditionalFormattingList()
        count = 0
        for cf in source.conditional_formatting:
            for rule in cf.rules:
                rebuilt.add(str(cf.sqref), deepcopy(rule))
            count += 1
        target.conditional_formatting = rebuilt
        copied["conditional_formats"] = count
    if source.freeze_panes:
        target.freeze_panes = source.freeze_panes
        copied["freeze_panes"] = source.freeze_panes
    if source.auto_filter is not None and source.auto_filter.ref:
        target.auto_filter.ref = source.auto_filter.ref
        copied["autofilter"] = source.auto_filter.ref
    for attr in ("print_title_rows", "print_title_cols"):
        value = getattr(source, attr, None)
        if value:
            try:
                setattr(target, attr, value)
            except Exception:  # noqa: BLE001
                pass
    retargeted = _retarget_self_references(target, source.title)
    if retargeted:
        copied["self_references_retargeted"] = retargeted
    not_copied = len(getattr(source, "_images", []) or []) + len(getattr(source, "_charts", []) or [])
    if not_copied:
        copied["not_copied"] = f"{not_copied} image(s)/chart(s) (openpyxl cannot copy drawings)"
    return copied


def _quote_sheet(title: str) -> str:
    if re.fullmatch(r"[A-Za-z_][A-Za-z0-9_.]*", title):
        return title
    return "'" + title.replace("'", "''") + "'"


def _retarget_self_references(ws: Any, old_title: str) -> int:
    from openpyxl.formula.tokenizer import Token, Tokenizer

    count = 0
    for row in ws.iter_rows():
        for cell in row:
            value = cell.value
            if not (isinstance(value, str) and value.startswith("=")) or old_title not in value:
                continue
            try:
                tokens = Tokenizer(value).items
            except Exception:  # noqa: BLE001
                continue
            changed = False
            out = []
            for token in tokens:
                text = token.value
                if token.type == Token.OPERAND and token.subtype == Token.RANGE and "!" in text:
                    sheet, ref = _split_sheet(text)
                    if _unquote(sheet) == old_title:
                        text = f"{_quote_sheet(ws.title)}!{ref}"
                        changed = True
                out.append(text)
            if changed:
                cell.value = "=" + "".join(out)
                count += 1
    return count


__all__ = [
    "copy_sheet_extras",
    "shift_formula",
    "shift_reference",
    "structural_shift",
    "translate_copied_formula",
]
