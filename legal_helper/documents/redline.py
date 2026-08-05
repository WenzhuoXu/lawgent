"""Native Word tracked-changes (``w:ins`` / ``w:del``) redlining and comments.

``edit_docx_text_document`` silently rewrites text; this module produces the
document a lawyer actually forwards: real revision marks that Word renders in
the Review pane (accept/reject-able), plus margin comments anchored with
``w:commentRangeStart`` / ``w:commentRangeEnd``. Comment authoring uses the
python-docx >= 1.2 comments API; revision marks are authored at the oxml
level because python-docx has no tracked-change API.

Scope notes (v1, mirrors ``edit_docx_text_document`` caveats):

- A ``find`` string must fall within one paragraph (body or table cell).
- Boundary runs split by a match are collapsed to ``rPr`` + one ``w:t``,
  so tabs/breaks inside a *split* run are dropped; fully-covered runs keep
  their children.
- Ops apply in order; text already inside a ``w:ins`` / ``w:del`` from an
  earlier op is never re-matched (only direct ``w:r`` children of ``w:p``
  are scanned).
"""

from __future__ import annotations

import copy
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
from xml.etree import ElementTree as ET

DEFAULT_AUTHOR = "Legal AI Helper"

_W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def _wq(name: str) -> str:
    return f"{{{_W_NS}}}{name}"


def _utc_stamp(timestamp: str | None) -> str:
    if timestamp:
        return timestamp
    return datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")


# ---------------------------------------------------------------------------
# oxml run surgery
# ---------------------------------------------------------------------------


def _direct_runs(p_el) -> list:
    """Direct ``w:r`` children of a ``w:p`` — excludes runs already nested in
    ``w:ins`` / ``w:del`` / ``w:hyperlink``, so revisions never nest."""
    return [child for child in p_el if child.tag == _wq("r")]


def _run_text(r_el) -> str:
    return "".join(t.text or "" for t in r_el.findall(_wq("t")))


def _set_run_text(r_el, text: str) -> None:
    """Collapse a run to ``rPr`` + a single ``w:t`` holding ``text``."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    rpr = r_el.find(qn("w:rPr"))
    for child in list(r_el):
        r_el.remove(child)
    if rpr is not None:
        r_el.append(rpr)
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    r_el.append(t)


def _split_run(r_el, offset: int):
    """Split a run at ``offset``; the left half is inserted before ``r_el``.

    Returns ``(left, right)`` where ``right`` is the original element.
    """
    text = _run_text(r_el)
    left = copy.deepcopy(r_el)
    _set_run_text(left, text[:offset])
    _set_run_text(r_el, text[offset:])
    r_el.addprevious(left)
    return left, r_el


def _isolate_span(p_el, start: int, end: int) -> list:
    """Split boundary runs so ``[start, end)`` aligns with whole direct runs.

    Offsets are into the concatenation of direct-run text. Returns the run
    elements that exactly cover the span, in document order.
    """
    covered: list = []
    pos = 0
    for r in list(_direct_runs(p_el)):
        text = _run_text(r)
        r_start, r_end = pos, pos + len(text)
        pos = r_end
        if r_end <= start or r_start >= end or r_start == r_end:
            continue
        local_start = max(start - r_start, 0)
        local_end = min(end - r_start, len(text))
        if local_start > 0:
            _, r = _split_run(r, local_start)
            local_end -= local_start
        if local_end < len(_run_text(r)):
            r, _ = _split_run(r, local_end)
        covered.append(r)
    return covered


def _make_tracked(tag: str, rev_id: int, author: str, date: str):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    el = OxmlElement(tag)
    el.set(qn("w:id"), str(rev_id))
    el.set(qn("w:author"), author)
    el.set(qn("w:date"), date)
    return el


def _wrap_del(covered: list, rev_id: int, author: str, date: str):
    """Move ``covered`` runs into a new ``w:del``, converting text to
    ``w:delText`` (required — Word rejects ``w:t`` inside ``w:del``)."""
    from docx.oxml.ns import qn

    del_el = _make_tracked("w:del", rev_id, author, date)
    covered[0].addprevious(del_el)
    for r in covered:
        del_el.append(r)  # lxml moves the element
        for t in r.findall(qn("w:t")):
            t.tag = qn("w:delText")
            t.set(qn("xml:space"), "preserve")
    return del_el


def _make_ins(text: str, rpr_source, rev_id: int, author: str, date: str):
    """Build ``w:ins`` holding one run; run formatting copied from
    ``rpr_source`` so the insertion matches surrounding text."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    ins = _make_tracked("w:ins", rev_id, author, date)
    r = OxmlElement("w:r")
    if rpr_source is not None:
        rpr = rpr_source.find(qn("w:rPr"))
        if rpr is not None:
            r.append(copy.deepcopy(rpr))
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    r.append(t)
    ins.append(r)
    return ins


def _mark_comment_range(first_el, last_el, comment_id: int) -> None:
    """Anchor comment ``comment_id`` on ``[first_el, last_el]``.

    Markers are placed as ``w:p``-level siblings (never inside ``w:ins`` /
    ``w:del``), matching what Word emits when commenting revised text.
    """
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    start = OxmlElement("w:commentRangeStart")
    start.set(qn("w:id"), str(comment_id))
    first_el.addprevious(start)

    end = OxmlElement("w:commentRangeEnd")
    end.set(qn("w:id"), str(comment_id))
    last_el.addnext(end)

    ref_run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    rstyle = OxmlElement("w:rStyle")
    rstyle.set(qn("w:val"), "CommentReference")
    rpr.append(rstyle)
    ref_run.append(rpr)
    ref = OxmlElement("w:commentReference")
    ref.set(qn("w:id"), str(comment_id))
    ref_run.append(ref)
    end.addnext(ref_run)


def _find_spans(haystack: str, needle: str) -> list[tuple[int, int]]:
    spans: list[tuple[int, int]] = []
    i = haystack.find(needle)
    while i >= 0:
        spans.append((i, i + len(needle)))
        i = haystack.find(needle, i + len(needle))
    return spans


def _next_rev_id(document_el) -> int:
    from docx.oxml.ns import qn

    max_id = 0
    for tag in ("w:ins", "w:del"):
        for el in document_el.iter(qn(tag)):
            try:
                max_id = max(max_id, int(el.get(qn("w:id"), "0")))
            except ValueError:
                continue
    return max_id + 1


def _iter_all_paragraphs(doc):
    for para in doc.paragraphs:
        yield para
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    yield para


# ---------------------------------------------------------------------------
# Public API
# ---------------------------------------------------------------------------


def redline_docx_document(
    source_path: Path,
    output_path: Path,
    edits: list[dict[str, Any]],
    *,
    author: str = DEFAULT_AUTHOR,
    initials: str = "AI",
    timestamp: str | None = None,
) -> dict[str, Any]:
    """Copy a .docx applying edits as native tracked changes + comments.

    Each edit is a dict with:

    - ``op`` — ``replace`` (default) | ``delete`` | ``insert_after`` |
      ``comment``
    - ``find`` — exact anchor text (required, within one paragraph)
    - ``replace`` — replacement text (``replace``; empty ⇒ pure deletion)
    - ``text`` — inserted text (``insert_after``) or comment body
      (``comment``)
    - ``comment`` — optional margin note attached to any revision op

    ``timestamp`` is an ISO-8601 UTC string (``2026-07-21T09:00:00Z``);
    defaults to now. Returns per-find match counts plus revision/comment
    totals; unmatched finds are reported in ``not_found``, never dropped
    silently.
    """
    from docx import Document

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document(str(source_path))
    date = _utc_stamp(timestamp)
    rev_id = _next_rev_id(doc.element)

    applied: list[dict[str, Any]] = []
    not_found: list[str] = []
    revision_count = 0
    comment_count = 0

    for edit in edits:
        op = str(edit.get("op") or "replace").strip()
        find = str(edit.get("find") or "")
        if not find:
            raise ValueError(f"redline edit requires a non-empty 'find': {edit!r}")
        if op not in ("replace", "delete", "insert_after", "comment"):
            raise ValueError(f"Unknown redline op: {op!r}")
        replace = str(edit.get("replace") or "")
        note = str(edit.get("text") or "") if op in ("insert_after", "comment") else ""
        margin_note = str(edit.get("comment") or "")
        if op == "comment":
            margin_note = note
        if op == "replace" and not replace:
            op = "delete"

        count = 0
        for para in _iter_all_paragraphs(doc):
            p_el = para._p
            full = "".join(_run_text(r) for r in _direct_runs(p_el))
            spans = _find_spans(full, find)
            # Apply right-to-left so earlier character offsets stay valid.
            for start, end in reversed(spans):
                covered = _isolate_span(p_el, start, end)
                if not covered:
                    continue
                anchor_first = anchor_last = None
                if op in ("replace", "delete"):
                    del_el = _wrap_del(covered, rev_id, author, date)
                    rev_id += 1
                    revision_count += 1
                    anchor_first = anchor_last = del_el
                    if op == "replace":
                        ins = _make_ins(replace, covered[0], rev_id, author, date)
                        rev_id += 1
                        revision_count += 1
                        del_el.addnext(ins)
                        anchor_last = ins
                elif op == "insert_after":
                    ins = _make_ins(note, covered[-1], rev_id, author, date)
                    rev_id += 1
                    revision_count += 1
                    covered[-1].addnext(ins)
                    anchor_first = anchor_last = ins
                else:  # comment — anchor only, no text change
                    anchor_first, anchor_last = covered[0], covered[-1]
                if margin_note:
                    comment = doc.comments.add_comment(
                        text=margin_note, author=author, initials=initials
                    )
                    _mark_comment_range(anchor_first, anchor_last, comment.comment_id)
                    comment_count += 1
                count += 1

        applied.append({"op": op, "find": find, "count": count})
        if count == 0:
            not_found.append(find)

    doc.save(str(output_path))
    return {
        "output_path": str(output_path),
        "author": author,
        "date": date,
        "applied": applied,
        "revision_count": revision_count,
        "comment_count": comment_count,
        "not_found": not_found,
    }


def extract_docx_revisions(path: Path) -> list[dict[str, Any]]:
    """List tracked changes in a .docx: type, author, date, text, paragraph.

    Read-side counterpart of ``redline_docx_document``; works on any .docx
    (stdlib zip + ElementTree, no python-docx dependency).
    """
    with zipfile.ZipFile(path) as zf:
        root = ET.fromstring(zf.read("word/document.xml"))

    out: list[dict[str, Any]] = []
    for p_idx, p in enumerate(root.iter(_wq("p"))):
        for el in p.iter():
            if el.tag == _wq("ins"):
                text = "".join(t.text or "" for t in el.iter(_wq("t")))
                kind = "insertion"
            elif el.tag == _wq("del"):
                text = "".join(t.text or "" for t in el.iter(_wq("delText")))
                kind = "deletion"
            else:
                continue
            out.append(
                {
                    "type": kind,
                    "id": el.get(_wq("id"), ""),
                    "author": el.get(_wq("author"), ""),
                    "date": el.get(_wq("date"), ""),
                    "text": text,
                    "paragraph": p_idx,
                }
            )
    return out


def preview_docx_revisions(path: Path, mode: str = "accept") -> str:
    """Render document text as if revisions were accepted or rejected.

    ``mode="accept"`` keeps insertions and drops deletions; ``"reject"``
    restores the original text. Lets callers (and tests) verify a redline
    round-trips to exactly the intended final / original document.
    """
    if mode not in ("accept", "reject"):
        raise ValueError(f"mode must be 'accept' or 'reject', got {mode!r}")
    with zipfile.ZipFile(path) as zf:
        root = ET.fromstring(zf.read("word/document.xml"))

    t_tag, del_text_tag = _wq("t"), _wq("delText")
    ins_tag, del_tag = _wq("ins"), _wq("del")

    def _collect(el, in_ins: bool, in_del: bool, parts: list[str]) -> None:
        if el.tag == ins_tag:
            in_ins = True
        elif el.tag == del_tag:
            in_del = True
        if el.tag == t_tag and not in_del:
            if not in_ins or mode == "accept":
                parts.append(el.text or "")
        elif el.tag == del_text_tag and mode == "reject":
            parts.append(el.text or "")
        for child in el:
            _collect(child, in_ins, in_del, parts)

    lines: list[str] = []
    for p in root.iter(_wq("p")):
        parts: list[str] = []
        _collect(p, False, False, parts)
        lines.append("".join(parts))
    return "\n".join(lines)
