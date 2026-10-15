"""python-docx writer: Markdown sections in, a PRC-ready Word document out.

**Page.** A Chinese document is A4 and an English one Letter, decided by the
share of Han characters in the content (``_text.detect_lang``); ``page_size``
overrides it. Every page carries a footer with live ``PAGE``/``NUMPAGES``
fields — "第 X 页 共 Y 页" or "Page X of Y". All 115 harness DOCX measured on
2026-10-07 were Letter and none had a page number.

**Typography.** Chinese content gets ``w:eastAsia`` fonts (黑体-class
headings, 宋体-class body), black headings, CJK line-break and punctuation
rules, and a 2-character first-line indent **on body prose only** — headings
sit 顶格, table cells and centred lines carry no indent, because an indent
inside a one-field cell wraps the field and an indent on a centred line pushes
it off centre. Explicit fonts strip the style's theme-font attributes: a
``w:eastAsiaTheme`` beside a ``w:eastAsia`` wins in Word, which is how a
heading set to 黑体 rendered in the theme's 宋体.

**GB/T 9704 (``layout="gbt9704"``).** Checked against GB/T 9704-2012 and the
verifier of mizzlelover/gongwen-gbt9704-skill (MIT): A4; 天头 37 mm, 订口
28 mm, 版心 156×225 mm (so right 26 mm, bottom 35 mm); 3号仿宋 body on a 28 pt
line pitch (22 lines); 层次序数 一、/（一）/1./（1） as Heading 1-4 in 黑体 /
楷体 / 仿宋 / 仿宋, unbolded ("字体名称已经体现字重"), each 空二字; title in
2号小标宋; page numbers "— n —" in 4号宋体, odd pages right and even pages
left, each 空一字, 7 mm below the type area. The file names the GB fonts; they
render on the reader's Word/WPS. A render on a host without them uses the
stand-ins in ``documents/review_fonts.conf`` and changes nothing in the file.

**Markdown** comes from :mod:`._markdown`: real hyperlinks, real Word
footnotes, H1-H6, strikethrough, ``<br>``, nested lists, GFM tables with
alignment, inline code, code blocks, block quotes. Every ordered list gets its
own ``w:num`` with a ``w:startOverride``, so a second list starts at 1 rather
than at 19 — "诉讼请求第3项" must not become item 19 — and nesting is kept as
``w:ilvl``. Table header rows repeat on every page and rows never split.

**Templates.** ``reference_docx`` opens a firm or court template, clears its
body and writes into its own styles, page setup, headers, footers and
numbering.

**Disclaimer.** Never in the body. ``disclaimer`` text goes into the core
properties (``description``); the caller decides the language and whether
there is one at all.
"""

from __future__ import annotations

import datetime as _dt
import io
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Iterable, Optional, Sequence

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.opc.constants import CONTENT_TYPE as CT
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Emu, Inches, Pt, RGBColor, Twips
from docx.text.paragraph import Paragraph
from docx.text.run import Run

from ._markdown import Parsed, allowed_tag, gb_ordinal_level, heading_levels, inline_text, parse
from ._text import Lang, bcp47, display_width, resolve_lang

_EA_BODY_FONT = "宋体"          # SimSun-class default body
_EA_GBT_BODY_FONT = "仿宋"      # GB/T 9704 7.3.3: 3号仿宋
_EA_HEADING_FONT = "黑体"       # SimHei-class headings; GB/T first layer
_EA_GBT_H2_FONT = "楷体"        # GB/T second layer
_EA_TITLE_FONT = "宋体"
_EA_GBT_TITLE_FONT = "方正小标宋简体"  # GB/T 7.3.1: 2号小标宋
_EA_SUBTITLE_FONT = "楷体"      # CJK convention instead of faux-italic
_EA_PAGE_NUMBER_FONT = "宋体"   # GB/T 7.5: 4号半角宋体
_LATIN_FONT = "Times New Roman"
_CODE_FONT = "Courier New"

#: GB/T 9704 structure levels → East Asian font (7.3.3).
_GBT_LEVEL_FONTS = {1: _EA_HEADING_FONT, 2: _EA_GBT_H2_FONT, 3: _EA_GBT_BODY_FONT, 4: _EA_GBT_BODY_FONT}
#: Space before each GB level, in twips, as calibrated by gongwen-gbt9704-skill.
_GBT_LEVEL_BEFORE = {1: 160, 2: 80, 3: 40, 4: 0}

_PAGE_TWIPS = {"A4": (11906, 16838), "Letter": (12240, 15840)}
#: GB/T 9704 5.2.1: 天头 37 mm, 订口 28 mm, 版心 156 x 225 mm. In twips exactly
#: as the gongwen verifier pins them (2098 / 1474 / 1984 / 1588).
_GBT_MARGINS = {"top": 2098, "right": 1474, "bottom": 1984, "left": 1588}
#: 7.5: the page number's rule sits 7 mm under the type area; 20 mm footer
#: distance is the value gongwen-gbt9704-skill calibrated in Word and WPS.
_GBT_FOOTER_DISTANCE = 1134
#: 5.2.2 "一般每面排22行": 28 pt is the conventional pitch (225 mm / 28 pt = 22 lines).
_GBT_LINE_PITCH = 560
#: 5.2.2 "每行排28个字": the 156 mm measure holds 27 三号 glyphs at their natural
#: 16 pt, so the character grid tightens the pitch to 156 mm / 28. ``charSpace``
#: is that pitch minus the Normal font size, in 1/4096 pt (the encoding Word and
#: LibreOffice's importer share); LibreOffice 7.3 then sets exactly 28 a line.
_GBT_CHARS_PER_LINE = 28
_GBT_CHAR_SPACE = round(
    ((11906 - 1588 - 1474) / 20 / _GBT_CHARS_PER_LINE - 16.0) * 4096
)
#: 7.5 "单页码居右空一字，双页码居左空一字": one 4号 (14 pt) character.
_GBT_PAGE_NUMBER_INSET = 280


# ---------------------------------------------------------------------------
# Style helpers
# ---------------------------------------------------------------------------


def _rpr_of(style):
    return style.element.get_or_add_rPr()


_THEME_ATTRS = ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme")


def _drop_theme_fonts(rfonts, *which: str) -> None:
    """Remove theme-font attributes so the explicit names take effect."""
    for attr in which or _THEME_ATTRS:
        key = qn(attr)
        if key in rfonts.attrib:
            del rfonts.attrib[key]


def _set_east_asia_font(style, font_name: str, lang: str = "zh-CN") -> None:
    """Set ``w:rFonts/@w:eastAsia`` + ``w:lang/@w:eastAsia`` on a style."""
    rpr = _rpr_of(style)
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:eastAsia"), font_name)
    _drop_theme_fonts(rfonts, "w:eastAsiaTheme")
    lang_el = rpr.find(qn("w:lang"))
    if lang_el is None:
        lang_el = OxmlElement("w:lang")
        rpr.append(lang_el)  # w:lang sorts near the end of CT_RPr
    lang_el.set(qn("w:eastAsia"), lang)


def _set_latin_font(style, font_name: str) -> None:
    rfonts = _rpr_of(style).get_or_add_rFonts()
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
        rfonts.set(qn(attr), font_name)
    _drop_theme_fonts(rfonts, "w:asciiTheme", "w:hAnsiTheme", "w:cstheme")


def _set_run_east_asia_font(run, font_name: str) -> None:
    rpr = run._r.get_or_add_rPr()
    rpr.get_or_add_rFonts().set(qn("w:eastAsia"), font_name)


# pPr children that must follow the CJK layout toggles (schema order).
_CJK_TOGGLE_SUCCESSORS = (
    "w:bidi", "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind",
    "w:contextualSpacing", "w:jc", "w:textAlignment", "w:outlineLvl", "w:rPr",
)


def _enable_cjk_paragraph_rules(style) -> None:
    """kinsoku line breaking, hanging punctuation, CJK/Latin auto-spacing."""
    ppr = style.element.get_or_add_pPr()
    for tag in ("w:kinsoku", "w:overflowPunct", "w:autoSpaceDE", "w:autoSpaceDN"):
        if ppr.find(qn(tag)) is None:
            ppr.insert_element_before(OxmlElement(tag), *_CJK_TOGGLE_SUCCESSORS)


def _set_first_line_indent(el, chars: int, char_size_pt: float) -> None:
    """首行缩进 in characters (``w:firstLineChars``), twips as fallback.

    ``el`` is a ``w:style`` or a ``w:p``: the indent belongs on ``Normal`` so
    that prose typed into the finished document inherits it, which means every
    construct that is not prose has to cancel it — at style level where a style
    exists to carry the exception, directly on the paragraph where none does.
    Both attributes go in every time because LibreOffice honours ``w:firstLine``
    and Word honours ``w:firstLineChars``.
    """
    ind = el.get_or_add_pPr().get_or_add_ind()
    ind.set(qn("w:firstLineChars"), str(chars * 100))
    ind.set(qn("w:firstLine"), str(int(chars * char_size_pt * 20)))


def _suppress_first_line_indent(el) -> None:
    """Cancel an inherited 首行缩进 on one style or paragraph.

    Centred text must not carry a first-line indent: the indent comes off the
    line box before the remainder is centred, so the line lands half the indent
    to the right of the page centre. A table cell must not carry one either —
    it holds one short field, and the indent buys a second line for it.
    """
    _set_first_line_indent(el, 0, 0.0)


def _body_size_pt(*, cjk: bool, gbt9704: bool) -> float:
    """三号 for GB/T 9704, 小四 for other CJK, 11pt for Latin.

    The one place the body size is decided, because the table columns convert
    ems to twips with the same number and ``Normal.font.size`` is ``None``
    until ``_set_default_style`` has run.
    """
    if gbt9704:
        return 16.0
    return 12.0 if cjk else 11.0


def _iter_heading_styles(doc, *, gbt9704: bool):
    # GB/T 9704 has four structure levels; a memo can nest a body heading
    # down to Heading 9, and each of those carries the template's theme font.
    for level in range(1, 5 if gbt9704 else 10):
        yield doc.styles[f"Heading {level}"]


def _apply_cjk_typography(doc, *, gbt9704: bool) -> None:
    body_font = _EA_GBT_BODY_FONT if gbt9704 else _EA_BODY_FONT
    normal = doc.styles["Normal"]
    body_size = Pt(_body_size_pt(cjk=True, gbt9704=gbt9704))
    normal.font.size = body_size
    normal.font.color.rgb = RGBColor(0, 0, 0)
    _set_east_asia_font(normal, body_font)
    _enable_cjk_paragraph_rules(normal)
    _set_first_line_indent(normal.element, 2, body_size.pt)
    # 两端对齐 is the default for Chinese body text in Word and WPS.
    normal.paragraph_format.alignment = WD_PARAGRAPH_ALIGNMENT.JUSTIFY
    if gbt9704:
        pf = normal.paragraph_format
        pf.line_spacing = Pt(28)  # 28 pt fixed: 22 lines in the 225 mm type area
        pf.space_after = Pt(0)

    for idx, style in enumerate(_iter_heading_styles(doc, gbt9704=gbt9704), start=1):
        if gbt9704:
            ea = _GBT_LEVEL_FONTS[idx]
            style.font.size = Pt(16)
            # The face carries the weight: GB/T sets no Word bold on 黑体/楷体/仿宋.
            style.font.bold = False
            pf = style.paragraph_format
            pf.space_before = Twips(_GBT_LEVEL_BEFORE[idx])
            pf.space_after = Pt(0)
            pf.line_spacing = Pt(28)
            pf.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
            _set_latin_font(style, _LATIN_FONT)
        else:
            ea = _EA_HEADING_FONT
        _set_east_asia_font(style, ea)
        style.font.color.rgb = RGBColor(0, 0, 0)
        # The template's Heading 4 and 6 are italic; a slanted Han glyph is a
        # synthetic distortion, not a typeface, and no CJK convention uses it.
        style.font.italic = False
        if not gbt9704:
            # 顶格 headings: the 首行缩进 is a body-prose rule, and a heading
            # that inherits it reads as one more paragraph. GB/T 9704 headings
            # are 空两格, so that layout keeps the inheritance.
            _suppress_first_line_indent(style.element)

    title_style = doc.styles["Title"]
    title_style.font.color.rgb = RGBColor(0, 0, 0)
    if gbt9704:
        _set_east_asia_font(title_style, _EA_GBT_TITLE_FONT)
        _set_latin_font(title_style, _LATIN_FONT)
        title_style.font.size = Pt(22)  # 二号
        title_style.font.bold = False
        pf = title_style.paragraph_format
        pf.line_spacing = Pt(28)
        pf.space_after = Twips(_GBT_LINE_PITCH)  # 标题下空一行
        _drop_paragraph_border(title_style.element)
        _no_grid_snap(title_style.element)
        # The template's Title adds 0.25 pt letter-spacing and a kerning floor.
        rpr = _rpr_of(title_style)
        for tag in ("w:spacing", "w:kern"):
            for el in rpr.findall(qn(tag)):
                rpr.remove(el)
    else:
        _set_east_asia_font(title_style, _EA_TITLE_FONT)
        title_style.font.bold = True
        # The template's Title carries a blue accent rule; a Chinese 起诉状 or
        # 法律意见书 has a black title and no rule under it.
        _drop_paragraph_border(title_style.element)
    _suppress_first_line_indent(title_style.element)


_SNAP_SUCCESSORS = (
    "w:spacing", "w:ind", "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap",
    "w:jc", "w:textDirection", "w:textAlignment", "w:textboxTightWrap", "w:outlineLvl",
    "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange",
)


def _no_grid_snap(el) -> None:
    """Keep a style off the character grid.

    Word sets each East Asian glyph into whole grid cells, so a 二号 title or
    a 小五 footnote on the 28-a-line grid would be letter-spaced to the body's
    pitch. Chinese templates switch "对齐到网格" off for exactly these styles.
    """
    ppr = el.get_or_add_pPr()
    snap = ppr.find(qn("w:snapToGrid"))
    if snap is None:
        snap = OxmlElement("w:snapToGrid")
        ppr.insert_element_before(snap, *_SNAP_SUCCESSORS)
    snap.set(qn("w:val"), "0")


def _drop_paragraph_border(el) -> None:
    ppr = el.find(qn("w:pPr"))
    if ppr is not None:
        for bdr in ppr.findall(qn("w:pBdr")):
            ppr.remove(bdr)


def _set_default_style(doc) -> None:
    style = doc.styles["Normal"]
    style.font.name = _LATIN_FONT
    style.font.size = Pt(_body_size_pt(cjk=False, gbt9704=False))
    pf = style.paragraph_format
    pf.line_spacing = 1.15
    pf.space_after = Pt(6)


def _tune_settings(doc, *, lang: Lang) -> None:
    """Two document settings the built-in template gets wrong for us.

    ``w:doNotExpandShiftReturn``: a justified line that ends in a ``<br>`` is
    otherwise stretched across the whole measure, one glyph per few
    millimetres. ``w:themeFontLang``: the template says ``ja-JP``, so any text
    left on a theme font resolves to the Japanese face in Word.
    """
    settings = doc.settings.element
    compat = settings.find(qn("w:compat"))
    if compat is None:
        compat = OxmlElement("w:compat")
        settings.append(compat)
    if compat.find(qn("w:doNotExpandShiftReturn")) is None:
        compat.insert(0, OxmlElement("w:doNotExpandShiftReturn"))
    if lang == "zh":
        theme_lang = settings.find(qn("w:themeFontLang"))
        if theme_lang is not None:
            theme_lang.set(qn("w:eastAsia"), "zh-CN")


# ---------------------------------------------------------------------------
# Page
# ---------------------------------------------------------------------------


def _resolve_page_size(page_size: Optional[str], lang: Lang, gbt9704: bool) -> Optional[str]:
    value = (page_size or "auto").strip().lower()
    if value in {"", "auto"}:
        return "A4" if (gbt9704 or lang == "zh") else "Letter"
    if value == "a4":
        return "A4"
    if value == "letter":
        if gbt9704:
            raise ValueError("layout='gbt9704' is an A4 layout (GB/T 9704-2012 5.1); drop page_size='Letter'")
        return "Letter"
    raise ValueError(f"Unknown page_size: {page_size!r} (expected 'A4', 'Letter' or 'auto')")


def _set_page(doc, *, size: str, cjk: bool, gbt9704: bool) -> None:
    width, height = _PAGE_TWIPS[size]
    for section in doc.sections:
        section.page_width = Twips(width)
        section.page_height = Twips(height)
        if gbt9704:
            section.top_margin = Twips(_GBT_MARGINS["top"])
            section.right_margin = Twips(_GBT_MARGINS["right"])
            section.bottom_margin = Twips(_GBT_MARGINS["bottom"])
            section.left_margin = Twips(_GBT_MARGINS["left"])
            section.footer_distance = Twips(_GBT_FOOTER_DISTANCE)
            grid = section._sectPr.find(qn("w:docGrid"))
            if grid is None:
                grid = OxmlElement("w:docGrid")
                section._sectPr.append(grid)
            grid.set(qn("w:type"), "linesAndChars")
            grid.set(qn("w:linePitch"), str(_GBT_LINE_PITCH))
            grid.set(qn("w:charSpace"), str(_GBT_CHAR_SPACE))
        elif cjk:
            # Word's zh-CN default page: 2.54 cm top and bottom, 3.17 cm sides.
            section.top_margin = Twips(1440)
            section.bottom_margin = Twips(1440)
            section.left_margin = Twips(1800)
            section.right_margin = Twips(1800)
        else:
            section.left_margin = Inches(1)
            section.right_margin = Inches(1)
            section.top_margin = Inches(1)
            section.bottom_margin = Inches(1)


def _field(instr: str, *, rpr_xml: str = "", placeholder: str = "1"):
    """A ``w:fldSimple`` (``PAGE``, ``NUMPAGES``) Word, WPS and LibreOffice all update."""
    return parse_xml(
        f'<w:fldSimple {nsdecls("w")} w:instr="{instr}"><w:r>{rpr_xml}'
        f"<w:t>{placeholder}</w:t></w:r></w:fldSimple>"
    )


def _text_run(text: str, rpr_xml: str = ""):
    return parse_xml(f'<w:r {nsdecls("w")}>{rpr_xml}<w:t xml:space="preserve">{_xml_escape(text)}</w:t></w:r>')


def _xml_escape(text: str) -> str:
    return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def _page_number_rpr(*, ea_font: str, size_pt: float) -> str:
    half_points = int(round(size_pt * 2))
    return (
        f'<w:rPr><w:rFonts w:ascii="{ea_font}" w:hAnsi="{ea_font}" w:eastAsia="{ea_font}" '
        f'w:cs="{ea_font}"/><w:sz w:val="{half_points}"/><w:szCs w:val="{half_points}"/></w:rPr>'
    )


def _fresh_footer_paragraph(footer) -> Paragraph:
    footer.is_linked_to_previous = False
    paragraph = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)
    return paragraph


def _write_page_numbers(doc, *, lang: Lang, cjk: bool, gbt9704: bool) -> None:
    """Footer page numbers on every page of every section."""
    for section in doc.sections:
        if gbt9704:
            # 单页码居右空一字，双页码居左空一字 needs separate odd and even footers.
            doc.settings.odd_and_even_pages_header_footer = True
            rpr = _page_number_rpr(ea_font=_EA_PAGE_NUMBER_FONT, size_pt=14)
            for footer, align, part_name in (
                (section.footer, "right", "/word/footerOdd.xml"),
                (section.even_page_footer, "left", "/word/footerEven.xml"),
            ):
                paragraph = _fresh_footer_paragraph(footer)
                paragraph.alignment = (
                    WD_PARAGRAPH_ALIGNMENT.RIGHT if align == "right" else WD_PARAGRAPH_ALIGNMENT.LEFT
                )
                ind = paragraph._p.get_or_add_pPr().get_or_add_ind()
                ind.set(qn("w:firstLineChars"), "0")
                ind.set(qn("w:firstLine"), "0")
                ind.set(qn(f"w:{align}"), str(_GBT_PAGE_NUMBER_INSET))
                paragraph._p.append(_text_run("— ", rpr))
                paragraph._p.append(_field("PAGE", rpr_xml=rpr))
                paragraph._p.append(_text_run(" —", rpr))
                # Named the way the gongwen verifier looks them up; part names
                # are arbitrary in OOXML, so this costs nothing in Word.
                footer.part.partname = PackURI(part_name)
            continue
        paragraph = _fresh_footer_paragraph(section.footer)
        paragraph.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        if cjk:
            _suppress_first_line_indent(paragraph._p)
        if lang == "zh":
            rpr = _page_number_rpr(ea_font=_EA_PAGE_NUMBER_FONT, size_pt=10.5)
            pieces = ["第 ", "PAGE", " 页 共 ", "NUMPAGES", " 页"]
        else:
            half = '<w:rPr><w:sz w:val="18"/><w:szCs w:val="18"/></w:rPr>'
            rpr = half
            pieces = ["Page ", "PAGE", " of ", "NUMPAGES", ""]
        for piece in pieces:
            if piece in {"PAGE", "NUMPAGES"}:
                paragraph._p.append(_field(piece, rpr_xml=rpr))
            elif piece:
                paragraph._p.append(_text_run(piece, rpr))


def _has_footer(doc) -> bool:
    return any(
        section._sectPr.find(qn("w:footerReference")) is not None for section in doc.sections
    )


def _apply_watermark(doc, text: str, *, cjk: bool) -> None:
    """Stamp a repeating watermark banner into every section's page header.

    A true rotated diagonal VML watermark renders inconsistently across
    Word/LibreOffice; a bold, centered, light-red header banner repeats on
    every page in both and reads unambiguously as ``DO NOT FILE``. With odd and
    even footers on (GB/T 9704), the even-page header gets the banner too.
    """
    banner_color = RGBColor(0xC9, 0x44, 0x42)  # terracotta, matches UI --dnf/accent
    headers = []
    for section in doc.sections:
        headers.append(section.header)
        if doc.settings.odd_and_even_pages_header_footer:
            headers.append(section.even_page_header)
    for header in headers:
        header.is_linked_to_previous = False
        hp = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
        hp.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        hp.text = ""
        if cjk:
            _suppress_first_line_indent(hp._p)
        run = hp.add_run(text)
        run.bold = True
        run.font.size = Pt(12)
        run.font.color.rgb = banner_color
        if cjk:
            _set_run_east_asia_font(run, _EA_HEADING_FONT)


# ---------------------------------------------------------------------------
# Tables
# ---------------------------------------------------------------------------

_CELL_PAD_TWIPS = 216       # Word's default 108-twip left + right cell margin
_MIN_CELL_EM = 3.0          # a column narrower than three ems cannot hold 条号


def _measured_width(text: str) -> float:
    """Cell width in ems."""
    return display_width(text or "")


def _text_width_twips(doc) -> int:
    # Length arithmetic degrades to int, so re-wrap before asking for twips.
    section = doc.sections[0]
    return Emu(section.page_width - section.left_margin - section.right_margin).twips


def _column_widths(
    rows: list[list[str]], ncols: int, total_twips: int, char_size_pt: float
) -> list[int]:
    """Column widths in twips, summing to exactly ``total_twips``.

    A column's ideal width is its longest cell, pulled towards the column's
    typical cell so that one outlier does not starve its neighbours. When the
    ideals fit, the slack is shared out proportionally, because a table that
    stops short of the text width reads as a mistake; when they do not, the
    overflow is taken only from the columns that are above the floor.
    """
    em = char_size_pt * 20
    ideal: list[float] = []
    for c in range(ncols):
        widths = [_measured_width(r[c]) for r in rows if c < len(r) and r[c].strip()]
        longest = max(widths, default=1.0)
        typical = sum(widths) / len(widths) if widths else 1.0
        ideal.append((0.65 * longest + 0.35 * typical) * em + _CELL_PAD_TWIPS)

    floor = min(_MIN_CELL_EM * em + _CELL_PAD_TWIPS, total_twips / ncols)
    if sum(ideal) <= total_twips:
        slack = total_twips - sum(ideal)
        weight = sum(ideal)
        widths = [w + slack * w / weight for w in ideal]
    else:
        widths = [max(floor, w) for w in ideal]
        for _ in range(ncols):
            excess = sum(widths) - total_twips
            if excess <= 0:
                break
            head = [i for i, w in enumerate(widths) if w > floor]
            room = sum(widths[i] - floor for i in head)
            if room <= 0:
                break
            for i in head:
                widths[i] -= excess * (widths[i] - floor) / room

    out = [int(w) for w in widths]
    out[out.index(max(out))] += total_twips - sum(out)
    return out


def _fit_table(
    table, rows: list[list[str]], *, total_twips: int, char_size_pt: float, cjk: bool
) -> None:
    """Fixed layout, content-weighted columns, and no inherited 首行缩进.

    Both ``w:gridCol`` and every cell's ``w:tcW`` are written: ``add_table``
    stamps an equal ``tcW`` on each cell, and a stale one overrides the grid.
    """
    widths = _column_widths(rows, len(table.columns), total_twips, char_size_pt)
    table.autofit = False  # emits w:tblLayout w:type="fixed" in schema order
    tbl_w = table._tbl.tblPr.find(qn("w:tblW"))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(total_twips))
    grid_cols = table._tbl.find(qn("w:tblGrid")).findall(qn("w:gridCol"))
    for grid_col, width in zip(grid_cols, widths):
        grid_col.set(qn("w:w"), str(width))
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Twips(width)
            if cjk:
                for paragraph in cell.paragraphs:
                    _suppress_first_line_indent(paragraph._p)


def _mark_rows(table, *, header_rows: int) -> None:
    """Repeat the header row on every page; never split a row across pages.

    0 of 356 measured tables repeated their header, so a table that ran onto a
    second page showed columns with no names.
    """
    for index, row in enumerate(table.rows):
        tr_pr = row._tr.get_or_add_trPr()
        if tr_pr.find(qn("w:cantSplit")) is None:
            tr_pr.append(OxmlElement("w:cantSplit"))
        if index < header_rows and tr_pr.find(qn("w:tblHeader")) is None:
            tr_pr.append(OxmlElement("w:tblHeader"))


def _set_table_borders(table) -> None:
    """Single grid lines, for a template that has no ``Table Grid`` style."""
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "000000")
        borders.append(el)
    tbl_pr.insert_element_before(
        borders, "w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook", "w:tblCaption",
        "w:tblDescription", "w:tblPrChange",
    )


# ---------------------------------------------------------------------------
# Numbering
# ---------------------------------------------------------------------------

_BULLETS = ("•", "◦", "▪")


def _numbering_root(doc):
    """The ``w:numbering`` element, creating the part a template may lack.

    python-docx's own ``numbering_part`` raises ``NotImplementedError`` when it
    has to create one, so a firm template without any list in it would make
    the first Markdown list fail.
    """
    from docx.parts.numbering import NumberingPart

    try:
        return doc.part.part_related_by(RT.NUMBERING).element
    except KeyError:
        element = parse_xml(f'<w:numbering {nsdecls("w")}/>')
        part = NumberingPart(PackURI("/word/numbering.xml"), CT.WML_NUMBERING, element, doc.part.package)
        doc.part.relate_to(part, RT.NUMBERING)
        return part.element


class _Numbering:
    """One ``w:num`` per list, so every list restarts where it begins.

    Two abstract definitions (ordered, bullet) per document; each list gets a
    fresh ``w:num`` pointing at one of them with a ``w:startOverride`` on the
    list's own level. Word counts per abstract definition, so a ``w:num``
    without the override would continue the previous list — which is exactly
    the 19-instead-of-1 defect the single ``List Number`` style produced.
    """

    def __init__(self, doc, *, cjk: bool, char_pt: float) -> None:
        self.root = _numbering_root(doc)
        self.cjk = cjk
        self.char_pt = char_pt
        self._abstract: dict[bool, int] = {}

    def indent(self, depth: int, ordered: bool) -> tuple[int, int]:
        """``(left, hanging)`` in twips for a list level."""
        if self.cjk:
            em = int(self.char_pt * 20)
            hanging = 3 * em if (ordered and depth % 3 == 1) else 2 * em
            left = sum((3 * em if (ordered and d % 3 == 1) else 2 * em) for d in range(depth + 1))
            return left, hanging
        return 720 * (depth + 1), 360

    def _level_xml(self, ilvl: int, ordered: bool) -> str:
        left, hanging = self.indent(ilvl, ordered)
        if ordered:
            if self.cjk:
                fmt, text = (("decimal", f"%{ilvl + 1}."), ("decimal", f"（%{ilvl + 1}）"),
                             ("decimalEnclosedCircle", f"%{ilvl + 1}"))[ilvl % 3]
            else:
                fmt, text = (("decimal", f"%{ilvl + 1}."), ("lowerLetter", f"(%{ilvl + 1})"),
                             ("lowerRoman", f"(%{ilvl + 1})"))[ilvl % 3]
        else:
            fmt, text = "bullet", _BULLETS[ilvl % len(_BULLETS)]
        chars = (
            f' w:hangingChars="{int(round(hanging / (self.char_pt * 20) * 100))}"' if self.cjk else ""
        )
        return (
            f'<w:lvl w:ilvl="{ilvl}"><w:start w:val="1"/><w:numFmt w:val="{fmt}"/>'
            f'<w:lvlText w:val="{text}"/><w:lvlJc w:val="left"/>'
            f'<w:pPr><w:ind w:left="{left}" w:hanging="{hanging}"{chars}/></w:pPr></w:lvl>'
        )

    def _ensure_abstract(self, ordered: bool) -> int:
        if ordered in self._abstract:
            return self._abstract[ordered]
        used = [int(a.get(qn("w:abstractNumId"))) for a in self.root.findall(qn("w:abstractNum"))]
        abstract_id = max(used, default=-1) + 1
        levels = "".join(self._level_xml(i, ordered) for i in range(9))
        element = parse_xml(
            f'<w:abstractNum {nsdecls("w")} w:abstractNumId="{abstract_id}">'
            f'<w:multiLevelType w:val="hybridMultilevel"/>{levels}</w:abstractNum>'
        )
        first_num = self.root.find(qn("w:num"))
        if first_num is not None:
            first_num.addprevious(element)
        else:
            existing = self.root.findall(qn("w:abstractNum"))
            if existing:
                existing[-1].addnext(element)
            else:
                self.root.insert(0, element)
        self._abstract[ordered] = abstract_id
        return abstract_id

    def new_list(self, *, ordered: bool, depth: int, start: int) -> int:
        num = self.root.add_num(self._ensure_abstract(ordered))
        override = num.add_lvlOverride(depth)
        override.add_startOverride(max(0, int(start)))
        return int(num.numId)


# ---------------------------------------------------------------------------
# Footnotes
# ---------------------------------------------------------------------------

_FOOTNOTES_SKELETON = (
    f'<w:footnotes {nsdecls("w", "r")}>'
    '<w:footnote w:type="separator" w:id="-1"><w:p><w:pPr><w:spacing w:after="0" '
    'w:line="240" w:lineRule="auto"/></w:pPr><w:r><w:separator/></w:r></w:p></w:footnote>'
    '<w:footnote w:type="continuationSeparator" w:id="0"><w:p><w:pPr><w:spacing w:after="0" '
    'w:line="240" w:lineRule="auto"/></w:pPr><w:r><w:continuationSeparator/></w:r></w:p></w:footnote>'
    "</w:footnotes>"
)


class _Footnotes:
    """Real Word footnotes, in ``word/footnotes.xml``.

    The part is kept as XML while the document is built and serialised into
    the package once at the end, so a template's existing footnotes part
    (separators and all) is extended rather than replaced.
    """

    def __init__(self, builder: "_Builder") -> None:
        self.builder = builder
        doc_part = builder.doc.part
        self.part: Part | None = None
        for rel in doc_part.rels.values():
            if rel.reltype == RT.FOOTNOTES and not rel.is_external:
                self.part = rel.target_part
                break
        if self.part is None:
            self.part = Part(
                PackURI("/word/footnotes.xml"), CT.WML_FOOTNOTES, _FOOTNOTES_SKELETON.encode("utf-8"),
                doc_part.package,
            )
            doc_part.relate_to(self.part, RT.FOOTNOTES)
        element = getattr(self.part, "_element", None)
        self.root = element if element is not None else parse_xml(self.part.blob)
        ids = [int(fn.get(qn("w:id"))) for fn in self.root.findall(qn("w:footnote"))]
        self.next_id = max([0, *ids]) + 1
        self.text_style = builder.style_id("Footnote Text", WD_STYLE_TYPE.PARAGRAPH)
        self.ref_style = builder.style_id("Footnote Reference", WD_STYLE_TYPE.CHARACTER)

    def add(self, tokens: list) -> int:
        fid = self.next_id
        self.next_id += 1
        note = OxmlElement("w:footnote")
        note.set(qn("w:id"), str(fid))
        self.root.append(note)
        box = _FootnoteBox(self, note)
        self.builder.blocks(tokens, 0, len(tokens), _Ctx(), box)
        if box.first:  # an empty definition still needs its mark
            box.add_paragraph()
        return fid

    def ref_rpr(self) -> str:
        if self.ref_style:
            return f'<w:rPr><w:rStyle w:val="{self.ref_style}"/></w:rPr>'
        return '<w:rPr><w:vertAlign w:val="superscript"/></w:rPr>'

    def flush(self) -> None:
        if getattr(self.part, "_element", None) is None:
            from lxml import etree

            self.part._blob = etree.tostring(self.root, xml_declaration=True, encoding="UTF-8", standalone=True)


# ---------------------------------------------------------------------------
# Containers: the body, or one footnote
# ---------------------------------------------------------------------------


class _BodyBox:
    is_body = True

    def __init__(self, builder: "_Builder") -> None:
        self.builder = builder
        self.part = builder.doc.part

    def add_paragraph(self, style: Optional[str] = None) -> Paragraph:
        paragraph = self.builder.doc.add_paragraph()
        if style:
            self.builder.apply_style(paragraph, style)
        return paragraph


class _FootnoteBox:
    is_body = False

    def __init__(self, footnotes: _Footnotes, note) -> None:
        self.footnotes = footnotes
        self.note = note
        self.part = footnotes.part
        self.first = True

    def add_paragraph(self, style: Optional[str] = None) -> Paragraph:
        p = OxmlElement("w:p")
        self.note.append(p)
        paragraph = Paragraph(p, self)
        if self.footnotes.text_style:
            paragraph._p.get_or_add_pPr().get_or_add_pStyle().val = self.footnotes.text_style
        if self.first:
            self.first = False
            p.append(parse_xml(
                f'<w:r {nsdecls("w")}>{self.footnotes.ref_rpr()}<w:footnoteRef/></w:r>'
            ))
            p.append(_text_run(" "))
        return paragraph


# ---------------------------------------------------------------------------
# The builder
# ---------------------------------------------------------------------------


@dataclass
class _ListState:
    ordered: bool
    depth: int
    num_id: Optional[int]
    counter: int
    first_paragraph_pending: bool = False


@dataclass
class _Ctx:
    lists: list[_ListState] = field(default_factory=list)
    quote_depth: int = 0


@dataclass
class _Fmt:
    bold: bool = False
    italic: bool = False
    strike: bool = False
    underline: bool = False
    sup: bool = False
    sub: bool = False
    code: bool = False
    link: bool = False
    ea_font: Optional[str] = None


def _matching_close(tokens: list, index: int) -> int:
    opener = tokens[index]
    close_type = opener.type[: -len("_open")] + "_close"
    for j in range(index + 1, len(tokens)):
        if tokens[j].type == close_type and tokens[j].level == opener.level:
            return j
    return len(tokens) - 1


_SAFE_LINK_SCHEMES = ("http://", "https://", "mailto:")
_IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".tif", ".tiff"}


class _Builder:
    def __init__(
        self,
        doc,
        *,
        lang: Lang,
        gbt9704: bool,
        templated: bool,
        char_size_pt: float,
        image_roots: Sequence[Path] = (),
    ) -> None:
        self.doc = doc
        self.lang = lang
        self.cjk = lang == "zh"
        self.gbt = gbt9704
        self.templated = templated
        self.char_pt = char_size_pt
        self.image_roots = [Path(r).resolve() for r in image_roots]
        self.text_width = _text_width_twips(doc)
        self._numbering: Optional[_Numbering] = None
        self._footnotes: Optional[_Footnotes] = None
        self._defs: dict[int, list] = {}
        self._fn_display: dict[tuple[int, int], int] = {}
        self._fn_count = 0
        self._section_index = 0
        self.heading_shift = 0
        self._style_cache: dict[tuple[str, Any], Optional[str]] = {}

    # -- styles -------------------------------------------------------------

    def style_id(self, name: str, kind=WD_STYLE_TYPE.PARAGRAPH, *, create: bool = True) -> Optional[str]:
        key = (name, kind)
        if key in self._style_cache:
            return self._style_cache[key]
        style = None
        try:
            style = self.doc.styles[name]
        except KeyError:
            if create:
                style = self._create_style(name, kind)
        sid = style.style_id if style is not None else None
        if sid is not None or create:
            # A miss from a look-only call must not stop a later call creating it.
            self._style_cache[key] = sid
        return sid

    def _create_style(self, name: str, kind):
        styles = self.doc.styles
        try:
            style = styles.add_style(name, kind)
        except Exception:  # noqa: BLE001 - a template can refuse; fall back to Normal
            return None
        if name == "Footnote Text":
            style.base_style = styles["Normal"]
            style.font.size = Pt(9 if self.cjk else 10)
            pf = style.paragraph_format
            pf.space_after = Pt(0)
            pf.space_before = Pt(0)
            pf.line_spacing = 1.0
            if self.cjk:
                _suppress_first_line_indent(style.element)
            if self.gbt:
                _no_grid_snap(style.element)
        elif name == "Footnote Reference":
            style.font.superscript = True
        elif name == "Hyperlink":
            style.font.color.rgb = RGBColor(0x05, 0x63, 0xC1)
            style.font.underline = True
        return style

    def apply_style(self, paragraph: Paragraph, name: str) -> None:
        sid = self.style_id(name, WD_STYLE_TYPE.PARAGRAPH, create=False)
        if sid:
            paragraph._p.get_or_add_pPr().get_or_add_pStyle().val = sid

    # -- lazily created parts ----------------------------------------------

    @property
    def numbering(self) -> _Numbering:
        if self._numbering is None:
            self._numbering = _Numbering(self.doc, cjk=self.cjk, char_pt=self.char_pt)
        return self._numbering

    @property
    def footnotes(self) -> _Footnotes:
        if self._footnotes is None:
            self._footnotes = _Footnotes(self)
        return self._footnotes

    def finish(self) -> None:
        if self._footnotes is not None:
            self._footnotes.flush()

    # -- sections -------------------------------------------------------------

    def section(self, heading: str, body: str) -> None:
        self._section_index += 1
        parsed: Parsed = parse(body)
        self._defs = parsed.footnotes
        box = _BodyBox(self)
        if heading:
            self.heading_paragraph(heading, 1, box)
        levels = heading_levels(parsed.tokens)
        base = 2 if heading else 1
        self.heading_shift = (base - min(levels)) if levels else 0
        self.blocks(parsed.tokens, 0, len(parsed.tokens), _Ctx(), box)

    # -- blocks ---------------------------------------------------------------

    def blocks(self, toks: list, i: int, end: int, ctx: _Ctx, box) -> None:
        while i < end:
            tok = toks[i]
            kind = tok.type
            if kind == "heading_open":
                self.heading_tokens(toks[i + 1], int(tok.tag[1]), ctx, box)
                i += 3
            elif kind == "paragraph_open":
                self.paragraph(toks[i + 1], ctx, box)
                i += 3
            elif kind in {"bullet_list_open", "ordered_list_open"}:
                close = _matching_close(toks, i)
                self.list_block(toks, i, close, ctx, box)
                i = close + 1
            elif kind == "blockquote_open":
                close = _matching_close(toks, i)
                ctx.quote_depth += 1
                self.blocks(toks, i + 1, close, ctx, box)
                ctx.quote_depth -= 1
                i = close + 1
            elif kind == "table_open":
                close = _matching_close(toks, i)
                self.table(toks, i, close, ctx, box)
                i = close + 1
            elif kind in {"fence", "code_block"}:
                self.code(tok.content, ctx, box)
                i += 1
            elif kind == "hr":
                self.rule(box)
                i += 1
            else:
                i += 1

    def heading_paragraph(self, text: str, level: int, box) -> None:
        """A heading given as plain text (a section heading)."""
        from markdown_it.token import Token

        inline = Token("inline", "", 0)
        inline.children = [Token("text", "", 0)]
        inline.children[0].content = text
        inline.content = text
        self._heading(inline, level, box)

    def heading_tokens(self, inline, md_level: int, ctx: _Ctx, box) -> None:
        if not box.is_body:
            paragraph = box.add_paragraph()
            self.inline(paragraph, inline.children, box, _Fmt(bold=True))
            return
        self._heading(inline, md_level + self.heading_shift, box)

    def _heading(self, inline, level: int, box) -> None:
        text = inline_text(inline.children)
        level = max(1, min(level, 9))
        if self.gbt:
            gb_level = gb_ordinal_level(text) or min(level, 4)
            self._gb_heading(inline, gb_level, box)
            return
        paragraph = box.add_paragraph(f"Heading {level}")
        # A template without this heading level still gets a visible heading.
        bold = self.style_id(f"Heading {level}", create=False) is None
        self.inline(paragraph, inline.children, box, _Fmt(bold=bold))

    def _gb_heading(self, inline, level: int, box) -> None:
        paragraph = box.add_paragraph(f"Heading {level}")
        if not self.templated:
            # Direct formatting duplicates the style on purpose: the gongwen
            # verifier, like a reviewer, reads the paragraph, not the cascade.
            _set_first_line_indent(paragraph._p, 2, 16.0)
        fmt = _Fmt(ea_font=None if self.templated else _GBT_LEVEL_FONTS[level])
        self.inline(paragraph, inline.children, box, fmt)

    def paragraph(self, inline, ctx: _Ctx, box) -> None:
        children = inline.children or []
        if ctx.lists:
            state = ctx.lists[-1]
            if state.first_paragraph_pending:
                state.first_paragraph_pending = False
                self._list_paragraph(inline, state, box)
            else:
                self._continuation_paragraph(inline, state, box)
            return
        if self.gbt and box.is_body and not ctx.quote_depth:
            level = gb_ordinal_level(inline_text(children))
            if level:
                self._gb_heading(inline, level, box)
                return
        paragraph = box.add_paragraph()
        fmt = _Fmt()
        if ctx.quote_depth:
            paragraph.paragraph_format.left_indent = Inches(0.4 * ctx.quote_depth)
            if self.cjk:
                _suppress_first_line_indent(paragraph._p)
                fmt.ea_font = _EA_SUBTITLE_FONT
            else:
                fmt.italic = True
        self.inline(paragraph, children, box, fmt)

    # -- lists ----------------------------------------------------------------

    def list_block(self, toks: list, i: int, close: int, ctx: _Ctx, box) -> None:
        opener = toks[i]
        ordered = opener.type == "ordered_list_open"
        start = 1
        if ordered:
            try:
                start = int(opener.attrGet("start") or 1)
            except (TypeError, ValueError):
                start = 1
        depth = len(ctx.lists)
        num_id = None
        if box.is_body and not self.gbt:
            num_id = self.numbering.new_list(ordered=ordered, depth=depth, start=start)
        state = _ListState(ordered=ordered, depth=depth, num_id=num_id, counter=start)
        ctx.lists.append(state)
        j = i + 1
        while j < close:
            if toks[j].type == "list_item_open":
                item_close = _matching_close(toks, j)
                state.first_paragraph_pending = True
                self.blocks(toks, j + 1, item_close, ctx, box)
                state.counter += 1
                j = item_close + 1
            else:
                j += 1
        ctx.lists.pop()

    def _list_paragraph(self, inline, state: _ListState, box) -> None:
        children = inline.children or []
        if self.gbt and box.is_body:
            # 公文 numbers its structure in the text, not with Word lists.
            if state.ordered:
                prefix = f"{state.counter}." if state.depth == 0 else f"（{state.counter}）"
                level = gb_ordinal_level(prefix) or 3
                self._gb_heading(_prefixed(inline, prefix), level, box)
            else:
                paragraph = box.add_paragraph()
                self.inline(paragraph, children, box, _Fmt())
            return
        if state.num_id is None:  # inside a footnote: a typed marker
            paragraph = box.add_paragraph()
            marker = f"{state.counter}. " if state.ordered else f"{_BULLETS[state.depth % 3]} "
            paragraph._p.append(_text_run(marker))
            self.inline(paragraph, children, box, _Fmt())
            return
        paragraph = box.add_paragraph()
        self.apply_style(paragraph, "List Paragraph")
        p_pr = paragraph._p.get_or_add_pPr()
        num_pr = p_pr.get_or_add_numPr()
        num_pr.get_or_add_ilvl().val = state.depth
        num_pr.get_or_add_numId().val = state.num_id
        self._list_indent(paragraph, state, hanging=True)
        self.inline(paragraph, children, box, _Fmt())

    def _continuation_paragraph(self, inline, state: _ListState, box) -> None:
        paragraph = box.add_paragraph()
        if box.is_body and not self.gbt:
            self.apply_style(paragraph, "List Paragraph")
            self._list_indent(paragraph, state, hanging=False)
        self.inline(paragraph, inline.children or [], box, _Fmt())

    def _list_indent(self, paragraph: Paragraph, state: _ListState, *, hanging: bool) -> None:
        """Direct indentation, so ``Normal``'s 首行缩进 cannot reach the list.

        Word gives a character-based indent precedence over a twips one, so
        the CJK track writes ``w:hangingChars`` beside ``w:hanging``: the style's
        ``w:firstLineChars`` would otherwise beat the numbering's hanging indent.
        """
        left, hang = self.numbering.indent(state.depth, state.ordered)
        ind = paragraph._p.get_or_add_pPr().get_or_add_ind()
        ind.set(qn("w:left"), str(left))
        if hanging:
            ind.set(qn("w:hanging"), str(hang))
            if self.cjk:
                ind.set(qn("w:hangingChars"), str(int(round(hang / (self.char_pt * 20) * 100))))
        else:
            ind.set(qn("w:firstLine"), "0")
            if self.cjk:
                ind.set(qn("w:firstLineChars"), "0")

    # -- tables ---------------------------------------------------------------

    def table(self, toks: list, i: int, close: int, ctx: _Ctx, box) -> None:
        rows: list[list[tuple[list, Optional[str]]]] = []
        header_rows = 0
        in_head = False
        for j in range(i + 1, close):
            tok = toks[j]
            if tok.type == "thead_open":
                in_head = True
            elif tok.type == "thead_close":
                in_head = False
            elif tok.type == "tr_open":
                rows.append([])
                if in_head:
                    header_rows += 1
            elif tok.type in {"th_open", "td_open"}:
                style = tok.attrGet("style") or ""
                align = style.split(":", 1)[1].strip() if "text-align:" in style else None
                children = toks[j + 1].children if toks[j + 1].type == "inline" else []
                if rows:
                    rows[-1].append((children or [], align))
        if not rows:
            return
        if not box.is_body:
            for row in rows:
                paragraph = box.add_paragraph()
                paragraph._p.append(_text_run(" | ".join(inline_text(c) for c, _ in row)))
            return
        ncols = max(len(r) for r in rows)
        table = self.doc.add_table(rows=len(rows), cols=ncols)
        if self.templated:
            if self.style_id("Table Grid", WD_STYLE_TYPE.TABLE, create=False):
                table.style = "Table Grid"
            else:
                _set_table_borders(table)
        else:
            # Plain black grid for CJK docs (GB-style), accent grid otherwise.
            table.style = "Table Grid" if self.cjk else "Light Grid Accent 1"
        plain: list[list[str]] = []
        for r_idx, row in enumerate(rows):
            plain.append([inline_text(c) for c, _ in row])
            for c_idx, (children, align) in enumerate(row[:ncols]):
                paragraph = table.rows[r_idx].cells[c_idx].paragraphs[0]
                if align in {"center", "right"}:
                    paragraph.alignment = (
                        WD_PARAGRAPH_ALIGNMENT.CENTER if align == "center" else WD_PARAGRAPH_ALIGNMENT.RIGHT
                    )
                self.inline(paragraph, children, box, _Fmt(bold=r_idx < header_rows))
        _mark_rows(table, header_rows=header_rows)
        _fit_table(
            table, plain, total_twips=self.text_width, char_size_pt=self.char_pt, cjk=self.cjk
        )

    # -- code and rules -------------------------------------------------------

    def code(self, content: str, ctx: _Ctx, box) -> None:
        paragraph = box.add_paragraph()
        if self.cjk:
            _suppress_first_line_indent(paragraph._p)
        paragraph.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
        lines = (content or "").rstrip("\n").split("\n")
        for index, line in enumerate(lines):
            run = Run(OxmlElement("w:r"), paragraph)
            paragraph._p.append(run._r)
            if index:
                run.add_break()
            run.add_text(line)
            run.font.name = _CODE_FONT
            run.font.size = Pt(max(self.char_pt - 1.5, 9))

    def rule(self, box) -> None:
        paragraph = box.add_paragraph()
        paragraph.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        if self.cjk:
            _suppress_first_line_indent(paragraph._p)
        paragraph._p.append(_text_run("* * *"))

    # -- inline ---------------------------------------------------------------

    def inline(self, paragraph: Paragraph, children: list, box, base: _Fmt) -> None:
        fmt = _Fmt(**vars(base))
        stack: list[tuple[str, _Fmt]] = []
        container = paragraph._p

        def push(kind: str) -> None:
            stack.append((kind, _Fmt(**vars(fmt))))

        def pop(kind: str) -> None:
            nonlocal fmt
            for index in range(len(stack) - 1, -1, -1):
                if stack[index][0] == kind:
                    fmt = stack[index][1]
                    del stack[index:]
                    return

        for child in children or []:
            kind = child.type
            if kind == "text":
                self._run(container, paragraph, child.content, fmt)
            elif kind == "code_inline":
                self._run(container, paragraph, child.content, _Fmt(**{**vars(fmt), "code": True}))
            elif kind in {"softbreak", "hardbreak"}:
                self._break(container, paragraph)
            elif kind == "strong_open":
                push("strong")
                fmt.bold = True
            elif kind == "strong_close":
                pop("strong")
            elif kind == "em_open":
                push("em")
                fmt.italic = True
            elif kind == "em_close":
                pop("em")
            elif kind == "s_open":
                push("s")
                fmt.strike = True
            elif kind == "s_close":
                pop("s")
            elif kind == "link_open":
                push("link")
                href = (child.attrGet("href") or "").strip()
                link = self._hyperlink(paragraph, href, box)
                if link is not None:
                    container = link
                    fmt.link = True
            elif kind == "link_close":
                pop("link")
                container = paragraph._p
            elif kind == "html_inline":
                tag = allowed_tag(child.content)
                if tag is None:
                    self._run(container, paragraph, child.content, fmt)
                    continue
                name, closing, _ = tag
                if name == "br":
                    self._break(container, paragraph)
                elif closing:
                    pop(name)
                else:
                    push(name)
                    if name in {"b", "strong"}:
                        fmt.bold = True
                    elif name in {"i", "em"}:
                        fmt.italic = True
                    elif name == "u":
                        fmt.underline = True
                    elif name in {"s", "del", "strike"}:
                        fmt.strike = True
                    elif name == "sup":
                        fmt.sup = True
                    elif name == "sub":
                        fmt.sub = True
            elif kind == "footnote_ref":
                self._footnote_ref(container, paragraph, child.meta or {}, box)
            elif kind == "image":
                self._image(container, paragraph, child, box, fmt)

    def _run(self, container, paragraph: Paragraph, text: str, fmt: _Fmt) -> Optional[Run]:
        if not text:
            return None
        run = Run(OxmlElement("w:r"), paragraph)
        container.append(run._r)
        run.add_text(text)
        if fmt.bold:
            run.bold = True
        if fmt.italic:
            run.italic = True
        if fmt.strike:
            run.font.strike = True
        if fmt.underline:
            run.underline = True
        if fmt.sup:
            run.font.superscript = True
        if fmt.sub:
            run.font.subscript = True
        if fmt.code:
            run.font.name = _CODE_FONT
        if fmt.link and not self.gbt:
            sid = self.style_id("Hyperlink", WD_STYLE_TYPE.CHARACTER)
            if sid:
                run._r.get_or_add_rPr().get_or_add_rStyle().val = sid
        if fmt.ea_font:
            _set_run_east_asia_font(run, fmt.ea_font)
        return run

    def _break(self, container, paragraph: Paragraph) -> None:
        run = Run(OxmlElement("w:r"), paragraph)
        container.append(run._r)
        run.add_break()

    def _hyperlink(self, paragraph: Paragraph, href: str, box):
        if not href.lower().startswith(_SAFE_LINK_SCHEMES):
            return None  # an anchor, a relative path or a script: keep the text only
        r_id = box.part.relate_to(href, RT.HYPERLINK, is_external=True)
        link = OxmlElement("w:hyperlink")
        link.set(qn("r:id"), r_id)
        link.set(qn("w:history"), "1")
        paragraph._p.append(link)
        return link

    def _footnote_ref(self, container, paragraph: Paragraph, meta: dict, box) -> None:
        local = int(meta.get("id", 0))
        key = (self._section_index, local)
        if not box.is_body:
            # A note inside a note: Word has no such thing; keep the marker.
            self._run(container, paragraph, f"[{meta.get('label', local + 1)}]", _Fmt(sup=True))
            return
        if key in self._fn_display:
            # A second reference to the same note repeats its number.
            self._run(container, paragraph, str(self._fn_display[key]), _Fmt(sup=True))
            return
        fid = self.footnotes.add(self._defs.get(local, []))
        self._fn_count += 1
        self._fn_display[key] = self._fn_count
        container.append(parse_xml(
            f'<w:r {nsdecls("w")}>{self.footnotes.ref_rpr()}<w:footnoteReference w:id="{fid}"/></w:r>'
        ))

    def _image(self, container, paragraph: Paragraph, child, box, fmt: _Fmt) -> None:
        alt = inline_text(child.children) or child.content or ""
        path = self._allowed_image(child.attrGet("src") or "") if box.is_body else None
        if path is None:
            self._run(container, paragraph, alt, fmt)
            return
        run = Run(OxmlElement("w:r"), paragraph)
        container.append(run._r)
        try:
            shape = run.add_picture(str(path))
        except Exception:  # noqa: BLE001 - an unreadable image degrades to its alt text
            container.remove(run._r)
            self._run(container, paragraph, alt, fmt)
            return
        limit = Twips(self.text_width)
        if shape.width > limit:
            ratio = limit / shape.width
            shape.width = int(shape.width * ratio)
            shape.height = int(shape.height * ratio)

    def _allowed_image(self, src: str) -> Optional[Path]:
        """An image the caller allowed: an existing file inside an image root."""
        if not src or "://" in src or not self.image_roots:
            return None
        candidate = Path(src).expanduser()
        if not candidate.is_absolute():
            candidate = self.image_roots[0] / candidate
        try:
            resolved = candidate.resolve()
        except OSError:
            return None
        if resolved.suffix.lower() not in _IMAGE_SUFFIXES or not resolved.is_file():
            return None
        if not any(root == resolved or root in resolved.parents for root in self.image_roots):
            return None
        return resolved


def _prefixed(inline, prefix: str):
    """A copy of an inline token with ``prefix`` typed in front."""
    from markdown_it.token import Token

    copy_ = Token("inline", "", 0)
    head = Token("text", "", 0)
    head.content = prefix
    copy_.children = [head, *(inline.children or [])]
    copy_.content = prefix + (inline.content or "")
    return copy_


# ---------------------------------------------------------------------------
# Templates and properties
# ---------------------------------------------------------------------------

_TEMPLATE_CT = "application/vnd.openxmlformats-officedocument.wordprocessingml.template.main+xml"
_DOCUMENT_CT = "application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"


def _open_template(path: Path):
    """Open a ``.docx`` or ``.dotx`` template with python-docx."""
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix not in {".docx", ".dotx"}:
        raise ValueError(f"reference_docx must be a .docx or .dotx file, got {path.name}")
    if suffix == ".docx":
        return Document(str(path))
    # A .dotx is a .docx whose main part declares the template content type.
    source = zipfile.ZipFile(path)
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as out:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = data.replace(_TEMPLATE_CT.encode(), _DOCUMENT_CT.encode())
            out.writestr(item, data)
    source.close()
    buffer.seek(0)
    return Document(buffer)


def _clear_body(doc) -> None:
    """Remove the template's content, keeping the final section properties.

    Footnotes, endnotes and comments anchored in that content go with it:
    left behind, they are notes with no reference, which LibreOffice prints.
    """
    body = doc.element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)
    _purge_orphans(doc)


_NOTE_KEEP_TYPES = {"separator", "continuationSeparator", "continuationNotice"}


def _purge_orphans(doc) -> None:
    from lxml import etree

    for rel in list(doc.part.rels.values()):
        if rel.is_external or rel.reltype not in {RT.FOOTNOTES, RT.ENDNOTES, RT.COMMENTS}:
            continue
        part = rel.target_part
        element = getattr(part, "_element", None)
        root = element if element is not None else etree.fromstring(part.blob)
        for child in list(root):
            if rel.reltype == RT.COMMENTS or child.get(qn("w:type")) not in _NOTE_KEEP_TYPES:
                root.remove(child)
        if element is None:
            part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def _set_properties(doc, *, title: str, lang: Lang, description: Optional[str]) -> None:
    """Title, language and the disclaimer, where Word shows them: File → Info."""
    props = doc.core_properties
    props.title = (title or "")[:255]
    props.language = bcp47(lang)
    props.comments = description or ""
    # Not "python-docx", and not a template author who never saw this file.
    props.author = ""
    props.last_modified_by = ""
    now = _dt.datetime.now(_dt.timezone.utc).replace(microsecond=0)
    props.created = now
    props.modified = now
    props.revision = 1


def write_docx_document(
    *,
    output_path: Path | str,
    title: str,
    sections: Iterable[tuple[str, str]],
    subtitle: str | None = None,
    disclaimer: str | None = None,
    lang: str | None = None,
    layout: str | None = None,
    watermark: str | None = None,
    page_size: str | None = None,
    reference_docx: Path | str | None = None,
    image_roots: Sequence[Path | str] = (),
) -> Path:
    """Render a legal-style .docx document.

    Each section is a ``(heading, body_markdown)`` pair; the heading becomes a
    level-1 heading and the body's own headings nest under it.

    ``lang``: ``"zh"``/``"zh-CN"``, ``"en"``, or ``None``/``"auto"`` to decide
    by Han share. ``layout``: ``"memo"`` (default) or ``"gbt9704"`` (implies
    Chinese). ``page_size``: ``"A4"``, ``"Letter"`` or ``None``/``"auto"`` (A4
    for Chinese). ``disclaimer`` is written to the document properties, not the
    body. ``reference_docx`` builds on a template's styles, page setup, headers
    and footers. ``image_roots`` lists the directories a Markdown image may be
    read from; an image anywhere else is rendered as its alt text.

    ``watermark`` stamps a repeating banner on every page (e.g.
    ``"DO NOT FILE · 禁止对外提交"``) — used when the citation audit returns a
    DO-NOT-FILE verdict so an exported draft cannot be mistaken for a filed
    document.
    """
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    layout_key = (layout or "memo").strip().lower()
    if layout_key not in {"memo", "gbt9704"}:
        raise ValueError(f"Unknown layout: {layout!r} (expected 'memo' or 'gbt9704')")
    gbt9704 = layout_key == "gbt9704"
    section_pairs = [(heading or "", body or "") for heading, body in sections]
    texts = [title, subtitle or ""] + [f"{h}\n{b}" for h, b in section_pairs]
    if gbt9704 and (lang or "auto").strip().lower() in {"", "auto"}:
        doc_lang: Lang = "zh"
    else:
        doc_lang = resolve_lang(lang, *texts)
    cjk = doc_lang == "zh"
    char_size_pt = _body_size_pt(cjk=cjk, gbt9704=gbt9704)

    templated = reference_docx is not None and str(reference_docx).strip() != ""
    if templated:
        doc = _open_template(Path(reference_docx))
        _clear_body(doc)
        explicit = (page_size or "auto").strip().lower()
        if explicit not in {"", "auto"}:
            size = _resolve_page_size(page_size, doc_lang, False)
            for section in doc.sections:
                section.page_width = Twips(_PAGE_TWIPS[size][0])
                section.page_height = Twips(_PAGE_TWIPS[size][1])
    else:
        doc = Document()
        _set_default_style(doc)
        _set_page(
            doc,
            size=_resolve_page_size(page_size, doc_lang, gbt9704),
            cjk=cjk,
            gbt9704=gbt9704,
        )
        if cjk:
            _apply_cjk_typography(doc, gbt9704=gbt9704)
        _tune_settings(doc, lang=doc_lang)

    builder = _Builder(
        doc,
        lang=doc_lang,
        gbt9704=gbt9704,
        templated=templated,
        char_size_pt=char_size_pt,
        image_roots=[Path(r) for r in image_roots],
    )

    if not templated or not _has_footer(doc):
        _write_page_numbers(doc, lang=doc_lang, cjk=cjk, gbt9704=gbt9704 and not templated)

    if watermark:
        # Repeat the verdict on every page header, and add a page-one banner.
        _apply_watermark(doc, watermark, cjk=cjk)
        banner = doc.add_paragraph()
        banner.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        if cjk:
            _suppress_first_line_indent(banner._p)
        br = banner.add_run(watermark)
        br.bold = True
        br.font.size = Pt(16)
        br.font.color.rgb = RGBColor(0xC9, 0x44, 0x42)
        if cjk:
            _set_run_east_asia_font(br, _EA_HEADING_FONT)

    if title:
        title_p = doc.add_paragraph()
        has_title_style = builder.style_id("Title", WD_STYLE_TYPE.PARAGRAPH, create=False) is not None
        if has_title_style:
            builder.apply_style(title_p, "Title")
        title_p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        if cjk and not has_title_style:
            _suppress_first_line_indent(title_p._p)
        run = title_p.add_run(title)
        if not has_title_style:
            run.bold = True
    if subtitle:
        sub = doc.add_paragraph()
        sub.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        if cjk:
            _suppress_first_line_indent(sub._p)
        r = sub.add_run(subtitle)
        if cjk:
            _set_run_east_asia_font(r, _EA_SUBTITLE_FONT)
        else:
            r.italic = True

    for heading, body in section_pairs:
        builder.section(heading, body)

    builder.finish()
    _set_properties(doc, title=title, lang=doc_lang, description=disclaimer)
    doc.save(str(output_path))
    return output_path


__all__ = ["write_docx_document"]
