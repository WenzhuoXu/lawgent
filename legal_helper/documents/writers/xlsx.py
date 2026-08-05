"""XLSX writer.

Beyond plain scalar rows the writer also accepts:

- Formula strings (``"=SUM(B2:B9)"``). Pair the writer with
  ``recalc_xlsx_workbook`` to evaluate them on save so downstream
  ``data_only=True`` reads see real values.
- Native charts (bar, line, pie, scatter) via openpyxl.chart.
- Conditional formatting rules: color scales, cell-is, formula-based,
  and data bars.
- Data-validation drop-downs and numeric/date ranges.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any


_CHART_CLASSES = {
    "bar": "BarChart",
    "line": "LineChart",
    "pie": "PieChart",
    "scatter": "ScatterChart",
}


def _build_chart(ws, spec: dict[str, Any]):
    from openpyxl.chart import BarChart, LineChart, PieChart, ScatterChart, Reference
    from openpyxl.utils.cell import range_boundaries

    kind = str(spec.get("type", "")).lower()
    cls_name = _CHART_CLASSES.get(kind)
    if not cls_name:
        return None
    chart_cls = {"BarChart": BarChart, "LineChart": LineChart, "PieChart": PieChart, "ScatterChart": ScatterChart}[cls_name]
    chart = chart_cls()
    title = spec.get("title") or ""
    if title:
        chart.title = title
    chart.style = 10
    series = spec.get("series") or []
    for s in series:
        rng = str(s.get("values_range") or "").strip()
        if not rng:
            continue
        min_col, min_row, max_col, max_row = range_boundaries(rng)
        ref = Reference(ws, min_col=min_col, min_row=min_row, max_col=max_col, max_row=max_row)
        chart.add_data(ref, titles_from_data=False)
        name = str(s.get("name") or "")
        if name and chart.series:
            from openpyxl.chart.series import SeriesLabel
            chart.series[-1].tx = SeriesLabel(v=name)
    cats_range = str(spec.get("categories_range") or "").strip()
    if cats_range and hasattr(chart, "set_categories"):
        min_col, min_row, max_col, max_row = range_boundaries(cats_range)
        cats = Reference(ws, min_col=min_col, min_row=min_row, max_col=max_col, max_row=max_row)
        chart.set_categories(cats)
    width = spec.get("width")
    height = spec.get("height")
    if isinstance(width, (int, float)):
        chart.width = float(width)
    if isinstance(height, (int, float)):
        chart.height = float(height)
    return chart


def _apply_conditional_format(ws, spec: dict[str, Any]) -> None:
    from openpyxl.formatting.rule import (
        ColorScaleRule,
        CellIsRule,
        FormulaRule,
        DataBarRule,
    )
    from openpyxl.styles import Font, PatternFill

    rng = str(spec.get("range") or "").strip()
    if not rng:
        return
    kind = str(spec.get("type", "")).lower()

    def _fill(color: str) -> PatternFill | None:
        color = color.lstrip("#").upper()
        if not color:
            return None
        if len(color) == 6:
            color = "FF" + color
        return PatternFill(start_color=color, end_color=color, fill_type="solid")

    def _font() -> Font | None:
        fg = str(spec.get("foreground") or "").lstrip("#").upper()
        bold = bool(spec.get("bold"))
        if not fg and not bold:
            return None
        if fg and len(fg) == 6:
            fg = "FF" + fg
        return Font(color=fg or None, bold=bold)

    if kind == "color_scale":
        rule = ColorScaleRule(
            start_type="min",
            start_color=(spec.get("color_min") or "F8696B").lstrip("#"),
            mid_type="percentile",
            mid_value=50,
            mid_color=(spec.get("color_mid") or "FFEB84").lstrip("#"),
            end_type="max",
            end_color=(spec.get("color_max") or "63BE7B").lstrip("#"),
        )
    elif kind == "data_bar":
        rule = DataBarRule(
            start_type="min",
            end_type="max",
            color=(spec.get("color") or "638EC6").lstrip("#"),
            showValue=True,
        )
    elif kind == "cell_is":
        operator = str(spec.get("operator") or "equal")
        formula = [str(f) for f in (spec.get("formula") or [])] or [""]
        rule = CellIsRule(
            operator=operator,
            formula=formula,
            fill=_fill(str(spec.get("background") or "")),
            font=_font(),
        )
    elif kind == "formula":
        formula = [str(f) for f in (spec.get("formula") or [])] or [""]
        rule = FormulaRule(
            formula=formula,
            fill=_fill(str(spec.get("background") or "")),
            font=_font(),
        )
    else:
        return
    ws.conditional_formatting.add(rng, rule)


def _apply_data_validation(ws, spec: dict[str, Any]) -> None:
    from openpyxl.worksheet.datavalidation import DataValidation

    rng = str(spec.get("range") or "").strip()
    if not rng:
        return
    dv = DataValidation(
        type=str(spec.get("type") or "list"),
        operator=str(spec.get("operator") or "between") or None,
        formula1=str(spec.get("formula1") or "") or None,
        formula2=str(spec.get("formula2") or "") or None,
        allow_blank=True,
        showErrorMessage=bool(spec.get("error_msg")),
        errorTitle=spec.get("error_title") or None,
        error=spec.get("error_msg") or None,
        showInputMessage=bool(spec.get("prompt")),
        promptTitle=spec.get("prompt_title") or None,
        prompt=spec.get("prompt") or None,
    )
    dv.add(rng)
    ws.add_data_validation(dv)


def write_xlsx_workbook(
    path: Path,
    sheets: list[
        tuple[str, list[list[Any]]]
        | dict[str, Any]
    ],
) -> None:
    """Write a workbook from ``sheets``.

    Each entry is either a legacy ``(name, rows)`` tuple or a dict with keys
    ``name`` / ``rows`` / ``charts`` / ``conditional_formats`` /
    ``data_validations`` / ``freeze_panes``. Formula strings in ``rows`` are
    written verbatim so a downstream recalc step can evaluate them.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    path.parent.mkdir(parents=True, exist_ok=True)
    safe_sheets = sheets or [("Sheet1", [])]
    wb = Workbook()
    wb.remove(wb.active)
    used_names: set[str] = set()
    for idx, sheet in enumerate(safe_sheets, start=1):
        if isinstance(sheet, tuple):
            name, rows = sheet
            charts: list[dict[str, Any]] = []
            cond_formats: list[dict[str, Any]] = []
            validations: list[dict[str, Any]] = []
            freeze = "A2"
        else:
            name = sheet.get("name") or f"Sheet{idx}"
            rows = sheet.get("rows") or []
            charts = sheet.get("charts") or []
            cond_formats = sheet.get("conditional_formats") or []
            validations = sheet.get("data_validations") or []
            freeze = sheet.get("freeze_panes", "A2")

        title = (name or f"Sheet{idx}")[:31]
        while title in used_names:
            suffix = f"_{idx}"
            title = f"{title[:31 - len(suffix)]}{suffix}"
        used_names.add(title)
        ws = wb.create_sheet(title=title)

        for row in rows:
            ws.append(row)
        if rows:
            for cell in ws[1]:
                cell.font = Font(bold=True)
                cell.fill = PatternFill(fill_type="solid", fgColor="EAF2F8")
            if freeze:
                ws.freeze_panes = freeze
            for column in ws.columns:
                max_len = max(
                    (len(str(cell.value)) for cell in column if cell.value is not None),
                    default=8,
                )
                ws.column_dimensions[column[0].column_letter].width = min(max(max_len + 2, 10), 60)

        for cf in cond_formats:
            _apply_conditional_format(ws, cf)
        for dv in validations:
            _apply_data_validation(ws, dv)
        for chart_spec in charts:
            chart = _build_chart(ws, chart_spec)
            if chart is not None:
                anchor = str(chart_spec.get("anchor") or "F2")
                ws.add_chart(chart, anchor)

    wb.save(path)
    wb.close()
