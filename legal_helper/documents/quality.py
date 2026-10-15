"""Deterministic conformance checks on a document we just wrote.

**This is a floor, not a quality score, and it cannot tell you whether the
document is any good.** It catches a shape 0.4in off the canvas, 8pt labels
and 1.9:1 contrast — defects that are objective, invisible in a thumbnail,
and free to find. It also passed, as "clean", a deck whose author called it empty with no
storytelling, because a card rectangle counts as content whether or not
anything is written in it.

So: a clean conformance result means nothing here is provably broken. Whether
the deck says anything is a judgement, and it belongs to the reader — and, for
a deck made by ``create_deck``, to ppt-master's own quality gate, which runs
before this does.

**A density metric was considered and deliberately not added.** The generated
deck averaged 97 em of text per slide; the reference decks in
``outputs/chat_uploads/`` average 82–206, so a density floor would have passed
it too. What distinguishes the real decks is the *distribution* — sparse
dividers and then a 500–660 em exhibit that does the work — and a floor on the
mean is satisfied by padding every slide, which is worse than an empty one.
Do not add one.

**A CJK language-tag check was measured and not added.** Both real
ppt-master exports (215 Han runs) tag every one ``zh-CN`` and set it in the
theme's East Asian font, so on the only deck path the check could never fire.

Two sources, deliberately:

* the file, which knows what was *specified* — shape boxes, run sizes, fills,
  table grids — and so can prove a shape is off-canvas or a label cannot fit;
* the render, which knows what a reader will *see* — how much of the canvas
  carries ink, and where the dead bands are — which no amount of reading the
  XML can tell you.

Findings are structured, not prose, because they are appended to tool results
that carry inline images and therefore skip the per-tool token budget
(``providers/*.py`` bypass ``apply_result_budget`` for image-bearing results).
``QualityReport.to_payload`` enforces its own cap instead.
"""

from __future__ import annotations

import re
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Literal, Sequence

from .graph_layout import display_width, wrap_label

Severity = Literal["critical", "major", "minor"]

#: Type floors for a 13.333in canvas seen from the back of a room.
BODY_FONT_FLOOR_PT = 11.0
NOTE_FONT_FLOOR_PT = 9.0

#: Two shapes may overlap by this fraction of the smaller one's area before it
#: reads as a collision rather than a deliberate nudge.
OVERLAP_TOLERANCE = 0.02

#: An empty box has to be big in BOTH directions to be a hole in the layout;
#: a 12x0.25in strip is a rule or a spacer, and flagging those buries the
#: finding that matters under noise.
EMPTY_SHAPE_MIN_AREA_IN2 = 3.0
EMPTY_SHAPE_MIN_EDGE_IN = 0.5

#: A shape thinner than this in its short dimension is a rule, a divider or a
#: connector, all of which are meant to sit under other shapes.
HAIRLINE_IN = 0.5

#: Line height multiplier used to estimate wrapped text height. Kept slightly
#: under PowerPoint's own single spacing so the estimate errs toward "fits": a
#: box sized exactly to its own line box would otherwise read as overflowing.
LINE_HEIGHT = 1.16

#: How far past its box text has to run before it is worth reporting. Text in
#: a PowerPoint text box spills rather than clips, so a few percent is invisible;
#: a quarter of a line is not.
OVERFLOW_WARN = 1.08
OVERFLOW_BLOCK = 1.25

#: Serialized size ceiling for a report handed back through a tool result.
PAYLOAD_CHAR_CAP = 2000
_MAX_EXAMPLES_PER_CHECK = 3


#: Checks that need pixels. Empty by design: occupancy is a proxy padding
#: satisfies, so the render stage reports metrics and finds nothing.
RENDER_CHECKS: frozenset[str] = frozenset()


@dataclass
class Finding:
    """One defect, with the fix stated so the model does not have to infer it.

    ``slide`` and ``shape`` are indices rather than prose so a finding names a
    place the model can go and look at; ``where`` repeats them, with the
    shape's name, for the payload the model reads.
    """

    check: str
    severity: Severity
    where: str
    detail: str
    fix: str
    #: 1-based slide number, 0 for a whole-deck finding.
    slide: int = 0
    #: 1-based z-order index of the offending shape, when there is one; a
    #: dotted path such as ``"4.2"`` for the second child of the group that is
    #: the fourth shape on the slide.
    shape: int | str | None = None

    def as_dict(self) -> dict[str, Any]:
        return {
            "check": self.check,
            "severity": self.severity,
            "where": self.where,
            "detail": self.detail,
            "fix": self.fix,
        }


#: ``conformance`` for a file type nothing here knows how to check. An XLSX
#: render used to report "pass" with no check having run, which reads as
#: "measured and fine" when nothing was measured.
NO_CHECKS = "no checks for this type"


@dataclass
class QualityReport:
    findings: list[Finding] = field(default_factory=list)
    metrics: dict[str, Any] = field(default_factory=dict)
    #: Whether any file check ran. ``False`` reports :data:`NO_CHECKS` (or
    #: "not checked" when the file could not be read) instead of "pass";
    #: ``None`` (a report built by hand) is judged by its findings alone.
    checked: bool | None = None

    @property
    def critical(self) -> list[Finding]:
        return [f for f in self.findings if f.severity == "critical"]

    @property
    def ok(self) -> bool:
        return not self.critical

    def _verdict(self) -> str:
        if self.findings:
            return "blocked" if self.critical else "warnings"
        if self.checked is False:
            if "file_check_error" in self.metrics:
                return "not checked: the file could not be read"
            return NO_CHECKS
        return "pass"

    def to_payload(self, *, cap: int = PAYLOAD_CHAR_CAP) -> dict[str, Any]:
        """Collapse to a bounded, machine-readable payload.

        Identical checks are grouped: twenty overflowing boxes are one finding
        with twenty locations, not twenty findings, which is both shorter and
        more useful — it says "the type scale is wrong", not "slide 7 again".
        """
        grouped: dict[tuple[str, str], list[Finding]] = {}
        for f in self.findings:
            grouped.setdefault((f.check, f.severity), []).append(f)
        order = {"critical": 0, "major": 1, "minor": 2}
        rows: list[dict[str, Any]] = []
        for (check, severity), group in sorted(
            grouped.items(), key=lambda kv: (order[kv[0][1]], kv[0][0])
        ):
            rows.append({
                "check": check,
                "severity": severity,
                "count": len(group),
                "where": [f.where for f in group[:_MAX_EXAMPLES_PER_CHECK]],
                "detail": group[0].detail,
                "fix": group[0].fix,
            })
        payload: dict[str, Any] = {
            # "conformance", never "quality": the only thing a pass means is
            # that nothing here is provably broken.
            "conformance": self._verdict(),
            "counts": {s: sum(1 for f in self.findings if f.severity == s)
                       for s in ("critical", "major", "minor") if any(f.severity == s for f in self.findings)},
            "metrics": self.metrics,
            "findings": rows,
        }
        while len(repr(payload)) > cap and payload["findings"]:
            payload["findings"].pop()
            payload["truncated"] = True
        return payload

    def summary(self) -> str:
        if not self.findings:
            verdict = self._verdict()
            if verdict != "pass":
                return f"conformance: {verdict}"
            return "conformance: pass (says nothing about whether it reads well)"
        parts = [f"{sum(1 for f in self.findings if f.severity == s)} {s}"
                 for s in ("critical", "major", "minor")
                 if any(f.severity == s for f in self.findings)]
        return "conformance: " + ", ".join(parts)


# ---------------------------------------------------------------------------
# Colour
# ---------------------------------------------------------------------------


def relative_luminance(rgb: tuple[int, int, int]) -> float:
    """WCAG 2.x relative luminance of an sRGB triple."""
    channels = []
    for raw in rgb:
        c = raw / 255.0
        channels.append(c / 12.92 if c <= 0.03928 else ((c + 0.055) / 1.055) ** 2.4)
    r, g, b = channels
    return 0.2126 * r + 0.7152 * g + 0.0722 * b


def contrast_ratio(fg: tuple[int, int, int], bg: tuple[int, int, int]) -> float:
    """WCAG 2.x contrast ratio, 1.0 (identical) to 21.0 (black on white)."""
    l1, l2 = relative_luminance(fg), relative_luminance(bg)
    lighter, darker = max(l1, l2), min(l1, l2)
    return (lighter + 0.05) / (darker + 0.05)


def contrast_floor(font_pt: float, bold: bool) -> float:
    """WCAG large-text exemption: 18pt, or 14pt when bold, drops to 3:1."""
    return 3.0 if font_pt >= 18.0 or (bold and font_pt >= 14.0) else 4.5


# ---------------------------------------------------------------------------
# Geometry helpers
# ---------------------------------------------------------------------------


def _area(box: tuple[float, float, float, float]) -> float:
    return max(box[2], 0.0) * max(box[3], 0.0)


def _intersection(a: tuple[float, float, float, float], b: tuple[float, float, float, float]) -> float:
    ox = min(a[0] + a[2], b[0] + b[2]) - max(a[0], b[0])
    oy = min(a[1] + a[3], b[1] + b[3]) - max(a[1], b[1])
    return max(ox, 0.0) * max(oy, 0.0)


def _contains(outer: tuple[float, float, float, float], inner: tuple[float, float, float, float],
              slack: float = 0.05) -> bool:
    return (outer[0] - slack <= inner[0] and outer[1] - slack <= inner[1]
            and outer[0] + outer[2] + slack >= inner[0] + inner[2]
            and outer[1] + outer[3] + slack >= inner[1] + inner[3])


def estimate_text_block(
    text: str, *, font_pt: float, width_in: float, wrap: bool = True
) -> tuple[float, float]:
    """Estimated ``(width_in, height_in)`` of ``text`` set at ``font_pt``.

    Shares :func:`~.graph_layout.display_width` with the diagram layout so the
    two never disagree about how wide a Han glyph is.
    """
    em_in = font_pt / 72.0
    if not wrap:
        return display_width(text) * em_in, font_pt * LINE_HEIGHT / 72.0
    max_em = max(width_in / em_in, 1.0)
    lines: list[str] = []
    for para in (text or "").splitlines() or [""]:
        lines.extend(wrap_label(para, max_em))
    widest = max((display_width(ln) for ln in lines), default=0.0) * em_in
    return widest, len(lines) * font_pt * LINE_HEIGHT / 72.0


# ---------------------------------------------------------------------------
# PPTX
# ---------------------------------------------------------------------------


def _run_colour(run: Any) -> tuple[int, int, int] | None:
    """Explicit RGB of a run, or None when it inherits from theme or layout.

    python-pptx raises rather than returning a value for a theme-coloured run,
    and an inherited colour is genuinely unknowable without resolving the whole
    theme chain — so an unresolved colour is skipped rather than guessed, which
    would manufacture contrast findings out of nothing.
    """
    try:
        colour = run.font.color
        if colour is None or colour.type is None:
            return None
        rgb = colour.rgb
    except Exception:
        return None
    return (rgb[0], rgb[1], rgb[2]) if rgb is not None else None


def _shape_fill_colour(shape: Any) -> tuple[int, int, int] | None:
    """Solid fill colour, or None when it is inherited, absent or translucent.

    ``fore_color`` raises rather than returning None on a no-fill shape, and a
    fill carrying an ``<a:alpha>`` composites against whatever is behind it —
    so its effective colour is unknowable and a contrast ratio computed from
    the nominal one is fiction: a node filled with the accent at 70%
    transparency reads as pale blue and measured 1.9:1 against the nominal navy.
    """
    from pptx.oxml.ns import qn

    try:
        fill = shape.fill
        if fill.type is None or fill.type != 1:  # MSO_FILL.SOLID
            return None
        if fill._xPr is not None and fill._xPr.findall(".//" + qn("a:alpha")):
            return None
        rgb = fill.fore_color.rgb
    except Exception:
        return None
    return (rgb[0], rgb[1], rgb[2]) if rgb is not None else None


#: Preset geometries that are a stroke rather than a box. A diagonal line's
#: bounding box is as large as the rectangle it spans, so a connector between
#: two nodes "overlaps" both of them by a fifth of their area — measuring
#: thinness cannot tell them apart, but the geometry can.
_STROKE_PRESETS = frozenset({
    "line", "straightConnector1", "bentConnector2", "bentConnector3",
    "bentConnector4", "bentConnector5", "curvedConnector2", "curvedConnector3",
    "curvedConnector4", "curvedConnector5",
})


def _preset(shape: Any) -> str:
    from pptx.oxml.ns import qn

    try:
        geom = shape._element.find(".//" + qn("a:prstGeom"))
        return str(geom.get("prst") or "") if geom is not None else ""
    except Exception:
        return ""


def _is_stroke(shape: Any) -> bool:
    return _preset(shape) in _STROKE_PRESETS


def _has_visible_fill(shape: Any) -> bool:
    try:
        return shape.fill.type is not None and shape.fill.type != 5  # MSO_FILL.BACKGROUND
    except Exception:
        return False


def _frame_run_sizes(frame: Any, default_pt: float = 18.0) -> list[tuple[float, bool, str]]:
    """``(size_pt, bold, text)`` per run, with the paragraph/default fallback."""
    out: list[tuple[float, bool, str]] = []
    for para in frame.paragraphs:
        para_pt = para.font.size.pt if para.font.size is not None else None
        for run in para.runs:
            if not run.text.strip():
                continue
            pt = run.font.size.pt if run.font.size is not None else (para_pt or default_pt)
            out.append((pt, bool(run.font.bold), run.text))
    return out


def _run_sizes(shape: Any, default_pt: float = 18.0) -> list[tuple[float, bool, str]]:
    return _frame_run_sizes(shape.text_frame, default_pt)


_EMU_PER_IN = 914400.0

#: ``(tx, ty, sx, sy)``: slide EMU = ``t + s * child EMU``, per axis.
_Transform = tuple[float, float, float, float]
_IDENTITY: _Transform = (0.0, 0.0, 1.0, 1.0)


@dataclass
class _Placed:
    """A leaf shape with its box in slide inches and where it sits."""

    shape: Any
    box: tuple[float, float, float, float]
    #: ``"4"`` for the fourth top-level shape, ``"4.2"`` for the second child
    #: of a group that is the fourth top-level shape.
    path: str
    #: Inches of slide per inch of the shape's own coordinates, per axis.
    scale: tuple[float, float]

    @property
    def index(self) -> int | str:
        return int(self.path) if "." not in self.path else self.path


def _group_transform(group: Any, parent: _Transform) -> _Transform:
    """Compose a group's child-to-parent mapping onto ``parent``.

    A group stores its children in its own coordinate space (``chOff`` /
    ``chExt``) and places that space at ``off`` / ``ext`` on its parent. A
    resized group therefore scales every child, and a lint that read child
    coordinates as slide coordinates would put them in the wrong place.
    """
    from pptx.oxml.ns import qn

    tx, ty, sx, sy = parent
    try:
        xfrm = group._element.find(qn("p:grpSpPr")).find(qn("a:xfrm"))
        off, ext = xfrm.find(qn("a:off")), xfrm.find(qn("a:ext"))
        ch_off, ch_ext = xfrm.find(qn("a:chOff")), xfrm.find(qn("a:chExt"))
        ox, oy = float(off.get("x")), float(off.get("y"))
        cx, cy = float(ext.get("cx")), float(ext.get("cy"))
        chx, chy = float(ch_off.get("x")), float(ch_off.get("y"))
        chcx, chcy = float(ch_ext.get("cx")), float(ch_ext.get("cy"))
    except Exception:  # noqa: BLE001 — a group without a full xfrm keeps its parent's mapping
        return parent
    kx = cx / chcx if chcx else 1.0
    ky = cy / chcy if chcy else 1.0
    return (tx + sx * (ox - chx * kx), ty + sy * (oy - chy * ky), sx * kx, sy * ky)


def _placed_shapes(shapes: Any, transform: _Transform = _IDENTITY, prefix: str = ""):
    """Every leaf shape on a slide, groups opened, in z-order.

    Group shapes used to be checked as opaque boxes, which is how a ppt-master
    deck — 38 groups holding all 112 of its text frames — always passed: there
    was no text the lint could see.
    """
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    for z, shape in enumerate(shapes, 1):
        path = f"{prefix}{z}"
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _placed_shapes(shape.shapes, _group_transform(shape, transform), f"{path}.")
            continue
        if shape.left is None or shape.top is None or shape.width is None or shape.height is None:
            continue
        tx, ty, sx, sy = transform
        box = (
            (tx + sx * shape.left) / _EMU_PER_IN,
            (ty + sy * shape.top) / _EMU_PER_IN,
            sx * shape.width / _EMU_PER_IN,
            sy * shape.height / _EMU_PER_IN,
        )
        yield _Placed(shape=shape, box=box, path=path, scale=(sx, sy))


def _where(slide: int, placed: _Placed) -> str:
    """A location the model can go to: slide, shape path, and the shape's name,
    which is what ``reshape_pptx``'s ``shape_name`` addresses inside a group."""
    name = str(getattr(placed.shape, "name", "") or "")[:40]
    return f"slide {slide}, shape {placed.path}" + (f" ({name})" if name else "")


#: How far past the slide's bottom edge a table must run before it is
#: reported: rounding in the row-height estimate stays inside it.
TABLE_OVERFLOW_TOLERANCE_IN = 0.05


def _table_finding(placed: _Placed, slide: int, canvas_h: float) -> Finding | None:
    """Rows that leave the slide because the table grew to fit its text.

    A table row is at least its declared height and grows to fit its text, so
    a table whose frame fits the slide can still push its last rows off the
    bottom — a 16-row table lost rows 9–16 that way, and the frame-only
    canvas check passed it. Each row is measured at the size of its tallest
    cell, the same CJK-aware estimate the text-fit check uses.
    """
    from pptx.util import Emu

    table = placed.shape.table
    rows = list(table.rows)
    if not rows:
        return None
    sx, sy = placed.scale
    col_w = [Emu(c.width).inches * sx for c in table.columns]
    top = placed.box[1]
    first_out: int | None = None
    for r, row in enumerate(rows):
        need = 0.0
        for c, cell in enumerate(row.cells):
            if cell.is_spanned or c >= len(col_w):
                continue
            text = cell.text_frame.text.strip()
            if not text:
                continue
            width = sum(col_w[c:c + max(cell.span_width, 1)])
            margins_w = sum((m or 0) for m in (cell.margin_left, cell.margin_right)) / _EMU_PER_IN * sx
            margins_h = sum((m or 0) for m in (cell.margin_top, cell.margin_bottom)) / _EMU_PER_IN * sy
            sizes = _frame_run_sizes(cell.text_frame)
            pt = max((s for s, _, _ in sizes), default=18.0)
            _, text_h = estimate_text_block(text, font_pt=pt, width_in=max(width - margins_w, 0.2))
            # A cell spanning several rows spreads its text over all of them.
            need = max(need, (text_h + margins_h) / max(cell.span_height, 1))
        bottom = top + max(Emu(row.height).inches * sy, need)
        if first_out is None and bottom > canvas_h + TABLE_OVERFLOW_TOLERANCE_IN:
            first_out = r + 1
        top = bottom
    if first_out is None:
        return None
    count = len(rows)
    lost = f"row {first_out}" if first_out == count else f"rows {first_out}–{count}"
    return Finding(
        "table_overflow", "critical", _where(slide, placed),
        f"{lost} of a {count}-row table run past the slide: the rows need "
        f"~{top - placed.box[1]:.2f}in from {placed.box[1]:.2f}in on a {canvas_h:.2f}in canvas",
        "Split the table across slides, cut or merge rows, or set smaller type; rows grow to fit "
        "their text, so the frame's own height does not bound the table.",
        slide=slide, shape=placed.index,
    )


def _ink_box(placed: _Placed) -> tuple[float, float, float, float]:
    """The part of an unfilled text shape that its glyphs can actually cover.

    An unfilled text box is invisible except for its text, and its box carries
    the line spacing above and below the glyphs. ppt-master sets a two-line
    display title as two such boxes whose line boxes overlap by a fifth while
    the lettering clears by a clean gap, so comparing boxes reported a
    collision a reader cannot see. A filled shape is its whole box.
    """
    shape, (x, y, w, h) = placed.shape, placed.box
    if not getattr(shape, "has_text_frame", False) or _has_visible_fill(shape):
        return placed.box
    text = shape.text_frame.text.strip()
    runs = _run_sizes(shape) if text else []
    if not runs:
        return placed.box
    pt = max(size for size, _, _ in runs)
    frame = shape.text_frame
    top_in = (frame.margin_top.inches if frame.margin_top is not None else 0.05) * placed.scale[1]
    bottom_in = (frame.margin_bottom.inches if frame.margin_bottom is not None else 0.05) * placed.scale[1]
    side_in = sum(
        (m.inches if m is not None else 0.1) for m in (frame.margin_left, frame.margin_right)
    ) * placed.scale[0]
    _, block_h = estimate_text_block(
        text, font_pt=pt, width_in=max(w - side_in, 0.2), wrap=frame.word_wrap is not False
    )
    # The last line contributes its glyphs, not the leading after it.
    ink_h = min(max(block_h - pt * (LINE_HEIGHT - 1.0) / 72.0, pt / 72.0), h)
    anchor = str(getattr(frame, "vertical_anchor", None) or "").upper()
    if "MIDDLE" in anchor:
        top = y + (h - ink_h) / 2.0
    elif "BOTTOM" in anchor:
        top = y + h - bottom_in - ink_h
    else:
        top = y + top_in
    return (x, top, w, ink_h)


def lint_pptx(path: str | Path) -> tuple[list[Finding], dict[str, Any]]:
    """Geometry, typography, contrast and table fit, read straight off the deck."""
    from pptx import Presentation
    from pptx.util import Emu

    prs = Presentation(str(path))
    canvas_w = Emu(prs.slide_width).inches
    canvas_h = Emu(prs.slide_height).inches
    findings: list[Finding] = []
    shape_total = 0
    nested_total = 0

    for index, slide in enumerate(prs.slides, 1):
        boxes: list[tuple[tuple[float, float, float, float], Any, _Placed]] = []
        for placed in _placed_shapes(slide.shapes):
            shape, box = placed.shape, placed.box
            shape_total += 1
            nested_total += "." in placed.path
            where = _where(index, placed)
            boxes.append((box, shape, placed))

            x, y, w, h = box
            if x < -0.02 or y < -0.02 or x + w > canvas_w + 0.02 or y + h > canvas_h + 0.02:
                findings.append(Finding(
                    "shape_off_canvas", "critical", where,
                    f"{shape.shape_type} at ({x:.2f}, {y:.2f}) {w:.2f}x{h:.2f}in "
                    f"leaves the {canvas_w:.2f}x{canvas_h:.2f}in canvas",
                    "Move or resize the shape inside the canvas; coordinates past the edge are "
                    "written, not clamped, so the shape is simply not on the slide.",
                    slide=index, shape=placed.index,
                ))

            if getattr(shape, "has_table", False):
                overflow = _table_finding(placed, index, canvas_h)
                if overflow is not None:
                    findings.append(overflow)
                continue

            if not shape.has_text_frame:
                continue
            text = shape.text_frame.text.strip()
            if not text:
                if (_area(box) >= EMPTY_SHAPE_MIN_AREA_IN2
                        and min(w, h) >= EMPTY_SHAPE_MIN_EDGE_IN
                        and not _has_visible_fill(shape)
                        # A diagonal connector's bounding box is as large as
                        # the rectangle it spans and its text frame is empty,
                        # so without this it reads as a hole in the layout.
                        and not _is_stroke(shape)):
                    findings.append(Finding(
                        "empty_text_box", "major", where,
                        f"an empty {w:.2f}x{h:.2f}in text box holds {_area(box):.1f}in² of the slide",
                        "Fill it or delete it — an unfilled placeholder is the dead space a reader "
                        "sees as an unfinished slide.",
                        slide=index, shape=placed.index,
                    ))
                continue

            runs = _run_sizes(shape)
            for pt, bold, run_text in runs:
                # A short box is a label, a chip or a footer, not body copy.
                floor = NOTE_FONT_FLOOR_PT if h < 0.7 else BODY_FONT_FLOOR_PT
                if pt < floor:
                    findings.append(Finding(
                        "type_too_small", "major", where,
                        f"{pt:.0f}pt run {run_text.strip()[:24]!r} (floor {floor:.0f}pt)",
                        f"Raise the run to at least {floor:.0f}pt.",
                        slide=index, shape=placed.index,
                    ))
                fg = None
                for para in shape.text_frame.paragraphs:
                    for run in para.runs:
                        if run.text == run_text:
                            fg = _run_colour(run)
                            break
                bg = _shape_fill_colour(shape)
                if fg and bg:
                    ratio = contrast_ratio(fg, bg)
                    need = contrast_floor(pt, bold)
                    if ratio < need:
                        findings.append(Finding(
                            "low_contrast", "major", where,
                            f"{ratio:.1f}:1 for {pt:.0f}pt text (WCAG needs {need:.1f}:1)",
                            "Darken the text or lighten the fill until the ratio clears the floor.",
                            slide=index, shape=placed.index,
                        ))

            if runs:
                biggest = max(pt for pt, _, _ in runs)
                # Read the real insets rather than assuming PowerPoint's
                # defaults: a box written with zero insets and sized to its
                # line box would otherwise be charged 0.1in it does not have.
                frame = shape.text_frame
                inset_w = sum(
                    (m.inches if m is not None else 0.1)
                    for m in (frame.margin_left, frame.margin_right)
                ) * placed.scale[0]
                inset_h = sum(
                    (m.inches if m is not None else 0.05)
                    for m in (frame.margin_top, frame.margin_bottom)
                ) * placed.scale[1]
                needed_w, needed_h = estimate_text_block(
                    text, font_pt=biggest, width_in=max(w - inset_w, 0.2),
                    wrap=shape.text_frame.word_wrap is not False,
                )
                if shape.text_frame.auto_size is None or int(getattr(shape.text_frame.auto_size, "real", 0) or 0) != 1:
                    ratio = needed_h / max(h - inset_h, 0.05)
                    if ratio > OVERFLOW_WARN:
                        # Only a filled shape clips its text; a bare text box
                        # spills, which is ugly rather than lossy.
                        blocking = ratio > OVERFLOW_BLOCK and _has_visible_fill(shape)
                        findings.append(Finding(
                            "text_overflow", "critical" if blocking else "major", where,
                            f"{needed_h:.2f}in of text at {biggest:.0f}pt in a {h:.2f}in box "
                            f"({ratio:.0%} of the room, {text[:24]!r})",
                            "Shorten the text, enlarge the box, or drop the type size.",
                            slide=index, shape=placed.index,
                        ))
                    elif needed_w > max(w - inset_w, 0.2) * OVERFLOW_WARN and shape.text_frame.word_wrap is False:
                        findings.append(Finding(
                            "text_overflow", "major", where,
                            f"{needed_w:.2f}in of unwrapped text in a {w:.2f}in box",
                            "Enable word wrap or widen the box.",
                            slide=index, shape=placed.index,
                        ))

        # Once per shape, not per pair: a busy slide has hundreds of pairs.
        inks = [_ink_box(placed) for _, _, placed in boxes]
        for i, (_, shape_a, _) in enumerate(boxes):
            box_a = inks[i]
            for j in range(i + 1, len(boxes)):
                _, shape_b, placed_b = boxes[j]
                box_b = inks[j]
                overlap = _intersection(box_a, box_b)
                if overlap <= 0:
                    continue
                if _contains(box_a, box_b) or _contains(box_b, box_a):
                    continue  # deliberate layering: a card behind its own text
                smaller = min(_area(box_a), _area(box_b))
                if smaller <= 0 or overlap / smaller <= OVERLAP_TOLERANCE:
                    continue
                # Rules, dividers and connectors are drawn under or between the
                # shapes they join, so their overlap is the design.
                if (min(box_a[2], box_a[3]) < HAIRLINE_IN
                        or min(box_b[2], box_b[3]) < HAIRLINE_IN
                        or _is_stroke(shape_a) or _is_stroke(shape_b)):
                    continue
                findings.append(Finding(
                    "shape_overlap", "major", _where(index, placed_b),
                    f"{shape_a.shape_type} and {shape_b.shape_type} overlap by "
                    f"{overlap / smaller:.0%} of the smaller shape",
                    "Separate them — overlapping filled shapes read as a layout error, not a design.",
                    slide=index, shape=placed_b.index,
                ))

    return findings, {
        "canvas_in": [round(canvas_w, 3), round(canvas_h, 3)],
        "slides": len(prs.slides._sldIdLst),
        "shapes": shape_total,
        "shapes_in_groups": nested_total,
    }


# ---------------------------------------------------------------------------
# Renders
# ---------------------------------------------------------------------------


#: A background this far from white is a deliberate panel, not paper.
FULL_BLEED_DELTA = 26


@dataclass
class PageInk:
    coverage: float
    ink: float
    dead_rows: float
    dead_band: float
    top: float
    bottom: float
    left: float
    right: float
    #: True when the page's own background is a colour rather than paper. The
    #: ink metric then measures only what sits *on* the panel, so a cover slide
    #: whose whole canvas is filled reads as 11% covered — which is how a
    #: correct full-bleed cover burned a repair round.
    full_bleed: bool = False


def page_ink(path: str | Path, *, threshold: int = 18) -> PageInk:
    """Where the ink is on one rendered page.

    The background is the modal value of the outer border ring rather than the
    four corners — a corner can land inside a full-bleed panel, and one wrong
    corner inverts every number downstream.
    """
    from PIL import Image

    im = Image.open(str(path)).convert("L")
    w, h = im.size
    px = im.load()
    ring = Counter()
    for x in range(0, w, max(w // 200, 1)):
        ring[px[x, 0]] += 1
        ring[px[x, h - 1]] += 1
    for y in range(0, h, max(h // 200, 1)):
        ring[px[0, y]] += 1
        ring[px[w - 1, y]] += 1
    bg = ring.most_common(1)[0][0]
    full_bleed = abs(bg - 255) > FULL_BLEED_DELTA

    step = max(min(w, h) // 400, 1)
    rows = h // step + 1
    row_ink = [0] * rows
    ink = total = 0
    minx, miny, maxx, maxy = w, h, -1, -1
    for y in range(0, h, step):
        for x in range(0, w, step):
            total += 1
            if abs(px[x, y] - bg) > threshold:
                ink += 1
                row_ink[y // step] += 1
                minx, miny = min(minx, x), min(miny, y)
                maxx, maxy = max(maxx, x), max(maxy, y)
    if maxx < 0:
        return PageInk(0.0, 0.0, 1.0, 1.0, 1.0, 0.0, 1.0, 0.0, full_bleed)

    longest = run = 0
    for count in row_ink:
        run = run + 1 if count == 0 else 0
        longest = max(longest, run)
    return PageInk(
        coverage=((maxx - minx) * (maxy - miny)) / float(w * h),
        ink=ink / float(total or 1),
        dead_rows=sum(1 for c in row_ink if c == 0) / float(rows),
        dead_band=longest / float(rows),
        top=miny / h, bottom=maxy / h, left=minx / w, right=maxx / w,
        full_bleed=full_bleed,
    )


def lint_renders(
    paths: Sequence[str | Path], *, label: str = "slide"
) -> tuple[list[Finding], dict[str, Any]]:
    """Report where the ink is. Deliberately produces no findings.

    A check earns its place here when a pass means "nothing is provably
    broken", and loses it the moment a pass could be read as "this is good".
    Every occupancy measure fails that test: each one can be passed by padding.
    """
    findings: list[Finding] = []
    pages: list[PageInk] = []
    for index, path in enumerate(paths, 1):
        try:
            metrics = page_ink(path)
        except Exception:
            continue
        pages.append(metrics)
    # No findings from this stage any more. Coverage, overfill and dead-band
    # were all measures of pixel occupancy, and occupancy is a proxy for effort
    # that is satisfied by making shapes bigger — which is exactly what
    # produced the deck its author called empty. Tested over 98 reference
    # slides: `canvas_underfilled` ranked our emptiest slide (0.81, three cards
    # holding 45 characters each) above the reference set's densest statutory
    # slide (0.66, ten articles with a holding for each), because a pale 6.7in
    # card is ink; `canvas_overfilled` fired four times and all four were
    # correct work; `dead_band` sat 0.01 from flagging eleven correct section
    # dividers, and its advice — fill the height — is the opposite of what a
    # divider wants. The numbers are still reported, but nothing is graded on
    # them.
    measured = [p for p in pages if not p.full_bleed]
    metrics: dict[str, Any] = {f"{label}s_measured": len(pages)}
    if len(measured) != len(pages):
        metrics["full_bleed_pages"] = len(pages) - len(measured)
    if measured:
        metrics |= {
            "coverage_mean": round(sum(p.coverage for p in measured) / len(measured), 3),
            "coverage_min": round(min(p.coverage for p in measured), 3),
            "dead_band_max": round(max(p.dead_band for p in measured), 3),
        }
    return findings, metrics


# ---------------------------------------------------------------------------
# DOCX
# ---------------------------------------------------------------------------


def _effective_first_line_indent(para: Any, styles: Any) -> float:
    """First-line indent in twips, resolved through the style chain.

    Reading only direct formatting is how this check scored zero on a memo
    whose table cells provably wrapped a statute title across two lines: the
    indent lived on ``Normal``, and every heading and cell paragraph inherited
    it through ``basedOn`` without carrying a ``w:ind`` of its own.
    """
    from docx.oxml.ns import qn

    def own(element: Any) -> float | None:
        pPr = element.find(qn("w:pPr")) if element is not None else None
        node = pPr.find(qn("w:ind")) if pPr is not None else None
        if node is None:
            return None
        chars = node.get(qn("w:firstLineChars"))
        if chars is not None:
            return float(chars) / 100.0 * 240.0  # hundredths of a char → twips at 12pt
        raw = node.get(qn("w:firstLine"))
        return float(raw) if raw is not None else None

    direct = own(para._p)
    if direct is not None:
        return direct
    style = para.style
    seen = set()
    while style is not None and style.style_id not in seen:
        seen.add(style.style_id)
        value = own(style.element)
        if value is not None:
            return value
        style = getattr(style, "base_style", None)
    try:
        return own(styles["Normal"].element) or 0.0
    except Exception:
        return 0.0


# --- Markdown residue (DOCX and PDF text) ----------------------------------

#: Markdown that reached a document as characters. Each one is a writer that
#: did not understand its input: 517 literal ``[text](url)`` links in 28 DOCX,
#: ``####`` headings in 5. Nothing in a legal instrument is written this way on
#: purpose, so a match is a defect rather than a style call.
_RESIDUE_PATTERNS: tuple[tuple[str, re.Pattern[str]], ...] = (
    ("**bold** markers", re.compile(r"\*\*[^*\s][^*\n]{0,200}?\*\*")),
    ("[text](url) link", re.compile(r"\[[^\]\n]{1,200}\]\((?:https?://|www\.)")),
    ("#### heading marks", re.compile(r"^\s{0,3}#{1,6}\s+\S", re.M)),
    ("|---| table rule", re.compile(r"\|\s*:?-{3,}:?\s*\|")),
    ("[^1] footnote marker", re.compile(r"\[\^[^\]\s]{1,20}\]")),
    ("~~strike~~ markers", re.compile(r"~~[^~\n]{1,200}~~")),
    ("<br> tag", re.compile(r"<\s*br\s*/?\s*>", re.I)),
)
_RESIDUE_FIX = (
    "Write the content again through write_docx / write_pdf, which render Markdown; "
    "do not paste Markdown source into an edit."
)


def markdown_residue(text: str) -> list[str]:
    """Names of the Markdown constructs present as literal text."""
    return [name for name, pattern in _RESIDUE_PATTERNS if pattern.search(text or "")]


def _residue_findings(text: str, where: str) -> list[Finding]:
    found = markdown_residue(text)
    if not found:
        return []
    snippet = " ".join((text or "").split())[:40]
    return [Finding(
        "markdown_residue", "major", where,
        f"literal {', '.join(found)} in {snippet!r}",
        _RESIDUE_FIX,
    )]


# --- List numbering --------------------------------------------------------

#: A level whose text is only its number ("1." "(1)" "1、") is a list. One that
#: carries words ("第%1条", "Article %1") or several levels ("%1.%2") is clause
#: numbering, which is meant to run on across chapter headings.
_PLAIN_LIST_TEXT_RE = re.compile(r"^[(（]?%\d[.)）、．]?$")
_BULLET_FORMATS = {"bullet", "none"}


def _numbering_model(doc: Any) -> tuple[dict[str, dict[int, tuple[str, int, str]]], dict[str, tuple[str, dict[int, int]]]]:
    """``(abstract levels, nums)`` from the numbering part.

    ``abstract levels`` maps abstractNumId → ilvl → (numFmt, start, lvlText);
    ``nums`` maps numId → (abstractNumId, {ilvl: startOverride}).
    """
    from docx.oxml.ns import qn

    rel = next(
        (r for r in doc.part.rels.values()
         if r.reltype.endswith("/numbering") and not r.is_external),
        None,
    )
    if rel is None:
        return {}, {}
    part = rel.target_part
    root = getattr(part, "element", None)
    if root is None:
        from lxml import etree

        root = etree.fromstring(part.blob)
    abstracts: dict[str, dict[int, tuple[str, int, str]]] = {}
    for abstract in root.findall(qn("w:abstractNum")):
        levels: dict[int, tuple[str, int, str]] = {}
        for lvl in abstract.findall(qn("w:lvl")):
            fmt = lvl.find(qn("w:numFmt"))
            start = lvl.find(qn("w:start"))
            text = lvl.find(qn("w:lvlText"))
            levels[int(lvl.get(qn("w:ilvl"), "0"))] = (
                fmt.get(qn("w:val")) if fmt is not None else "decimal",
                int(start.get(qn("w:val"))) if start is not None else 1,
                text.get(qn("w:val")) if text is not None else "",
            )
        abstracts[abstract.get(qn("w:abstractNumId"))] = levels
    nums: dict[str, tuple[str, dict[int, int]]] = {}
    for num in root.findall(qn("w:num")):
        abstract_id = num.find(qn("w:abstractNumId"))
        overrides: dict[int, int] = {}
        for override in num.findall(qn("w:lvlOverride")):
            start = override.find(qn("w:startOverride"))
            if start is not None:
                overrides[int(override.get(qn("w:ilvl"), "0"))] = int(start.get(qn("w:val")))
        nums[num.get(qn("w:numId"))] = (
            abstract_id.get(qn("w:val")) if abstract_id is not None else "",
            overrides,
        )
    return abstracts, nums


def _paragraph_numbering(para: Any) -> tuple[str, int] | None:
    """``(numId, ilvl)`` from the paragraph or its style chain."""
    from docx.oxml.ns import qn

    def own(element: Any) -> tuple[str | None, int | None] | None:
        ppr = element.find(qn("w:pPr")) if element is not None else None
        num_pr = ppr.find(qn("w:numPr")) if ppr is not None else None
        if num_pr is None:
            return None
        num_id = num_pr.find(qn("w:numId"))
        ilvl = num_pr.find(qn("w:ilvl"))
        return (
            num_id.get(qn("w:val")) if num_id is not None else None,
            int(ilvl.get(qn("w:val"))) if ilvl is not None else None,
        )

    found = own(para._p)
    num_id, ilvl = found if found else (None, None)
    style = para.style
    seen: set[str] = set()
    while (num_id is None or ilvl is None) and style is not None and style.style_id not in seen:
        seen.add(style.style_id)
        inherited = own(style.element)
        if inherited:
            num_id = num_id if num_id is not None else inherited[0]
            ilvl = ilvl if ilvl is not None else inherited[1]
        style = getattr(style, "base_style", None)
    if not num_id or num_id == "0":
        return None
    return num_id, ilvl or 0


def _is_heading(para: Any) -> bool:
    name = (getattr(para.style, "name", "") or "").lower()
    return name.startswith("heading") or name == "title"


def _list_continuation_findings(doc: Any) -> list[Finding]:
    """An ordered list that keeps counting after a heading instead of restarting.

    Modelled on Word, which counts per abstract numbering definition: a new
    ``w:num`` restarts only when it carries a ``w:startOverride``. The writer's
    single ``List Number`` style used to number 62 of 115 documents' lists on
    from the previous section — the first item of a list showed a median of 19
    — and in a filing "诉讼请求第3项" means item 3, not item 19.
    """
    abstracts, nums = _numbering_model(doc)
    if not nums:
        return []
    counters: dict[str, dict[int, int]] = {}
    seen_nums: set[str] = set()
    last_item: dict[tuple[str, int], int] = {}
    last_heading = -1
    heading_text = ""
    findings: list[Finding] = []
    for index, para in enumerate(doc.paragraphs):
        if _is_heading(para):
            last_heading, heading_text = index, para.text.strip()
            continue
        numbering = _paragraph_numbering(para)
        if numbering is None:
            continue
        num_id, ilvl = numbering
        abstract_id, overrides = nums.get(num_id, ("", {}))
        levels = abstracts.get(abstract_id)
        if not levels:
            continue
        fmt, start, text = levels.get(ilvl, ("decimal", 1, f"%{ilvl + 1}."))
        counter = counters.setdefault(abstract_id, {})
        if num_id not in seen_nums:
            seen_nums.add(num_id)
            for level, value in overrides.items():
                counter[level] = value - 1
        for deeper in [lvl for lvl in counter if lvl > ilvl]:
            counter[deeper] = levels.get(deeper, ("decimal", 1, ""))[1] - 1
        counter.setdefault(ilvl, start - 1)
        counter[ilvl] += 1
        value = counter[ilvl]
        key = (abstract_id, ilvl)
        previous = last_item.get(key)
        last_item[key] = index
        if fmt in _BULLET_FORMATS or not _PLAIN_LIST_TEXT_RE.match(text or ""):
            continue
        if previous is not None and previous < last_heading < index and value > start:
            findings.append(Finding(
                "list_numbering_continues", "critical", f"paragraph {index + 1}",
                f"{para.text.strip()[:24]!r} is numbered {value}: the list continues from before "
                f"heading {heading_text[:24]!r} instead of starting again",
                "Give each list its own w:num with a w:startOverride (write_docx does this), or "
                "restart numbering at this item.",
            ))
    return findings


# --- Tables across pages ---------------------------------------------------


def _pdf_page_texts(pdf: str | Path) -> list[str]:
    """Each page's text with all whitespace removed, for substring probes."""
    import pypdfium2 as pdfium

    document = pdfium.PdfDocument(str(pdf))
    try:
        texts = []
        for page in document:
            textpage = page.get_textpage()
            texts.append("".join(textpage.get_text_range().split()))
        return texts
    finally:
        document.close()


def _row_probe(row: Any) -> str:
    texts = ["".join(cell.text.split()) for cell in row.cells]
    longest = max(texts, key=len, default="")
    return longest[:12]


def _table_page_findings(doc: Any, pages: list[str]) -> list[Finding]:
    """A table that runs onto another page without a repeating header row."""
    from docx.oxml.ns import qn

    findings: list[Finding] = []
    for t_index, table in enumerate(doc.tables, 1):
        rows = table.rows
        if len(rows) < 2:
            continue
        tr_pr = rows[0]._tr.find(qn("w:trPr"))
        if tr_pr is not None and tr_pr.find(qn("w:tblHeader")) is not None:
            continue
        head, tail = _row_probe(rows[0]), _row_probe(rows[-1])
        if len(head) < 2 or len(tail) < 2:
            continue
        first = next((i for i, text in enumerate(pages) if head in text), None)
        if first is None:
            continue
        last = next((i for i in range(first, len(pages)) if tail in pages[i]), None)
        if last is None or last == first:
            continue
        findings.append(Finding(
            "table_header_not_repeated", "major", f"table {t_index}",
            f"runs from page {first + 1} to page {last + 1} and its header row does not repeat, "
            "so the later page shows columns with no names",
            "Mark the first row as a repeating header row (w:tblHeader); write_docx does this "
            "for every table.",
        ))
    return findings


def _is_gbt9704_page(doc: Any) -> bool:
    """GB/T 9704 page: 天头 37±1 mm and 订口 28±1 mm, where headings are 空二字."""
    section = doc.sections[0]
    try:
        top = section.top_margin.mm
        left = section.left_margin.mm
    except AttributeError:
        return False
    return 36.0 <= top <= 38.0 and 27.0 <= left <= 29.0


def lint_docx(path: str | Path, *, pdf: str | Path | None = None) -> tuple[list[Finding], dict[str, Any]]:
    """Checks for the defects that actually show up in generated memos.

    ``pdf`` is the document's render, when one exists; the page-crossing table
    check needs it, because only a layout knows where a page ends.
    """
    from docx import Document
    from docx.oxml.ns import qn

    from .writers._text import display_width

    doc = Document(str(path))
    findings: list[Finding] = []

    for t_index, table in enumerate(doc.tables, 1):
        where = f"table {t_index}"
        indented = 0
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    if _effective_first_line_indent(para, doc.styles) > 0:
                        indented += 1
        if indented:
            findings.append(Finding(
                "cell_first_line_indent", "major", where,
                f"{indented} cell paragraph(s) carry a first-line indent",
                "Zero w:ind on paragraphs inside w:tbl — a cell is one short field, and the indent "
                "pushes its text into a second line.",
            ))

        grid = table._tbl.find(qn("w:tblGrid"))
        widths = [int(c.get(qn("w:w")) or 0) for c in grid.findall(qn("w:gridCol"))] if grid is not None else []
        if len(widths) > 1 and len(set(widths)) == 1:
            demand = []
            for col in range(len(widths)):
                longest = 0.0
                for row in table.rows:
                    if col < len(row.cells):
                        longest = max(longest, display_width(row.cells[col].text.strip()))
                demand.append(longest or 1.0)
            if max(demand) / max(min(demand), 1e-6) >= 2.5:
                findings.append(Finding(
                    "table_columns_unweighted", "major", where,
                    "every column has the same width while content demand varies "
                    f"{max(demand) / max(min(demand), 1e-6):.1f}x",
                    "Weight w:gridCol by the CJK-aware content width so long citations get the room "
                    "and short codes do not hoard it.",
                ))

    headings = [p for p in doc.paragraphs if (p.style.name or "").lower().startswith("heading")]
    # GB/T 9704 7.3.3 sets every structure level 空二字, so there the indent is
    # the standard, not a leak of the body rule.
    gbt9704 = _is_gbt9704_page(doc)
    for index, para in enumerate(headings):
        if not gbt9704 and _effective_first_line_indent(para, doc.styles) > 0:
            findings.append(Finding(
                "heading_indent", "minor", f"heading {index + 1}",
                f"{para.text.strip()[:20]!r} carries a first-line indent",
                "Headings sit on the margin; the CJK body indent must not reach them.",
            ))

    # In a 公文 the 1./（1） levels are run-in headings that carry their own text,
    # so consecutive ones are items, not empty sections.
    body = [] if gbt9704 else [p for p in doc.paragraphs if p.text.strip()]
    for a, b in zip(body, body[1:]):
        if ((a.style.name or "").lower().startswith("heading")
                and (b.style.name or "").lower().startswith("heading")
                and (a.style.name or "") == (b.style.name or "")):
            findings.append(Finding(
                "empty_section", "minor", a.text.strip()[:24],
                "a heading is immediately followed by a sibling heading with no body",
                "Write the section or drop the heading.",
            ))

    findings += _list_continuation_findings(doc)

    for index, para in enumerate(doc.paragraphs, 1):
        if _is_code_paragraph(para):
            continue
        findings += _residue_findings(para.text, f"paragraph {index}")
    for t_index, table in enumerate(doc.tables, 1):
        for r_index, row in enumerate(table.rows, 1):
            for c_index, cell in enumerate(row.cells, 1):
                findings += _residue_findings(cell.text, f"table {t_index} row {r_index} cell {c_index}")

    metrics: dict[str, Any] = {"tables": len(doc.tables), "paragraphs": len(doc.paragraphs)}
    if pdf is not None:
        try:
            pages = _pdf_page_texts(pdf)
            metrics["pages"] = len(pages)
            findings += _table_page_findings(doc, pages)
        except Exception as exc:  # noqa: BLE001 - the render is optional evidence
            metrics["page_check_error"] = f"{type(exc).__name__}: {exc}"
    return findings, metrics


_MONO_FONTS = ("courier", "consolas", "menlo", "mono")


def _is_code_paragraph(para: Any) -> bool:
    """A code listing may show Markdown on purpose; leave it alone."""
    runs = [r for r in para.runs if r.text.strip()]
    return bool(runs) and all(
        any(mono in (r.font.name or "").lower() for mono in _MONO_FONTS) for r in runs
    )


# ---------------------------------------------------------------------------
# XLSX
# ---------------------------------------------------------------------------

#: Per-check ceiling on findings built for one workbook; the count is exact.
_XLSX_FINDING_CAP = 200


_DECIMALS_RE = re.compile(r"0\.(0+)")
#: Cells read for the width check; a ledger is far below it, a data dump is not
#: worth a slow render for.
_XLSX_WIDTH_SCAN_CAP = 50_000


def _number_text(value: float, number_format: str) -> str:
    """The text Excel shows for a number under a common format (approximate)."""
    fmt = number_format or "General"
    m = _DECIMALS_RE.search(fmt)
    decimals = len(m.group(1)) if m else 0
    if "%" in fmt:
        return f"{value * 100:.{decimals}f}%"
    if fmt == "General":
        text = f"{value:.10g}"
    elif "," in fmt:
        text = f"{value:,.{decimals}f}"
    else:
        text = f"{value:.{decimals}f}"
    return ("¥" if "¥" in fmt else "$" if "$" in fmt else "") + text


def _excel_text_width(text: str, bold: bool) -> float:
    """Width in Excel's column units (one unit = one digit of the body font)."""
    width = sum(0.5 if ch in ",.:;'" else 2.0 if ord(ch) > 0x2E7F else 1.0 for ch in text)
    return width * (1.08 if bold else 1.0) + 0.75


def lint_xlsx(path: str | Path) -> tuple[list[Finding], dict[str, Any]]:
    """Error values, formulas without a cached value, and numbers wider than their column.

    All three are read off the file: an error cell is ``#REF!``/``#DIV/0!``/…
    in the cached value a reader sees; a formula with no cached value shows as
    blank to every reader but Excel (pandas, previews, ``read_document``); and
    a number whose formatted text is wider than its column shows as ``###`` in
    Excel. The last is measured in the workbook's own units, because a preview
    drawn with stand-in fonts at LibreOffice's column conversion shows ``###``
    where Excel does not, and the reverse.
    """
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter

    from .writers._xlsx_recalc import ERROR_TOKENS

    values_wb = load_workbook(str(path), data_only=True, read_only=True)
    formulas_wb = load_workbook(str(path), data_only=False, read_only=True)
    findings: list[Finding] = []
    errors = uncached = formulas = too_narrow = 0
    try:
        # Column widths and fonts need the full (non-streaming) reader.
        styled_wb = load_workbook(str(path), data_only=True)
    except Exception:  # noqa: BLE001 — the width check is best effort
        styled_wb = None
    scanned = 0
    if styled_wb is not None:
        for ws in styled_wb.worksheets:
            default_width = (ws.sheet_format.defaultColWidth or 8.43) if ws.sheet_format else 8.43
            # A <col min max> span is stored once, under its first letter.
            widths: dict[str, float] = {}
            for dim in ws.column_dimensions.values():
                if dim.width and dim.min and dim.max:
                    for index in range(dim.min, dim.max + 1):
                        widths[get_column_letter(index)] = dim.width
            for row in ws.iter_rows():
                if scanned > _XLSX_WIDTH_SCAN_CAP:
                    break
                for cell in row:
                    scanned += 1
                    value = cell.value
                    if isinstance(value, bool) or not isinstance(value, (int, float)):
                        continue
                    letter = get_column_letter(cell.column)
                    widths.setdefault(letter, default_width)
                    shown = _number_text(float(value), cell.number_format)
                    need = _excel_text_width(shown, bool(cell.font and cell.font.b))
                    if need > widths[letter] and not any(
                        cell.coordinate in rng for rng in ws.merged_cells.ranges
                    ):
                        too_narrow += 1
                        if too_narrow <= _XLSX_FINDING_CAP:
                            findings.append(Finding(
                                "number_wider_than_column", "major", f"{ws.title}!{cell.coordinate}",
                                f"{shown!r} needs about {need:.1f} of column {letter}'s width {widths[letter]:g}; "
                                "Excel shows ###",
                                f"Widen column {letter} (reshape_xlsx set_col_width).",
                            ))
        try:
            styled_wb.close()
        except Exception:  # noqa: BLE001
            pass
    try:
        for ws_f, ws_v in zip(formulas_wb.worksheets, values_wb.worksheets):
            for row_f, row_v in zip(ws_f.iter_rows(), ws_v.iter_rows()):
                for cell_f, cell_v in zip(row_f, row_v):
                    coordinate = getattr(cell_f, "coordinate", None)
                    if coordinate is None:
                        continue
                    value = cell_v.value
                    is_formula = getattr(cell_f, "data_type", None) == "f" or (
                        isinstance(cell_f.value, str) and cell_f.value.startswith("=")
                    )
                    formulas += is_formula
                    where = f"{ws_f.title}!{coordinate}"
                    if isinstance(value, str) and value in ERROR_TOKENS:
                        errors += 1
                        if errors <= _XLSX_FINDING_CAP:
                            findings.append(Finding(
                                "formula_error_value", "critical", where,
                                f"{value} in {str(cell_f.value)[:40]!r}",
                                "Fix the formula or its inputs; a reader sees the error, not a number.",
                            ))
                    elif is_formula and value is None:
                        uncached += 1
                        if uncached <= _XLSX_FINDING_CAP:
                            findings.append(Finding(
                                "formula_without_cached_value", "major", where,
                                f"{str(cell_f.value)[:40]!r} has no computed value stored",
                                "Write the workbook with write_xlsx (it computes and stores every "
                                "formula's value), or open and save it once in a spreadsheet program.",
                            ))
    finally:
        values_wb.close()
        formulas_wb.close()
    return findings, {
        "sheets": len(formulas_wb.sheetnames),
        "formulas": formulas,
        "error_cells": errors,
        "uncached_formulas": uncached,
        "numbers_wider_than_column": too_narrow,
    }


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------


def lint_pdf(path: str | Path) -> tuple[list[Finding], dict[str, Any]]:
    """Markdown residue in the text a reader of the PDF sees, by page."""
    import pypdfium2 as pdfium

    document = pdfium.PdfDocument(str(path))
    findings: list[Finding] = []
    try:
        pages = len(document)
        for index, page in enumerate(document, 1):
            text = page.get_textpage().get_text_range()
            findings += _residue_findings(text, f"page {index}")
    finally:
        document.close()
    return findings, {"pages": pages}


# ---------------------------------------------------------------------------
# Entry point
# ---------------------------------------------------------------------------


def lint_artifact(
    path: str | Path,
    renders: Sequence[str | Path] | None = None,
    *,
    render_label: str = "slide",
    pdf: str | Path | None = None,
) -> QualityReport:
    """Lint whatever ``path`` is, plus its renders when they are available.

    ``pdf`` is the rendered PDF of an Office file, when the caller has one; the
    DOCX page checks read it. Every check is best effort: a linter that can
    fail a document write would be worse than no linter, so an unreadable file
    yields a report saying so rather than an exception. A file type with no
    checks reports :data:`NO_CHECKS`, never "pass".
    """
    path = Path(path)
    findings: list[Finding] = []
    metrics: dict[str, Any] = {}
    suffix = path.suffix.lower()
    checked = False

    try:
        if suffix == ".pptx":
            f, m = lint_pptx(path)
            findings += f
            metrics |= m
            checked = True
        elif suffix == ".docx":
            f, m = lint_docx(path, pdf=pdf)
            findings += f
            metrics |= m
            checked = True
        elif suffix in {".xlsx", ".xlsm"}:
            f, m = lint_xlsx(path)
            findings += f
            metrics |= m
            checked = True
        elif suffix == ".pdf":
            f, m = lint_pdf(path)
            findings += f
            metrics |= m
            checked = True
    except Exception as exc:
        metrics["file_check_error"] = f"{type(exc).__name__}: {exc}"

    if renders:
        try:
            f, m = lint_renders(renders, label=render_label)
            findings += f
            metrics |= m
        except Exception as exc:
            metrics["render_check_error"] = f"{type(exc).__name__}: {exc}"

    return QualityReport(findings=findings, metrics=metrics, checked=checked)


__all__ = [
    "BODY_FONT_FLOOR_PT",
    "FULL_BLEED_DELTA",
    "Finding",
    "HAIRLINE_IN",
    "NO_CHECKS",
    "RENDER_CHECKS",
    "PageInk",
    "QualityReport",
    "Severity",
    "TABLE_OVERFLOW_TOLERANCE_IN",
    "contrast_floor",
    "contrast_ratio",
    "estimate_text_block",
    "lint_artifact",
    "lint_docx",
    "lint_pdf",
    "lint_pptx",
    "lint_renders",
    "lint_xlsx",
    "markdown_residue",
    "page_ink",
    "relative_luminance",
]
